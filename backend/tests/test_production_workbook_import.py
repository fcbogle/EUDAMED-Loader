"""Production importer fixtures are synthetic and use isolated SQLite/file roots."""
from copy import deepcopy
from pathlib import Path
import json
import sqlite3
from unittest.mock import patch

import lxml.etree as etree
from openpyxl import load_workbook
import pytest

from app.config import get_settings
from app.services.production_import_files import publish_current_pair
from app.services.production_workbook_import import ProductionWorkbookImporter
from app.services.production_workbook_preparation import prepare_workbook, NS
from app.services.xml_projection import DeviceXmlProjectionBuilder
from app.services.xml_rendering import EudamedMessageRenderer
from app.services.testing_state_store import TestingStateStore as StateStore
from app.services.xml_generation import XmlGenerationService
from app.services.registration_summary import RegistrationSummaryService
from scripts.benchmark_xml import synthetic_record
from test_production_workbook_preparation import workbook, child


@pytest.fixture
def production(tmp_path, monkeypatch):
    base=get_settings()
    settings=base.model_copy(update=dict(environment='prod',data_root=tmp_path,
        testing_state_db_path=tmp_path/'application.sqlite3',testing_state_backup_dir=tmp_path/'backups',
        production_import_workbook=tmp_path/'out/production-import.xlsx',production_import_audit=tmp_path/'out/production-import.audit.json',
        eudamed_message_schema_version='3.0.30',schema_dir=base.schema_dir.parent/'prod-3.0.30',
        eudamed_manufacturer_srn_override='GB-MF-000000001',eudamed_suppress_authorised_representative=True,
        eudamed_authorised_representative_srn_override=None))
    monkeypatch.setattr('app.config.get_settings',lambda: settings)
    modules=['production_workbook_import','workbook_import','testing_state_store','canonical_validation','basic_udi_reference',
             'normalization','xml_projection','xml_generation','registration_summary','xml_validation','testing_read_model','canonical_review','workbook_import_selector']
    for module in modules:
        monkeypatch.setattr(f'app.services.{module}.get_settings',lambda: settings)
    for folder in ('templates','xml','out'):
        (tmp_path/folder).mkdir()
    parent=tmp_path/'BasicUDIs.xlsx'
    workbook(parent,['Issuing Entity','Basic UDI-DI code','Device Model','Risk class'],[['GS1','BASICSHARED','Synthetic model','Class IIa']],sheet='Upload')
    workbook(tmp_path/'templates/Template.xlsx',['Issuing Entity e.g. GS1','UDI-DI code','Basic UDI-DI','Reference/ Catalogue number e.g . Taken from Product code in second page of DoC'],
             [['GS1','99000000000001','BASICSHARED','NEW-ONE']],sheet='Synthetic model')
    exported=DeviceXmlProjectionBuilder().build_device_record(synthetic_record(0))
    pushed=etree.fromstring(EudamedMessageRenderer(settings).render_message(exported))
    device=pushed.find('message:payload/device:Device',NS)
    assert device is not None
    basic=device.find('device:MDRBasicUDI',NS);udi=device.find('device:MDRUDIDIData',NS)
    for node in (basic,udi,udi.find('udi:marketInfos',NS)):
        for name,value in (('state','REGISTERED'),('version','2'),('versionDate','2026-09-11T10:00:00Z')):
            existing=node.find('entity:'+name,NS)
            if existing is None: child(node,'entity',name,value)
            else: existing.text=value
    root=etree.Element('{'+NS['message']+'}PullResponse',nsmap=pushed.nsmap)
    child(root,'message','creationDateTime','2026-09-11T10:00:00Z')
    payload=child(root,'message','payload');payload.append(device)
    for name,value in (('pageNumber','0'),('numberOfPages','1'),('pageSize','50')):child(root,'message',name,value)
    (tmp_path/'xml/export.xml').write_bytes(etree.tostring(root,xml_declaration=True,encoding='utf-8'))
    paths=dict(template_dir=tmp_path/'templates',xml_dir=tmp_path/'xml',parent_reference=parent,output_dir=tmp_path/'out')
    def prepare():
        result=prepare_workbook(**paths)
        publish_current_pair(Path(result['workbook']),Path(result['audit']))
        return result
    prepare()
    ProductionWorkbookImporter._pending.clear()
    return settings,paths,prepare


def counts(settings):
    with sqlite3.connect(settings.testing_state_db_path) as c:
        return tuple(c.execute('SELECT COUNT(*) FROM '+t).fetchone()[0] for t in ('import_batch','source_row','device_subject','canonical_device_record','testing_events'))


def test_readonly_assessment_and_atomic_baseline_import(production):
    settings,_,_=production
    assessment=ProductionWorkbookImporter().assess()
    assert not settings.testing_state_db_path.exists()
    assert assessment['eligible_count']==2 and assessment['new_count']==2
    result=ProductionWorkbookImporter().run_import(assessment_token=assessment['assessment_token'])
    assert result['created_count']==2 and counts(settings)==(1,2,2,2,1)
    with sqlite3.connect(settings.testing_state_db_path) as c:
        assert c.execute('SELECT status,event_kind,message_type FROM testing_events').fetchone()==('IMPORTED','BASELINE_IMPORT','PRODUCTION_EXPORT.SNAPSHOT')
        assert c.execute('SELECT post_success,latest_successful_post_version,latest_successful_patch_version,latest_successful_version FROM testing_subjects').fetchone()==(0,None,None,'2')
        assert c.execute('SELECT COUNT(*) FROM generated_packages').fetchone()[0]==0
    store=StateStore()
    assert store.has_successful_primary_udi_post(product_family='Synthetic model',product_variant='Synthetic model',primary_udi_di='99000000000000')
    assert len(store.posted_entries(product_family='Synthetic model',product_variant='Synthetic model',basic_udi_di='BASICSHARED'))==1
    summary=RegistrationSummaryService().summary()
    assert summary['counts']['registered_devices']==1 and summary['counts']['awaiting_devices']==1
    assert sum(summary['event_counts'].values())==0


def test_reimport_noop_additive_missing_retention_and_changed_existing_report(production):
    settings,paths,prepare=production
    service=ProductionWorkbookImporter();a=service.assess();service.run_import(assessment_token=a['assessment_token'])
    before=counts(settings);a=service.assess();result=service.run_import(assessment_token=a['assessment_token'])
    assert result['already_imported'] and counts(settings)==before
    workbook(paths['template_dir']/'Template.xlsx',['Issuing Entity e.g. GS1','UDI-DI code','Basic UDI-DI','Reference/ Catalogue number e.g . Taken from Product code in second page of DoC'],
             [['GS1','99000000000002','BASICSHARED','NEW-TWO']],sheet='Synthetic model')
    prepare();a=service.assess()
    assert a['new_count']==1 and a['retained_count']==1
    service.run_import(assessment_token=a['assessment_token'])
    assert counts(settings)==(2,3,3,3,1)
    workbook(paths['template_dir']/'Template.xlsx',['Issuing Entity e.g. GS1','UDI-DI code','Basic UDI-DI','Reference/ Catalogue number e.g . Taken from Product code in second page of DoC'],
             [['GS1','99000000000002','BASICSHARED','CHANGED']],sheet='Synthetic model')
    prepare();a=service.assess();assert not a['can_import'] and len(a['differences'])==1
    before=counts(settings)
    with pytest.raises(ValueError,match='differ'):service.run_import(assessment_token=a['assessment_token'])
    assert counts(settings)==before


def test_stale_confirmation_and_rollback(production,monkeypatch):
    settings,_,_=production
    service=ProductionWorkbookImporter();a=service.assess()
    with pytest.raises(ValueError,match='required'): service.run_import()
    original=service._baseline
    def fail(*args):
        original(*args)
        raise RuntimeError('Synthetic baseline failure')
    monkeypatch.setattr(service,'_baseline',fail)
    with pytest.raises(RuntimeError,match='Synthetic'):service.run_import(assessment_token=a['assessment_token'])
    assert counts(settings)==(0,0,0,0,0)
    a=service.assess()
    settings.production_import_audit.write_text('{}')
    with pytest.raises(ValueError,match='mismatch'):service.run_import(assessment_token=a['assessment_token'])
    assert counts(settings)==(0,0,0,0,0)


def test_pair_locks_and_review_skip(production):
    settings,_,_=production
    book=load_workbook(settings.production_import_workbook)
    page=book['To Register'];keys={c.value:c.column for c in page[1]}
    page.cell(2,keys['Review Required'],'Yes');page.cell(2,keys['Review Reasons'],'Deferred synthetic data')
    book.save(settings.production_import_workbook);book.close()
    audit=json.loads(settings.production_import_audit.read_text())
    next(r for r in audit['output_rows'] if r['tab']=='To Register')['review_reasons']=['Deferred synthetic data']
    import hashlib
    audit['workbook_sha256']=hashlib.sha256(settings.production_import_workbook.read_bytes()).hexdigest()
    settings.production_import_audit.write_text(json.dumps(audit))
    service=ProductionWorkbookImporter();a=service.assess()
    assert a['skipped_count']==1 and a['eligible_count']==1
    service.run_import(assessment_token=a['assessment_token'])
    assert counts(settings)==(1,1,1,1,1)
    lock=settings.production_import_workbook.with_name('~$production-import.xlsx');lock.touch()
    with pytest.raises(ValueError,match='locked'): service.assess()
    lock.unlink()


def test_imported_baseline_can_generate_patch_and_market_info(production):
    settings,_,_=production
    importer=ProductionWorkbookImporter();a=importer.assess();importer.run_import(assessment_token=a['assessment_token'])
    service=XmlGenerationService(require_import=True)
    preview=service.preview_generated_patch_scenario(product_family='Synthetic model',product_variant='Synthetic model',catalogue_number='SYN-00000',scenario_id='trade_name_edit',patch_version='3',scenario_inputs={'new_trade_name':'Updated synthetic name'})
    assert preview.derived_patch_validation.valid, preview.derived_patch_validation.errors
    assert preview.context.base_state_source=='production_export'
    assert preview.context.base_version=='2'
    assert 'Imported accepted' in preview.context.base_state_label
    assert '99000000000000' in preview.derived_patch_xml
    market=service.preview_market_info_put(product_family='Synthetic model',product_variant='Synthetic model',catalogue_number='SYN-00000',market_info_version='3',market_countries=[('DE',True),('FR',False)])
    assert market.validation.valid, market.validation.errors


def acknowledgement(xml: str, version: str, *, code='99000000000000') -> bytes:
    root=etree.fromstring(xml.encode())
    root.tag='{'+NS['message']+'}PushResponse'
    payload=root.find('message:payload',NS)
    root.remove(payload)
    response=child(root,'message','responseEntity')
    child(response,'message','responseCode','SUCCESS')
    child(response,'message','entityCode',code)
    child(response,'message','entityVersion',version)
    return etree.tostring(root,xml_declaration=True,encoding='utf-8')


def test_import_zip_success_xml_duplicate_delayed_and_next_patch(production):
    from io import BytesIO
    from zipfile import ZipFile
    from app.services.testing_success_xml import TestingSuccessXmlService as SuccessService
    settings,_,_=production
    importer=ProductionWorkbookImporter();a=importer.assess();importer.run_import(assessment_token=a['assessment_token'])
    service=XmlGenerationService(require_import=True)
    args=dict(product_family='Synthetic model',product_variant='Synthetic model',catalogue_number='SYN-00000',scenario_id='trade_name_edit',patch_version='3',scenario_inputs={'new_trade_name':'Accepted new name'})
    _,archive=service.download_generated_patch_scenario(**args)
    with ZipFile(BytesIO(archive)) as package:
        xmls=[package.read(n).decode() for n in package.namelist() if n.endswith('.xml')]
    derived=next(x for x in xmls if 'Accepted new name' in x)
    ack=acknowledgement(derived,'3')
    success=SuccessService();result=success.record_success_xml(xml_bytes=ack,source_file_name='synthetic-success.xml')
    assert result.recorded_event_count==1
    with sqlite3.connect(settings.testing_state_db_path) as c:
        version,state=c.execute('SELECT latest_successful_version,latest_successful_state_json FROM testing_subjects').fetchone()
        assert version=='3' and json.loads(state)['trade_name']=='Accepted new name'
    before=counts(settings)
    success.record_success_xml(xml_bytes=ack,source_file_name='duplicate.xml')
    assert counts(settings)==before
    next_preview=service.preview_generated_patch_scenario(**{**args,'patch_version':'4','scenario_inputs':{'new_trade_name':'Next name'}})
    assert next_preview.context.base_version=='3' and 'Accepted new name' in next_preview.base_xml
    assert 'DE' in next_preview.derived_patch_xml
    # An older version response remains history without rolling back accepted state.
    success.record_success_xml(xml_bytes=acknowledgement(derived,'2'),source_file_name='older.xml')
    assert StateStore().latest_successful_patch_state(product_family='Synthetic model',product_variant='Synthetic model',catalogue_number='SYN-00000').state.version=='3'
    summary=RegistrationSummaryService().summary()
    assert summary['counts']['registered_devices']==1
    assert summary['event_counts']['UDI_DI.PATCH']==1


def test_imported_devices_are_in_bulk_patch_market_info_and_read_models(production):
    from app.services.testing_read_model import TestingReadModelService
    settings,_,_=production
    importer=ProductionWorkbookImporter();a=importer.assess();importer.run_import(assessment_token=a['assessment_token'])
    read=TestingReadModelService()
    summaries=read.list_subject_summaries()
    assert len(summaries)==1 and summaries[0].imported_baseline_present
    assert not summaries[0].post_success and summaries[0].latest_success_message_type is None
    assert read.subject_history(summaries[0].id).subject.imported_baseline_present
    entries=read.bulk_patch_posted_entries(product_family='Synthetic model',product_variant='Synthetic model',basic_udi_di='BASICSHARED')
    assert len(entries)==1 and entries[0]['current_state']['version']=='2'
    service=XmlGenerationService(require_import=True)
    common=dict(product_family='Synthetic model',product_variant='Synthetic model',basic_udi_di='BASICSHARED',record_count=1,selected_catalogue_numbers=['SYN-00000'])
    preview=service.preview_bulk_patch(**common,scenario_id='trade_name_edit',scenario_inputs={'new_trade_name':'Bulk name'})
    assert preview.selected_chunk_validation.valid and preview.included_record_count==1
    market=service.preview_bulk_market_info(**common,market_countries=[('DE',True),('FR',False)])
    assert market.selected_chunk_validation.valid


@pytest.mark.parametrize('version', ['1', '3'])
def test_imported_versions_generate_the_next_patch(production, version):
    _,paths,prepare=production
    path=paths['xml_dir']/'export.xml'
    root=etree.fromstring(path.read_bytes())
    for node in root.findall('.//entity:version',NS):
        node.text=version
    path.write_bytes(etree.tostring(root))
    prepare()
    importer=ProductionWorkbookImporter(); a=importer.assess(); importer.run_import(assessment_token=a['assessment_token'])
    preview=XmlGenerationService(require_import=True).preview_generated_patch_scenario(product_family='Synthetic model',product_variant='Synthetic model',catalogue_number='SYN-00000',scenario_id='trade_name_edit',patch_version=str(int(version)+1),scenario_inputs={'new_trade_name':'Next version'})
    assert preview.context.base_version==version and preview.derived_patch_validation.valid


def test_expired_configuration_and_database_changes_require_reassessment(production):
    settings,_,_=production
    importer=ProductionWorkbookImporter()
    a=importer.assess(); ProductionWorkbookImporter._pending[a['assessment_token']].expires_at=0
    with pytest.raises(ValueError,match='confirmation'):
        importer.run_import(assessment_token=a['assessment_token'])
    a=importer.assess()
    importer.settings=importer.settings.model_copy(update={'eudamed_manufacturer_srn_override':'GB-MF-000000002'})
    with pytest.raises(ValueError,match='configuration changed'):
        importer.run_import(assessment_token=a['assessment_token'])
    importer=ProductionWorkbookImporter(); a=importer.assess()
    StateStore()
    with pytest.raises(ValueError,match='Database state changed'):
        importer.run_import(assessment_token=a['assessment_token'])
    with sqlite3.connect(settings.testing_state_db_path) as connection:
        assert connection.execute('SELECT COUNT(*) FROM testing_events').fetchone()[0]==0


def test_pair_publication_rollback_and_hash_validation(production,monkeypatch):
    import app.services.production_import_files as files
    settings,_,prepare=production
    result=prepare()
    old=(settings.production_import_workbook.read_bytes(),settings.production_import_audit.read_bytes())
    staged_workbook, staged_audit=Path(result['workbook']),Path(result['audit'])
    staged_workbook.write_bytes(staged_workbook.read_bytes()+b'\n')
    import hashlib
    updated=json.loads(staged_audit.read_bytes())
    updated['workbook_sha256']=hashlib.sha256(staged_workbook.read_bytes()).hexdigest()
    staged_audit.write_text(json.dumps(updated))
    original=files.os.replace
    calls=0
    def fail_second(source,target):
        nonlocal calls
        calls+=1
        if calls==2:
            raise OSError('synthetic publication failure')
        return original(source,target)
    monkeypatch.setattr(files.os,'replace',fail_second)
    with pytest.raises(OSError,match='publication failure'):
        publish_current_pair(Path(result['workbook']),Path(result['audit']))
    # The published pair before the attempted replacement is restored coherently.
    assert old==(settings.production_import_workbook.read_bytes(),settings.production_import_audit.read_bytes())
    manifest=json.loads(settings.production_import_audit.read_bytes())
    import hashlib
    assert manifest['workbook_sha256']==hashlib.sha256(settings.production_import_workbook.read_bytes()).hexdigest()
    assert not (settings.production_import_workbook.parent/'.production-import.lock').exists()
    Path(result['audit']).write_text('{}')
    with pytest.raises(ValueError,match='companion audit'):
        publish_current_pair(Path(result['workbook']),Path(result['audit']))


def test_market_success_from_export_preserves_countries_in_next_patch(production):
    from io import BytesIO
    from zipfile import ZipFile
    from app.services.testing_success_xml import TestingSuccessXmlService as SuccessService
    _,_,_=production
    importer=ProductionWorkbookImporter();a=importer.assess();importer.run_import(assessment_token=a['assessment_token'])
    service=XmlGenerationService(require_import=True)
    identity=dict(product_family='Synthetic model',product_variant='Synthetic model',catalogue_number='SYN-00000')
    _,archive=service.download_market_info_put(**identity,market_info_version='3',market_countries=[('DE',True),('FR',False)])
    with ZipFile(BytesIO(archive)) as package:
        xml=next(package.read(n).decode() for n in package.namelist() if n.endswith('.xml'))
    result=SuccessService().record_success_xml(xml_bytes=acknowledgement(xml,'3'),source_file_name='market-success.xml')
    assert result.recorded_event_count==1
    preview=service.preview_generated_patch_scenario(**identity,scenario_id='trade_name_edit',patch_version='3',scenario_inputs={'new_trade_name':'After market acceptance'})
    assert preview.derived_patch_validation.valid and '>FR<' in preview.derived_patch_xml
    assert StateStore().latest_successful_patch_state(**identity).state.version=='2'


def test_imported_baselines_are_available_through_ui_assessments(production):
    from app.services.operation_assessment import OperationAssessmentService
    _,_,_=production
    importer=ProductionWorkbookImporter();a=importer.assess();importer.run_import(assessment_token=a['assessment_token'])
    identity=dict(product_family='Synthetic model',product_variant='Synthetic model',catalogue_number='SYN-00000')
    service=OperationAssessmentService()
    assert service.assess_single_patch(**identity).status=='available'
    assert service.assess_single_market_info(**identity).status=='available'
    assert service.assess_single_post(**identity).status=='blocked'
    readiness=service.record_readiness()
    accepted=next(r for r in readiness if r['catalogue_number']=='SYN-00000')
    assert accepted['patch_ready'] and accepted['market_info_ready'] and not accepted['post_ready']
    preview=service.xml_service.preview_post_registration(**identity,accepted_baseline=True)
    assert preview.post_validation.valid


def test_server_profile_selects_the_importer_and_has_no_fallback(production,monkeypatch):
    from app.services import workbook_import_selector as selector
    settings,_,_=production
    monkeypatch.setattr(selector,'WorkbookImportService',lambda:'dev-importer')
    monkeypatch.setattr(selector,'ProductionWorkbookImporter',lambda:'prod-importer')
    assert selector.workbook_importer()=='prod-importer'
    monkeypatch.setattr(selector,'get_settings',lambda:settings.model_copy(update={'environment':'dev'}))
    assert selector.workbook_importer()=='dev-importer'


def test_populated_unsupported_export_fields_block_xml_without_skipping_import(production):
    _,paths,prepare=production
    path=paths['xml_dir']/'export.xml'
    root=etree.fromstring(path.read_bytes())
    device=root.find('.//device:MDRUDIDIData',NS)
    node=etree.SubElement(device,'{'+NS['udi']+'}deviceMarking');node.text='Additional accepted marking'
    path.write_bytes(etree.tostring(root));prepare()
    importer=ProductionWorkbookImporter();assessment=importer.assess()
    assert assessment['eligible_count']==2 and assessment['xml_blocked_count']==2
    assert any('deviceMarking' in reason for r in assessment['xml_blocked_rows'] for reason in r['reasons'])


def test_production_api_requires_assessment_before_commit(production,monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app
    settings,_,_=production
    monkeypatch.setattr('app.routers.profiling.get_settings',lambda:settings)
    client=TestClient(app)
    response=client.post('/api/workbook-imports/run',json={})
    assert response.status_code==400 and not settings.testing_state_db_path.exists()
    assessment=client.post('/api/workbook-imports/assess')
    assert assessment.status_code==200 and not settings.testing_state_db_path.exists()
    result=client.post('/api/workbook-imports/run',json={'assessment_token':assessment.json()['assessment_token']})
    assert result.status_code==200 and result.json()['created_count']==2


def test_production_canonical_review_uses_imported_data_without_legacy_files(production,monkeypatch):
    from app.services.canonical_review import CanonicalReviewService
    from app.routers.canonical import canonical_review
    settings,_,_=production
    def forbidden(*args,**kwargs):
        raise AssertionError('Production must not load legacy reference workbooks')
    monkeypatch.setattr('app.services.basic_udi_reference.BasicUdiReferenceService.list_variant_mappings',forbidden)
    assert CanonicalReviewService().load_review_bundle().variant_mappings==[]
    assert not settings.testing_state_db_path.exists()
    from app.services.workbook_import import WorkbookImportService
    WorkbookImportService()
    StateStore()
    assert canonical_review()['variant_mappings']==[]
    importer=ProductionWorkbookImporter();assessment=importer.assess();importer.run_import(assessment_token=assessment['assessment_token'])
    before=counts(settings)
    mappings=CanonicalReviewService().load_review_bundle().variant_mappings
    assert len(mappings)==2
    accepted=next(m for m in mappings if m.submission_operation=='PATCH')
    assert accepted.basic_udi_di=='BASICSHARED' and accepted.first_eu_market_country=='DE'
    assert counts(settings)==before


def test_production_completeness_recovers_export_status_without_reimport_or_state_changes(production):
    from app.services.workbook_import import WorkbookImportService
    settings, _, _ = production
    importer = ProductionWorkbookImporter()
    importer.run_import(assessment_token=importer.assess()['assessment_token'])
    service = WorkbookImportService()
    with sqlite3.connect(settings.testing_state_db_path) as connection:
        row = connection.execute("SELECT id, record_json FROM canonical_device_record WHERE record_json LIKE '%Supplied EUDAMED export snapshot%' LIMIT 1").fetchone()
        record = json.loads(row[1])
        fields = {field['canonical_path']: field for field in record['fields']}
        assert fields['device_record.market_availability.market_status']['value'] == fields['device_record.status']['value']
        fields['device_record.market_availability.market_status']['value'] = None
        record['blockers'].append('UDI-DI Market Status is not populated.')
        record['completeness']['missing_required_fields'] += 1
        record['completeness']['mapped_required_fields'] -= 1
        legacy_json = json.dumps(record)
        connection.execute('UPDATE canonical_device_record SET record_json=? WHERE id=?', (legacy_json, row[0]))
        evidence_before = connection.execute('SELECT state_after_json FROM testing_events').fetchall()
    before = counts(settings)
    bundle = service.canonical_validation_bundle_from_sqlite()
    recovered = next(r for r in bundle.records if r.primary_udi_di == record['primary_udi_di'])
    current = {f.canonical_path: f for f in recovered.fields}
    assert current['device_record.market_availability.market_status'].value == current['device_record.status'].value
    assert 'UDI-DI Market Status is not populated.' not in recovered.blockers
    assert len(recovered.blockers) == 9
    assert recovered.completeness.missing_required_fields == 9
    assert current['basic_device.clinical_investigation'].value is None
    assert current['device_record.cmr_present'].value is None
    assert recovered.xml_blockers == record['xml_blockers']
    assert counts(settings) == before
    with sqlite3.connect(settings.testing_state_db_path) as connection:
        assert connection.execute('SELECT record_json FROM canonical_device_record WHERE id=?', (row[0],)).fetchone()[0] == legacy_json
        assert connection.execute('SELECT state_after_json FROM testing_events').fetchall() == evidence_before
    # The existing Dev projection must retain its original completeness values.
    service.settings = settings.model_copy(update={'environment': 'dev'})
    dev_record = next(r for r in service.canonical_validation_bundle_from_sqlite().records if r.primary_udi_di == record['primary_udi_di'])
    assert next(f for f in dev_record.fields if f.canonical_path == 'device_record.market_availability.market_status').value is None
    assert len(dev_record.blockers) == 10


def test_export_completeness_does_not_guess_or_replace_values(production):
    from app.services.production_completeness import populate_export_market_status
    from app.services.workbook_import import WorkbookImportService
    importer = ProductionWorkbookImporter()
    importer.run_import(assessment_token=importer.assess()['assessment_token'])
    record = WorkbookImportService().canonical_validation_bundle_from_sqlite().records[0]
    fields = {field.canonical_path: field for field in record.fields}
    source = fields['device_record.status']
    target = fields['device_record.market_availability.market_status']
    target.value = 'EXISTING'
    assert not populate_export_market_status(record)
    assert target.value == 'EXISTING'
    target.value = None
    source.source_detail = 'Template field'
    assert not populate_export_market_status(record)
    source.source_detail = 'Supplied EUDAMED export snapshot'
    source.value = None
    assert not populate_export_market_status(record)
    assert target.value is None


def test_startup_completes_partial_production_database_before_health_checks(production,monkeypatch):
    from app import main
    from app.services.workbook_import import WorkbookImportService
    settings,_,_=production
    monkeypatch.setattr(main,'get_settings',lambda:settings)
    service=WorkbookImportService()
    missing={issue.table_name for issue in service.database_health_summary().issues if issue.code=='missing_table'}
    assert missing=={'testing_subjects','testing_events','generated_packages','reviewed_post_baselines'}
    main.log_testing_state_context()
    assert service.database_health_summary().issues==[]
    assert counts(settings)==(0,0,0,0,0)
    main.log_testing_state_context()
    assert service.database_health_summary().issues==[]
    assert counts(settings)==(0,0,0,0,0)
    assessment=ProductionWorkbookImporter().assess()
    assert assessment['new_count']==2 and counts(settings)==(0,0,0,0,0)
