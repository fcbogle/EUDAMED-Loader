"""Real batch/preparation HTTP flow against synthetic files and no database services."""
import io
import json
from pathlib import Path
import zipfile

from fastapi import FastAPI
from fastapi.testclient import TestClient
import pytest

from app.config import get_settings
from app.routers import dev_input_batches as routes
from app.services import dev_batch_preparation as module
from app.services.dev_input_batches import fingerprint
from test_dev_input_batches import project, create


@pytest.fixture
def setup(project, monkeypatch):
    batch = create(project)
    sentinel = project / "operational.sqlite3"
    sentinel.write_bytes(b"operational database must not be opened")
    settings = get_settings().model_copy(update={"environment": "dev", "dev_input_batch_root": batch.parent,
        "eudamed_manufacturer_srn_override": "GB-MF-000000001", "testing_state_db_path": sentinel})
    monkeypatch.setattr(routes, "get_settings", lambda: settings)
    app = FastAPI(); app.include_router(routes.router, prefix="/api")
    return TestClient(app), module.DevBatchPreparation(settings), batch, sentinel, settings


def test_real_check_prepare_and_paired_download_without_database_access(setup):
    client, service, batch, sentinel, _ = setup
    before = sentinel.read_bytes()
    prefix = f"/api/dev-input-batches/{batch.name}"
    inventory = client.get("/api/dev-input-batches").json()["batches"]
    assert inventory[0]["file_integrity"] == "verified"
    response = client.post(prefix + "/inspect")
    assert response.status_code == 200
    checked = response.json(); assert checked["inspection_complete"]
    response = client.post(prefix + "/prepare", json={"inspection_token": checked["inspection_token"]})
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["summary"]["registered_tab_rows"] == 1
    assert result["workbook_name"].startswith("dev-import-")
    assert result["database_changed"] is False
    artifact = prefix + "/prepared/" + result["artifact_id"]
    assert client.get(artifact + "/workbook").content.startswith(b"PK")
    audit = client.get(artifact + "/audit").json()
    assert audit["workbook_sha256"] == result["workbook_sha256"]
    assert audit["dev_input_batch"]["batch_id"] == batch.name
    package = client.get(artifact + "/package")
    assert package.status_code == 200
    with zipfile.ZipFile(io.BytesIO(package.content)) as archive:
        assert set(archive.namelist()) == {result["workbook_name"], result["audit_name"]}
    assert sentinel.read_bytes() == before
    assert not (batch / "prepared/.preparing.lock").exists()


def test_stale_check_token_and_changed_sources_cannot_generate_outputs(setup):
    client, service, batch, _, _ = setup
    token = service.inspect(batch.name)["inspection_token"]
    with pytest.raises(ValueError, match="stale"):
        service.prepare(batch.name, "wrong")
    (batch / "templates/Template.xlsx").write_bytes(b"changed")
    with pytest.raises(ValueError, match="manifest"):
        service.prepare(batch.name, token)
    assert not list((batch / "prepared").iterdir())


def test_production_endpoints_are_forbidden(setup, monkeypatch):
    client, _, batch, sentinel, settings = setup
    monkeypatch.setattr(routes, "get_settings", lambda: settings.model_copy(update={"environment": "prod"}))
    before = sentinel.read_bytes()
    assert client.get("/api/dev-input-batches").status_code == 403
    assert client.post(f"/api/dev-input-batches/{batch.name}/inspect").status_code == 403
    assert client.post(f"/api/dev-input-batches/{batch.name}/prepare", json={"inspection_token": "x"}).status_code == 403
    assert sentinel.read_bytes() == before
    assert not list((batch / "prepared").iterdir())


def test_excluded_workbook_retained_but_not_added_to_preparation(setup):
    _, service, batch, _, settings = setup
    import shutil
    name = settings.excluded_excel_workbook_names[0]
    target = batch / "templates" / name
    shutil.copy2(batch / "templates/Template.xlsx", target)
    path = batch / "manifest.json"; manifest = json.loads(path.read_text())
    entry = dict(next(row for row in manifest["files"] if row["role"] == "templates"))
    entry["path"] = "templates/" + name
    manifest["files"].append(entry); path.write_text(json.dumps(manifest))
    check = service.inspect(batch.name)
    assert check["templates"]["duplicate_identity_occurrences"] == 1
    assert check["requested_inspection_scope"]["duplicate_identity_occurrences"] == 0
    result = service.prepare(batch.name, check["inspection_token"])
    record, _, audit = service.artifact(batch.name, result["artifact_id"])
    evidence = json.loads(audit.read_text())
    assert result["summary"]["template_rows"] == 1
    assert evidence["dev_input_batch"]["excluded_template_workbooks"] == [name]
    assert target.is_file()


def test_failed_preparation_leaves_no_partial_pair(setup, monkeypatch):
    _, service, batch, _, _ = setup
    checked = service.inspect(batch.name)
    original = module.prepare_workbook
    def interrupted(**kwargs):
        result = original(**kwargs)
        (batch / "templates/Template.xlsx").write_bytes(b"concurrent change")
        return result
    monkeypatch.setattr(module, "prepare_workbook", interrupted)
    with pytest.raises(ValueError, match="manifest"):
        service.prepare(batch.name, checked["inspection_token"])
    assert not list((batch / "prepared").iterdir())


def test_output_and_batch_path_guards_and_download_integrity(setup):
    client, service, batch, _, _ = setup
    with pytest.raises(ValueError, match="batch ID"):
        service.describe("../../prod")
    checked = service.inspect(batch.name)
    result = service.prepare(batch.name, checked["inspection_token"])
    _, workbook, _ = service.artifact(batch.name, result["artifact_id"])
    workbook.write_bytes(b"tampered workbook")
    assert client.get(f"/api/dev-input-batches/{batch.name}/prepared/{result['artifact_id']}/package").status_code == 409
    assert not list(batch.parent.glob(".preparing-*"))


def test_prepared_symlink_cannot_redirect_output(setup):
    _, service, batch, _, _ = setup
    checked = service.inspect(batch.name)
    (batch / "prepared").rmdir()
    other = batch.parent / "other-output"; other.mkdir()
    (batch / "prepared").symlink_to(other, target_is_directory=True)
    with pytest.raises(ValueError, match="symlink"):
        service.prepare(batch.name, checked["inspection_token"])
    assert not list(other.iterdir())


def test_export_absence_is_not_labelled_unregistered_in_dev(setup):
    _, service, batch, _, _ = setup
    from test_production_workbook_preparation import workbook as make_workbook
    from openpyxl import load_workbook
    path = batch / "templates/Template.xlsx"
    make_workbook(path, ["Issuing Entity e.g. GS1", "UDI-DI code", "Basic UDI-DI"],
                  [["GS1", "00000000000001", "PARENT"], ["GS1", "00000000000002", "PARENT"]])
    manifest_path = batch / "manifest.json"; manifest = json.loads(manifest_path.read_text())
    for entry in manifest["files"]:
        if entry["role"] == "templates":
            entry.update(sha256=fingerprint(path), size_bytes=path.stat().st_size)
    manifest_path.write_text(json.dumps(manifest))
    checked = service.inspect(batch.name)
    result = service.prepare(batch.name, checked["inspection_token"])
    _, output, audit = service.artifact(batch.name, result["artifact_id"])
    book = load_workbook(output, read_only=True)
    headers, row = list(book['To Register'].iter_rows(values_only=True))
    assert row[headers.index('Registration Status')] == 'Not evidenced by supplied exports'
    book.close()
    assert 'not proof of non-registration' in json.loads(audit.read_text())['dev_input_batch']['scope_limitation']
