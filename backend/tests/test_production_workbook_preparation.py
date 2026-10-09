"""Synthetic workbook/export tests; no production files or databases are changed."""
from __future__ import annotations

import json
from pathlib import Path

from lxml import etree
from openpyxl import Workbook, load_workbook
import pytest

from app.services.production_workbook_preparation import NS, prepare_workbook


def workbook(path: Path, headers: list, rows: list, sheet: str = "Model", preamble: int = 0) -> None:
    book = Workbook()
    page = book.active
    page.title = sheet
    for _ in range(preamble):
        page.append(["Template instructions"])
    page.append(headers)
    for row in rows:
        page.append(row)
    book.save(path)
    book.close()


def child(parent, prefix, name, text=None):
    node = etree.SubElement(parent, f"{{{NS[prefix]}}}{name}")
    node.text = text
    return node


def export(path: Path, code="00000000000001", *, missing_market=False, page=0, pages=1,
           parent_link="PARENT", trade_name="Accepted name", encoding="utf-8") -> None:
    root = etree.Element(f"{{{NS['message']}}}PullResponse", nsmap=NS)
    child(root, "message", "creationDateTime", "2026-09-11T10:00:00Z")
    payload = child(root, "message", "payload")
    device = child(payload, "device", "Device")
    basic = child(device, "device", "MDRBasicUDI")
    udi = child(device, "device", "MDRUDIDIData")
    for node in (basic, udi):
        child(node, "entity", "state", "REGISTERED")
        child(node, "entity", "version", "2")
    model = child(basic, "basic", "modelName")
    child(model, "common", "name", "Model")
    for node, prefix, identifier, value in ((basic, "basic", "identifier", "PARENT"),
                                          (udi, "udi", "identifier", code),
                                          (udi, "udi", "basicUDIIdentifier", parent_link)):
        ident = child(node, prefix, identifier)
        child(ident, "common", "DICode", value)
        child(ident, "common", "issuingEntityCode", "GS1")
    child(udi, "udi", "referenceNumber", trade_name)
    market = child(udi, "udi", "marketInfos")
    if not missing_market:
        child(market, "entity", "state", "REGISTERED")
        child(market, "entity", "version", "3")
    for country in ("DE", "FR"):
        item = child(market, "udi", "marketInfo")
        child(item, "udi", "country", country)
        child(item, "udi", "originalPlacedOnTheMarket", "true" if country == "DE" else "false")
    child(root, "message", "pageNumber", str(page))
    child(root, "message", "numberOfPages", str(pages))
    child(root, "message", "pageSize", "50")
    path.write_bytes(etree.tostring(root, encoding=encoding, xml_declaration=True))


@pytest.fixture
def inputs(tmp_path):
    template, xml, output = [tmp_path / name for name in ("templates", "xml", "out")]
    template.mkdir()
    xml.mkdir()
    parent = tmp_path / "BasicUDIs.xlsx"
    workbook(parent, ["Issuing Entity", "Basic UDI-DI code", "Device Model", "Risk class"],
             [["GS1", "PARENT", "Model", "Class I"]], sheet="Upload")
    workbook(template / "Template.xlsx", ["Issuing Entity e.g. GS1", "UDI-DI code", "Trade Name"],
             [["GS1", "00000000000001", "Proposed name"], ["GS1", "00000000000002", "Other name"]],
             preamble=4)
    export(xml / "export.xml")
    return dict(template_dir=template, xml_dir=xml, parent_reference=parent, output_dir=output)


def read_rows(path, sheet):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        rows = iter(book[sheet].iter_rows(values_only=True))
        headers = next(rows)
        return [dict(zip(headers, r)) for r in rows]
    finally:
        book.close()


def test_overlap_separates_proposed_and_accepted_and_preserves_all_sources(inputs):
    result = prepare_workbook(**inputs)
    assert result["matched_identities"] == 1
    registered = read_rows(result["workbook"], "Registered")
    pending = read_rows(result["workbook"], "To Register")
    assert len(registered) == len(pending) == 1
    assert registered[0]["UDI-DI"] == "00000000000001"
    assert registered[0]["Proposed: Trade Name"] == "Proposed name"
    assert registered[0]["Accepted Device: referenceNumber"] == "Accepted name"
    assert registered[0]["Device Version"] == "2"
    assert registered[0]["Market Info Version"] == "3"
    assert pending[0]["Registration Status"] == "Not registered"
    assert pending[0]["Proposed Basic UDI-DI"] == "PARENT"
    audit = json.loads(Path(result["audit"]).read_text())
    assert len(audit["template_occurrences"]) == 2
    assert audit["template_occurrences"][0]["source"]["header_row"] == 5
    assert "DE" in registered[0]["Accepted Device: marketInfos"]
    assert "FR" in registered[0]["Accepted Device: marketInfos"]
    assert not list(inputs["output_dir"].rglob("*.sqlite3"))


def test_conflicting_duplicate_rows_are_preserved_without_choosing_a_value(inputs):
    workbook(inputs["template_dir"] / "Duplicate.xlsx", ["Issuing Entity", "UDI-DI code", "Trade Name"],
             [["GS1", "00000000000002", "Conflicting name"]])
    result = prepare_workbook(**inputs)
    row = read_rows(result["workbook"], "To Register")[0]
    assert row["Registration Status"] == "Needs review"
    assert row["Source Occurrence Count"] == 2
    assert row["Proposed: Trade Name"] is None
    assert "Conflicting template values" in row["Review Reasons"]
    assert len(json.loads(Path(result["audit"]).read_text())["template_occurrences"]) == 3


def test_missing_market_version_remains_unknown_on_registered_tab(inputs):
    export(inputs["xml_dir"] / "export.xml", missing_market=True)
    result = prepare_workbook(**inputs)
    row = read_rows(result["workbook"], "Registered")[0]
    assert row["Registration Status"] == "Registered"
    assert row["Market Info Version"] is None
    assert row["Review Required"] == "Yes"
    assert "Missing exported Market Info" in row["Review Reasons"]


def test_unknown_parent_and_unnamed_column_remain_visible(inputs):
    workbook(inputs["template_dir"] / "New.xlsx", ["Issuing Entity", "UDI-DI code", None],
             [["GS1", "00000000000003", "Extra business value"]], sheet="Unknown")
    result = prepare_workbook(**inputs)
    row = next(r for r in read_rows(result["workbook"], "To Register") if r["UDI-DI"] == "00000000000003")
    assert row["Proposed: Unnamed column 3"] == "Extra business value"
    assert row["Proposed Basic UDI-DI"] is None
    assert "Missing Basic UDI-DI" in row["Review Reasons"]


def test_xml_parent_link_mismatch_stops_without_outputs(inputs):
    export(inputs["xml_dir"] / "export.xml", parent_link="WRONG")
    with pytest.raises(ValueError, match="conflicting exported identity"):
        prepare_workbook(**inputs)
    assert not inputs["output_dir"].exists()


def test_missing_export_page_stops_without_outputs(inputs):
    export(inputs["xml_dir"] / "export.xml", pages=2)
    with pytest.raises(ValueError, match="populated export pages"):
        prepare_workbook(**inputs)
    assert not inputs["output_dir"].exists()


def test_encoding_override_is_explicit_and_recorded(inputs):
    path = inputs["xml_dir"] / "export.xml"
    export(path, trade_name="Child’s device")
    path.write_bytes(path.read_bytes().decode("utf-8").encode("windows-1252"))
    with pytest.raises(etree.XMLSyntaxError):
        prepare_workbook(**inputs)
    result = prepare_workbook(**inputs, encoding_overrides={"export.xml": "windows-1252"})
    row = read_rows(result["workbook"], "Registered")[0]
    assert row["Accepted Device: referenceNumber"] == "Child’s device"
    assert row["Encoding Overrides"] == "windows-1252"


def test_rerun_retains_prior_workbook_and_does_not_modify_inputs(inputs):
    before = {p: p.read_bytes() for p in inputs["template_dir"].glob("*.xlsx")}
    first = prepare_workbook(**inputs)
    previous = Path(first["workbook"]).read_bytes()
    second = prepare_workbook(**inputs)
    assert first["workbook"] != second["workbook"]
    assert Path(first["workbook"]).read_bytes() == previous
    assert all(p.read_bytes() == contents for p, contents in before.items())


def test_formula_cache_absence_is_reported_and_raw_formula_is_preserved(inputs):
    workbook(inputs["template_dir"] / "Formula.xlsx", ["Issuing Entity", "UDI-DI code", "Trade Name"],
             [["GS1", "00000000000003", '=CONCAT("A","B")']])
    result = prepare_workbook(**inputs)
    row = next(r for r in read_rows(result["workbook"], "To Register") if r["UDI-DI"] == "00000000000003")
    assert "Formula has no cached value" in row["Review Reasons"]
    audit = json.loads(Path(result["audit"]).read_text())
    source = next(r for r in audit["template_occurrences"] if r["key"][1] == "00000000000003")
    assert source["formulas"]["Trade Name"] == '=CONCAT("A","B")'


def test_empty_first_cell_and_blank_rows_preserve_excel_row_numbers(inputs):
    workbook(inputs["template_dir"] / "MissingIssuer.xlsx", ["Issuing Entity", "UDI-DI code", "Trade Name"],
             [[None, "00000000000003", "Missing issuer"], [None, None, None],
              ["GS1", "00000000000004", "After blank row"]], preamble=4)
    result = prepare_workbook(**inputs)
    audit = json.loads(Path(result["audit"]).read_text())
    sources = {r["key"][1]: r for r in audit["template_occurrences"]
               if r["source"]["file"] == "MissingIssuer.xlsx"}
    assert sources["00000000000003"]["source"]["row"] == 6
    assert "Missing device identity" in sources["00000000000003"]["reasons"]
    assert sources["00000000000004"]["source"]["row"] == 8
    assert len(sources) == 2


def test_duplicate_export_identity_requires_review_before_output(inputs):
    export(inputs["xml_dir"] / "export.xml", pages=2)
    export(inputs["xml_dir"] / "second.xml", page=1, pages=2)
    with pytest.raises(ValueError, match="Duplicate exported device identities"):
        prepare_workbook(**inputs)


def test_oversized_source_cell_is_not_silently_truncated(inputs):
    # Repeated XML structures can exceed Excel capacity even when each text node fits.
    path = inputs["xml_dir"] / "export.xml"
    tree = etree.parse(str(path))
    udi = tree.find("message:payload/device:Device/device:MDRUDIDIData", NS)
    repeated = child(udi, "udi", "largeStructure")
    for _ in range(5):
        child(repeated, "udi", "text", "x" * 8000)
    path.write_bytes(etree.tostring(tree))
    with pytest.raises(ValueError, match="Excel cell capacity"):
        prepare_workbook(**inputs)
    assert not inputs["output_dir"].exists()


def test_parent_reference_in_template_folder_is_excluded(inputs):
    path = inputs["template_dir"] / "BasicUDIs.xlsx"
    path.write_bytes(inputs["parent_reference"].read_bytes())
    inputs["parent_reference"] = path
    result = prepare_workbook(**inputs)
    assert result["template_workbooks"] == 1
    assert result["template_rows"] == 2


def test_approved_countries_replace_tentative_proposed_values_and_preserve_sources(inputs):
    workbook(inputs["parent_reference"],
             ["Issuing Entity", "Basic UDI-DI code", "Device Model",
              "Member State of the placing on the EU market of the Device:",
              "Member States where device is or is to be made available on the market:",
              "URL for additional information (as electronic instructions for use):"],
             [["GS1", "PARENT", "Model", "Germany", "France?", "Not available yet"]], sheet="Upload")
    before = inputs["parent_reference"].read_bytes()
    result = prepare_workbook(**inputs, parent_market_overrides={("GS1", "PARENT"): ["Germany", "France"]})
    row = read_rows(result["workbook"], "To Register")[0]
    assert row["Proposed Parent: Member States where device is or is to be made available on the market:"] == "Germany; France"
    assert "market countries missing" not in row["Review Reasons"]
    assert "URL unavailable" in row["Review Reasons"]
    assert inputs["parent_reference"].read_bytes() == before
    audit = json.loads(Path(result["audit"]).read_text())
    assert audit["parent_reference_entries"][0]["fields"]["Member States where device is or is to be made available on the market:"] == "France?"
    entries = read_rows(result["workbook"], "Summary")
    assert any(r["Category"] == "Approved setting" and "First placement: Germany" in r["Details"] for r in entries)
    assert next(r for r in entries if r["Category"] == "Optional-field exception")["Affected Rows"] == 2


def test_summary_retains_missing_market_exception_and_prior_removed_device(inputs):
    first = prepare_workbook(**inputs)
    workbook(inputs["template_dir"] / "Template.xlsx", ["Issuing Entity", "UDI-DI code", "Trade Name"],
             [["GS1", "00000000000001", "Proposed name"]])
    export(inputs["xml_dir"] / "export.xml", missing_market=True)
    result = prepare_workbook(**inputs, previous_audit=Path(first["audit"]))
    book = load_workbook(result["workbook"], read_only=True)
    assert book.sheetnames == ["To Register", "Registered", "Summary"]
    book.close()
    entries = read_rows(result["workbook"], "Summary")
    assert next(r for r in entries if r["Category"] == "Removed from current templates")["UDI-DI / Issuer"] == "GS1:00000000000002"
    exception = next(r for r in entries if r["Category"] == "Device exception")
    assert "Market Info" in exception["Details"]
    assert exception["Data Location"] == "Registered row 2"
    assert result["removed_template_identities"] == 1


def test_unknown_parent_override_fails_before_output(inputs):
    with pytest.raises(ValueError, match="Invalid parent market override"):
        prepare_workbook(**inputs, parent_market_overrides={("GS1", "MISSING"): ["Germany"]})
    assert not inputs["output_dir"].exists()
