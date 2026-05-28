from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from app.config import get_settings
from app.services.basic_udi_reference import BasicUdiReferenceRow, BasicUdiReferenceService
from app.services.canonical_review import CanonicalReviewService
from app.services.normalization import NormalizationRepository
from app.validation_models import (
    CanonicalValidationBundle,
    CanonicalValidationFieldValue,
    CanonicalValidationRecord,
    CompletenessSnapshot,
    DeferredValidationScopeSummary,
    FamilyValidationSummary,
    MarketAvailabilityItemPreview,
    MatchStatus,
    SourceFieldCoverageEntry,
    SourceFieldCoverageStatus,
    SourceFieldCoverageSummary,
    StructuredListItemPreview,
    ValidationBlockerSummary,
    VariantValidationSummary,
    ValueSourceType,
)

HEADER_SENTINEL = "UDI-DI code"

ISSUING_ENTITY_HEADER = "Issuing Entity e.g. GS1"
UDI_DI_HEADER = "UDI-DI code e.g. taken from 2nd page of DoC (Note: needs to be 14 digits long, add zero to front of code)"
UDI_DI_HEADERS = (
    UDI_DI_HEADER,
    "UDI-DI code e.g. taken from ist page of DoC (Note: needs to be 14 digits long, add zero to front of code)",
)
SECONDARY_UDI_HEADER = "UDI-DI from another entity (secondary) applicable e.g. No"
EMDN_HEADER = "Enter a nomenclature code (EMDN code) e.g starts with Y06xxxx then click Find"
TRADE_NAME_HEADER = "Trade Name applicable e.g. YES (use  Decription detail from second page of DoC"
LANGUAGE_HEADER = "Select the language e.g English"
CATALOGUE_HEADERS = (
    "Reference/ Catalogue number e.g . Taken from Product code in second page of DoC",
    "Reference/ Catalogue number e.g . Taken from Product code on second page of DoC",
)
QUANTITY_HEADER = "Quantity of device e.g. 1"
UDI_PI_HEADER = "Type of UDI-PI e.g. select Serial number and Manufacturing Date"
STATUS_HEADER = "UDI-DI status e.g. On the EU market"
STERILE_HEADER = "Device labelled as sterile e.g. NO"
LATEX_HEADER = "Containing latex e.g. NO"
DIRECT_MARKING_HEADER = "Is the device directly marked? E.g. NO"
SINGLE_USE_HEADER = "Labelled as single use e.g.NO"
CLINICAL_SIZE_APPLICABLE_HEADER = "Clinical size applicable e.g. NO"
MAX_REUSES_APPLICABLE_HEADER = "Maximum number of reuses applicable e.g. NO"
STERILIZATION_HEADER = "Need for sterilsation before use e.g. NO"
REPROCESSED_HEADER = "Reprocessed single use device e.g. NO"
ANNEX_XVI_HEADER = "Intended purpose other than medical (Annex XVI) e.g. NO"
DESIGNED_BY_ANOTHER_HEADER = "Is the device designed and manufactured by another legal or natural person? E.g. NO"
CLINICAL_INVESTIGATION_HEADER = "Clinical Investigation e.g. NO"
HUMAN_TISSUE_HEADER = "Presence of human tissues or cells, or their derivatives e.g. NO"
ANIMAL_TISSUE_HEADER = "Presence of animal tissues or cells, or their derivatives e.g. NO"
MEDICINAL_SUBSTANCE_HEADER = "Presence of a substance which, if used separately, may be considered to be a medicinal product e.g. NO"
HUMAN_BLOOD_SUBSTANCE_HEADER = "Presence of a substance which, if used separately, may be considered to be a medicinal product derived from human blood or human plasma e.g. NO"
CMR_HEADER = "Labelled for presence of Carcinogenic, Mutagenic and toxic to Reproduction (CMR) substances of category 1A or 1B e.g. NO"
ENDOCRINE_DISRUPTOR_HEADER = "Labelled for presence of substances with endocrine-disrupting properties e.g. NO"
FIRST_EU_MARKET_HEADER = "Member state where first placed on the EU market e.g. Germany"
STORAGE_APPLICABLE_HEADER = "Storage/handling conditions, if applicable e.g. YES"
STORAGE_TYPE_HEADER = "Storage /handling conditions type e.g. Lower limit of temp"
STORAGE_DESCRIPTION_ONE_HEADER = "Description e.g. taken from  IFU Technical Data page Storage Temp range e.g. -15C"
STORAGE_TYPE_TWO_HEADER = "Add another Storage/handling condition e.g. Upper limit of temp"
STORAGE_DESCRIPTION_TWO_HEADER = "Description e.g. taken from  IFU Technical Data page Storage Temp range e.g. +50C"
WARNING_APPLICABLE_HEADER = "Critical warnings or contra-indications, if applicable e.g. Yes"
WARNING_TYPE_HEADER = "Critical warning type e.g.  Consult instructions for use"
SECONDARY_IDENTIFIER_CODE_HEADER = "Secondary Identifier - Code"
SECONDARY_IDENTIFIER_ENTITY_HEADER = "Secondary Identifier - Issuing Entity"

FAMILY_LABELS: dict[str, str] = {
    "Template for Echelon family EUDAMED.xlsx": "Echelon",
    "Template for Elan products EUDAMED.xlsx": "Elan",
    "Template for Elite family EUDAMED.xlsx": "Elite",
    "Template for Epirus_Esprit EUDAMED.xlsx": "Epirus / Esprit",
    "Template for Navigator_Javelin_Linx EUDAMED.xlsx": "Navigator / Javelin / Linx",
}


@dataclass(frozen=True)
class SourceHeaderProfile:
    display_name: str
    normalized_name: str
    sheets: tuple[str, ...]


@dataclass(frozen=True)
class SourceRow:
    sheet_name: str
    row_index: int
    values: dict[str, object | None]


class CanonicalValidationService:
    def __init__(self) -> None:
        self.settings = get_settings()
        self.normalization_repository = NormalizationRepository()
        self.reference_service = BasicUdiReferenceService()

    def build_validation_bundle(self) -> CanonicalValidationBundle:
        variant_mappings = self.reference_service.list_variant_mappings()
        reference_rows = self.reference_service.rows_by_device_model()
        mapping_lookup = {(mapping.workbook, mapping.sheet): mapping for mapping in variant_mappings}

        records: list[CanonicalValidationRecord] = []
        deferred_scope_summaries: list[DeferredValidationScopeSummary] = []
        source_headers: list[SourceHeaderProfile] = []
        total_source_records = 0

        for workbook_path in sorted(self.settings.excel_dir.glob("*.xlsx")):
            workbook_rows = self._load_source_rows(workbook_path)
            workbook_headers = self._load_source_headers(workbook_path)
            total_source_records += len(workbook_rows)

            if workbook_path.name not in self.settings.excluded_excel_workbook_names:
                source_headers.extend(workbook_headers)

            rows_by_sheet: dict[str, list[SourceRow]] = {}
            for row in workbook_rows:
                rows_by_sheet.setdefault(row.sheet_name, []).append(row)

            for sheet_name, sheet_rows in sorted(rows_by_sheet.items()):
                mapping = mapping_lookup.get((workbook_path.name, sheet_name))
                if mapping is None or mapping.match_status != "matched" or mapping.device_model is None:
                    reason = (
                        mapping.notes[0]
                        if mapping and mapping.notes
                        else "No active variant mapping is available for this workbook sheet."
                    )
                    deferred_scope_summaries.append(
                        DeferredValidationScopeSummary(
                            workbook=workbook_path.name,
                            sheet_name=sheet_name,
                            record_count=len(sheet_rows),
                            reason=reason,
                        )
                    )
                    continue

                reference_row = reference_rows.get(mapping.device_model)
                if reference_row is None:
                    deferred_scope_summaries.append(
                        DeferredValidationScopeSummary(
                            workbook=workbook_path.name,
                            sheet_name=sheet_name,
                            record_count=len(sheet_rows),
                            reason="The matched Device Model did not resolve to a reference row in BasicUDIs.xlsx.",
                        )
                    )
                    continue

                family_label = FAMILY_LABELS.get(workbook_path.name, workbook_path.stem)
                for row in sheet_rows:
                    records.append(
                        self._build_record(
                            workbook_name=workbook_path.name,
                            product_family=family_label,
                            row=row,
                            reference_row=reference_row,
                        )
                    )

        source_field_coverage = self._build_source_field_coverage(self._merge_header_profiles(source_headers))
        tracked_required_fields = (
            sum(1 for field in records[0].fields if field.required)
            if records
            else 0
        )
        tracked_xml_required_fields = (
            sum(1 for field in records[0].fields if field.xml_required)
            if records
            else 0
        )

        return CanonicalValidationBundle(
            family_scope="In-scope non-accessories families",
            scope_note=(
                "Canonical validation now runs across the in-scope non-accessories product families using "
                "variant-level Basic UDI linkage from the authoritative BasicUDIs.xlsx workbook."
            ),
            validation_note=(
                "Canonical completeness and XML readiness are now shown separately. XML readiness applies "
                "the current XML-facing required field set across POST and PATCH variants, while "
                "operation-specific rules can still be tightened later once QMS confirms any divergent requirements."
            ),
            total_source_records=total_source_records,
            validation_subset_records=len(records),
            excluded_records=sum(summary.record_count for summary in deferred_scope_summaries),
            matched_reference_records=len(records),
            tracked_required_fields=tracked_required_fields,
            tracked_xml_required_fields=tracked_xml_required_fields,
            ready_records=sum(1 for record in records if record.completeness.status == "complete"),
            blocked_records=sum(1 for record in records if record.completeness.status == "incomplete"),
            xml_ready_records=sum(1 for record in records if record.xml_readiness.status == "complete"),
            xml_blocked_records=sum(1 for record in records if record.xml_readiness.status == "incomplete"),
            family_summaries=self._build_family_summaries(records),
            variant_summaries=self._build_variant_summaries(records),
            blocker_summaries=self._build_blocker_summaries(records),
            source_field_total=len(source_field_coverage),
            source_field_coverage_summaries=self._build_source_field_coverage_summaries(source_field_coverage),
            source_field_coverage=source_field_coverage,
            sample_records=self._build_sample_records(records),
            deferred_scope_summaries=deferred_scope_summaries,
            records=records,
        )

    def _build_record(
        self,
        *,
        workbook_name: str,
        product_family: str,
        row: SourceRow,
        reference_row: BasicUdiReferenceRow,
    ) -> CanonicalValidationRecord:
        source_values = row.values
        issuing_entity = self._issuing_entity_code(source_values.get(ISSUING_ENTITY_HEADER))
        primary_udi_di = self._first_string_value(source_values, UDI_DI_HEADERS)
        basic_udi_identifier = self._di_identifier(issuing_entity=issuing_entity, di_code=reference_row.basic_udi_di)
        secondary_identifier = self._di_identifier(
            issuing_entity=self._issuing_entity_code(source_values.get(SECONDARY_IDENTIFIER_ENTITY_HEADER)),
            di_code=self._string_value(source_values.get(SECONDARY_IDENTIFIER_CODE_HEADER)),
        )
        storage_condition_items = self._assemble_storage_condition_items(source_values)
        critical_warning_items = self._assemble_critical_warning_items(source_values)
        market_availability_items = self._assemble_market_availability_items(reference_row)
        production_identifier = self._production_identifier(source_values.get(UDI_PI_HEADER))

        fields = [
            self._field(
                canonical_path="manufacturer.issuing_entity",
                business_label="Issuing Entity",
                value=issuing_entity,
                source="workbook",
                source_detail=f"{row.sheet_name}!{ISSUING_ENTITY_HEADER}",
                xml_required=True,
            ),
            self._field(
                canonical_path="manufacturer.manufacturer_srn",
                business_label="Manufacturer SRN",
                value=reference_row.manufacturer_srn,
                source=reference_row.manufacturer_srn_source or "missing",
                source_detail=(
                    "uat-eudamed_mdr_products_tracekey_sample_data.xlsx!MDR-Basic Manufacturer Code"
                    if reference_row.manufacturer_srn_source == "legacy_basic_udi_reference"
                    else "No confirmed workbook or authoritative Basic UDI source column is currently available."
                ),
                review_note=(
                    "Supplemented from the legacy tracekey workbook because the current authoritative BasicUDIs.xlsx "
                    "workbook does not carry manufacturer SRN."
                    if reference_row.manufacturer_srn_source == "legacy_basic_udi_reference"
                    else "Still required for XML generation, but not currently present in BasicUDIs.xlsx or the in-scope device workbooks."
                ),
                required=False,
                xml_required=True,
            ),
            self._field(
                canonical_path="manufacturer.designed_by_another_legal_entity",
                business_label="Designed By Another Legal Entity",
                value=self._normalized_boolean(source_values.get(DESIGNED_BY_ANOTHER_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{DESIGNED_BY_ANOTHER_HEADER}",
                required=False,
            ),
            self._field(
                canonical_path="basic_device.regulation",
                business_label="Regulation",
                value=self._regulation_code(reference_row.applicable_regulation),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Applicable regulation",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.basic_udi_di",
                business_label="Basic UDI-DI",
                value=reference_row.basic_udi_di,
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Basic UDI-DI code",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.device_model",
                business_label="Device Model",
                value=reference_row.device_model,
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Device Model",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.device_type",
                business_label="Device Type",
                value=self._device_type(reference_row.device_type),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Is it a System or Procedure Pack which is a Device in itself?",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.special_device_type",
                business_label="Special Device Type",
                value=reference_row.special_device_type,
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Special device type",
                required=False,
            ),
            self._field(
                canonical_path="basic_device.risk_class",
                business_label="Risk Class",
                value=reference_row.risk_class,
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Risk class",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.authorised_representative_srn",
                business_label="Authorised Representative SRN",
                value=reference_row.authorised_representative_srn,
                source=reference_row.authorised_representative_srn_source or "missing",
                source_detail=(
                    "uat-eudamed_mdr_products_tracekey_sample_data.xlsx!MDR-Basic Authorised Representative Actor Code"
                    if reference_row.authorised_representative_srn_source == "legacy_basic_udi_reference"
                    else "No confirmed workbook or authoritative Basic UDI source column is currently available."
                ),
                review_note=(
                    "Supplemented from the legacy tracekey workbook because the current authoritative BasicUDIs.xlsx "
                    "workbook does not carry authorised representative SRN."
                    if reference_row.authorised_representative_srn_source == "legacy_basic_udi_reference"
                    else "Optional in XML unless the manufacturer is non-EU, but retained here as an explicit reviewable gap."
                ),
                required=False,
            ),
            self._field(
                canonical_path="basic_device.nomenclature_code",
                business_label="EMDN Code",
                value=self._string_value(source_values.get(EMDN_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{EMDN_HEADER}",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.intended_purpose_summary",
                business_label="Intended Purpose Summary",
                value=None,
                source="missing",
                source_detail="No dedicated intended-purpose summary field is currently confirmed in the in-scope workbook templates.",
                review_note="Retained as a visible canonical gap until QMS confirms a source field or an approved enrichment rule.",
                required=False,
            ),
            self._field(
                canonical_path="basic_device.annex_xvi_other_purpose",
                business_label="Annex XVI Other Purpose",
                value=self._normalized_boolean(source_values.get(ANNEX_XVI_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{ANNEX_XVI_HEADER}",
            ),
            self._field(
                canonical_path="basic_device.clinical_investigation",
                business_label="Clinical Investigation",
                value=self._normalized_boolean(source_values.get(CLINICAL_INVESTIGATION_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{CLINICAL_INVESTIGATION_HEADER}",
            ),
            self._field(
                canonical_path="basic_device.human_tissues_cells",
                business_label="Human Tissues Cells",
                value=self._normalized_boolean(source_values.get(HUMAN_TISSUE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{HUMAN_TISSUE_HEADER}",
                review_note="Used as the current XML-facing surrogate until an authoritative Basic UDI-level source is confirmed.",
                required=False,
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.animal_tissues_cells",
                business_label="Animal Tissues Cells",
                value=self._normalized_boolean(source_values.get(ANIMAL_TISSUE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{ANIMAL_TISSUE_HEADER}",
                review_note="Used as the current XML-facing surrogate until an authoritative Basic UDI-level source is confirmed.",
                required=False,
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.human_product_check",
                business_label="Human Product Check",
                value=self._normalized_boolean(source_values.get(HUMAN_BLOOD_SUBSTANCE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{HUMAN_BLOOD_SUBSTANCE_HEADER}",
                review_note="Used as the current XML-facing surrogate until an authoritative Basic UDI-level source is confirmed.",
                required=False,
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.medicinal_product_check",
                business_label="Medicinal Product Check",
                value=self._normalized_boolean(source_values.get(MEDICINAL_SUBSTANCE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{MEDICINAL_SUBSTANCE_HEADER}",
                review_note="Used as the current XML-facing surrogate until an authoritative Basic UDI-level source is confirmed.",
                required=False,
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.implantable",
                business_label="Implantable",
                value=self._normalized_boolean(reference_row.implantable),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Implantable",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.measuring_function",
                business_label="Measuring Function",
                value=self._normalized_boolean(reference_row.measuring_function),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Measuring function",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.reusable_surgical_instrument",
                business_label="Reusable Surgical Instrument",
                value=self._normalized_boolean(reference_row.reusable_surgical_instrument),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Reusable surgical instrument",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.reusable",
                business_label="Reusable",
                value=self._normalized_boolean(reference_row.reusable_surgical_instrument),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Reusable surgical instrument",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.active",
                business_label="Active Device",
                value=self._normalized_boolean(reference_row.active_device),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Active device",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.administering_medicinal_product",
                business_label="Administering Medicinal Product",
                value=self._normalized_boolean(reference_row.administering_medicinal_product),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Device intended to administer and/or remove medicinal product",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.administering_medicine",
                business_label="Administering Medicine",
                value=self._normalized_boolean(reference_row.administering_medicinal_product),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Device intended to administer and/or remove medicinal product",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.device_model_applicable",
                business_label="Device Model Applicable",
                value=self._normalized_yes_no(reference_row.device_model_applicable),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Device model applicable",
                required=False,
            ),
            self._field(
                canonical_path="basic_device.additional_information_url",
                business_label="Additional Information URL",
                value=reference_row.additional_information_url,
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!URL for additional information (as electronic instructions for use):",
                required=False,
            ),
            self._field(
                canonical_path="basic_device.submission_operation",
                business_label="Submission Operation",
                value=reference_row.submission_operation,
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Operation",
                xml_required=True,
            ),
            self._field(
                canonical_path="basic_device.source_version_marker",
                business_label="Source Version Marker",
                value=reference_row.source_version_marker,
                source="basic_udi_reference" if reference_row.source_version_marker else "missing",
                source_detail="BasicUDIs.xlsx!Version",
                review_note="Preserved as workbook metadata until QMS confirms whether it should populate a true EUDAMED entity-version field.",
                required=False,
            ),
            self._field(
                canonical_path="basic_device.type",
                business_label="Basic Device Type",
                value=self._device_type(reference_row.device_type),
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Is it a System or Procedure Pack which is a Device in itself?",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.basic_device_ref",
                business_label="Basic Device Reference",
                value=reference_row.basic_udi_di,
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Basic UDI-DI code",
                review_note="Explicitly links the UDI-DI row to the matched Basic UDI-DI variant context.",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.identifier",
                business_label="UDI-DI Identifier",
                value=self._di_identifier(issuing_entity=issuing_entity, di_code=primary_udi_di),
                source="derived",
                source_detail=f"{row.sheet_name}!{ISSUING_ENTITY_HEADER} + UDI-DI code",
                review_note="Composed as a DI identifier using workbook issuing entity and primary UDI-DI code.",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.primary_udi_di",
                business_label="Primary UDI-DI",
                value=primary_udi_di,
                source="workbook",
                source_detail=f"{row.sheet_name}!UDI-DI code",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.catalogue_number",
                business_label="Catalogue Number",
                value=self._catalogue_number(source_values),
                source="workbook",
                source_detail=f"{row.sheet_name}!Reference/ Catalogue number",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.basic_udi_identifier",
                business_label="Basic UDI Identifier",
                value=basic_udi_identifier,
                source="derived",
                source_detail="BasicUDIs.xlsx!Basic UDI-DI code + workbook issuing entity",
                review_note="Composed as a DI identifier using the matched Basic UDI-DI and workbook issuing entity.",
                required=False,
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.trade_name",
                business_label="Trade Name",
                value=self._string_value(source_values.get(TRADE_NAME_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{TRADE_NAME_HEADER}",
                required=False,
            ),
            self._field(
                canonical_path="device_record.language",
                business_label="Language",
                value=self._string_value(source_values.get(LANGUAGE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{LANGUAGE_HEADER}",
                required=False,
            ),
            self._field(
                canonical_path="device_record.status",
                business_label="UDI-DI Status",
                value=self._status_code(source_values.get(STATUS_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{STATUS_HEADER}",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.quantity",
                business_label="Quantity",
                value=self._string_value(source_values.get(QUANTITY_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{QUANTITY_HEADER}",
                required=False,
            ),
            self._field(
                canonical_path="device_record.udi_pi_type",
                business_label="UDI-PI Type",
                value=self._string_value(source_values.get(UDI_PI_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{UDI_PI_HEADER}",
                required=False,
            ),
            self._field(
                canonical_path="device_record.clinical_size_applicable",
                business_label="Clinical Size Applicable",
                value=self._normalized_boolean(source_values.get(CLINICAL_SIZE_APPLICABLE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{CLINICAL_SIZE_APPLICABLE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.production_identifier",
                business_label="Production Identifier",
                value=production_identifier,
                source="normalized" if production_identifier else "missing",
                source_detail=f"{row.sheet_name}!{UDI_PI_HEADER}",
                review_note="Optional XML-facing value normalized from the UDI-PI workbook selection.",
                required=False,
            ),
            self._field(
                canonical_path="device_record.number_of_reuses",
                business_label="Number Of Reuses",
                value=self._number_of_reuses(source_values),
                source="derived",
                source_detail=f"{row.sheet_name}!{MAX_REUSES_APPLICABLE_HEADER}",
                review_note="Derived from the workbook applicability field for current canonical validation.",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.max_reuses_applicable",
                business_label="Maximum Reuses Applicable",
                value=self._normalized_boolean(source_values.get(MAX_REUSES_APPLICABLE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{MAX_REUSES_APPLICABLE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.secondary_udi_di_applicable",
                business_label="Secondary UDI-DI Applicable",
                value=self._normalized_yes_no(source_values.get(SECONDARY_UDI_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{SECONDARY_UDI_HEADER}",
                required=False,
            ),
            self._field(
                canonical_path="device_record.secondary_identifier",
                business_label="Secondary Identifier",
                value=secondary_identifier,
                source="workbook" if secondary_identifier else "missing",
                source_detail=(
                    f"{row.sheet_name}!{SECONDARY_IDENTIFIER_ENTITY_HEADER} + {SECONDARY_IDENTIFIER_CODE_HEADER}"
                ),
                review_note="Only populated when the workbook provides both a secondary issuing entity and a secondary identifier code.",
                required=False,
            ),
            self._field(
                canonical_path="device_record.sterile",
                business_label="Sterile",
                value=self._normalized_boolean(source_values.get(STERILE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{STERILE_HEADER}",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.sterilization",
                business_label="Sterilization Before Use",
                value=self._normalized_boolean(source_values.get(STERILIZATION_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{STERILIZATION_HEADER}",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.sterilisation_before_use",
                business_label="Sterilisation Before Use",
                value=self._normalized_boolean(source_values.get(STERILIZATION_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{STERILIZATION_HEADER}",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.contains_latex",
                business_label="Contains Latex",
                value=self._normalized_boolean(source_values.get(LATEX_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{LATEX_HEADER}",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.reprocessed",
                business_label="Reprocessed",
                value=self._normalized_boolean(source_values.get(REPROCESSED_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{REPROCESSED_HEADER}",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.reprocessed_single_use",
                business_label="Reprocessed Single Use Device",
                value=self._normalized_boolean(source_values.get(REPROCESSED_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{REPROCESSED_HEADER}",
                xml_required=True,
            ),
            self._field(
                canonical_path="device_record.direct_marking",
                business_label="Direct Marking",
                value=self._normalized_boolean(source_values.get(DIRECT_MARKING_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{DIRECT_MARKING_HEADER}",
                required=False,
            ),
            self._field(
                canonical_path="device_record.single_use",
                business_label="Single Use",
                value=self._normalized_boolean(source_values.get(SINGLE_USE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{SINGLE_USE_HEADER}",
                required=False,
            ),
            self._field(
                canonical_path="device_record.cmr_present",
                business_label="CMR Present",
                value=self._normalized_boolean(source_values.get(CMR_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{CMR_HEADER}",
            ),
            self._field(
                canonical_path="device_record.endocrine_disruptor_present",
                business_label="Endocrine Disruptor Present",
                value=self._normalized_boolean(source_values.get(ENDOCRINE_DISRUPTOR_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{ENDOCRINE_DISRUPTOR_HEADER}",
            ),
            self._field(
                canonical_path="device_record.human_tissue_present",
                business_label="Human Tissue Present",
                value=self._normalized_boolean(source_values.get(HUMAN_TISSUE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{HUMAN_TISSUE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.animal_tissue_present",
                business_label="Animal Tissue Present",
                value=self._normalized_boolean(source_values.get(ANIMAL_TISSUE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{ANIMAL_TISSUE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.medicinal_substance_present",
                business_label="Medicinal Substance Present",
                value=self._normalized_boolean(source_values.get(MEDICINAL_SUBSTANCE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{MEDICINAL_SUBSTANCE_HEADER}",
                required=False,
            ),
            self._field(
                canonical_path="device_record.blood_plasma_derivative_present",
                business_label="Blood Plasma Derivative Present",
                value=self._normalized_boolean(source_values.get(HUMAN_BLOOD_SUBSTANCE_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{HUMAN_BLOOD_SUBSTANCE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.market_availability",
                business_label="Market Availability",
                value="present" if market_availability_items else None,
                source="basic_udi_reference" if market_availability_items else "missing",
                source_detail="BasicUDIs.xlsx!Member State of the placing on the EU market of the Device + available market country list",
                review_note="Represented as a nested canonical object with a market status and first-EU-market-country subfield.",
                required=False,
            ),
            self._field(
                canonical_path="device_record.market_availability.market_status",
                business_label="UDI-DI Market Status",
                value=self._status_code(source_values.get(STATUS_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{STATUS_HEADER}",
            ),
            self._field(
                canonical_path="device_record.market_availabilities",
                business_label="Market Availability Items",
                value=str(len(market_availability_items)) if market_availability_items else None,
                source="basic_udi_reference",
                source_detail="BasicUDIs.xlsx!Member States where device is or is to be made available on the market:",
                review_note="Represented as repeated market-availability items in the canonical layer.",
                required=False,
            ),
            self._field(
                canonical_path="device_record.market_availability.first_eu_market_country",
                business_label="First EU Market Country",
                value=reference_row.first_eu_market_country or self._string_value(source_values.get(FIRST_EU_MARKET_HEADER)),
                source="basic_udi_reference" if reference_row.first_eu_market_country else "workbook",
                source_detail="BasicUDIs.xlsx!Member State of the placing on the EU market of the Device:",
                required=False,
            ),
            self._field(
                canonical_path="device_record.base_quantity",
                business_label="Base Quantity",
                value=self._string_value(source_values.get(QUANTITY_HEADER)),
                source="workbook",
                source_detail=f"{row.sheet_name}!{QUANTITY_HEADER}",
                review_note="Optional XML-facing quantity field currently taken directly from the workbook quantity column.",
                required=False,
            ),
            self._field(
                canonical_path="device_record.storage_conditions",
                business_label="Storage Conditions",
                value=str(len(storage_condition_items)) if storage_condition_items else None,
                source="derived" if storage_condition_items else "missing",
                source_detail=f"{row.sheet_name}!{STORAGE_APPLICABLE_HEADER}",
                required=self._normalized_yes_no(source_values.get(STORAGE_APPLICABLE_HEADER)) == "Yes",
            ),
            self._field(
                canonical_path="device_record.warnings",
                business_label="Critical Warnings",
                value=str(len(critical_warning_items)) if critical_warning_items else None,
                source="derived" if critical_warning_items else "missing",
                source_detail=f"{row.sheet_name}!{WARNING_APPLICABLE_HEADER}",
                review_note="Represented as repeated warning items in the canonical layer.",
                required=self._normalized_yes_no(source_values.get(WARNING_APPLICABLE_HEADER)) == "Yes",
            ),
        ]

        blockers = [f"{field.business_label} is not populated." for field in fields if field.required and field.value is None]
        completeness = self._completeness_snapshot(fields)
        xml_blockers = [
            f"{field.business_label} is not populated for XML generation."
            for field in fields
            if field.xml_required and field.value is None
        ]
        xml_readiness = self._xml_readiness_snapshot(fields)

        return CanonicalValidationRecord(
            source_workbook=workbook_name,
            product_family=product_family,
            product_variant=reference_row.device_model,
            source_sheet=row.sheet_name,
            source_row_index=row.row_index,
            trade_name=self._string_value(source_values.get(TRADE_NAME_HEADER)),
            primary_udi_di=primary_udi_di,
            catalogue_number=self._catalogue_number(source_values),
            issuing_entity=issuing_entity,
            submission_operation=reference_row.submission_operation,
            reference_match_status="matched",
            completeness=completeness,
            xml_readiness=xml_readiness,
            blockers=blockers,
            xml_blockers=xml_blockers,
            storage_condition_items=storage_condition_items,
            critical_warning_items=critical_warning_items,
            market_availability_items=market_availability_items,
            fields=fields,
        )

    def _load_source_rows(self, workbook_path: Path) -> list[SourceRow]:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        rows: list[SourceRow] = []
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            header_row_index: int | None = None
            headers: list[str] | None = None
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = list(row)
                if headers is None and any(
                    isinstance(value, str) and HEADER_SENTINEL in value for value in values if value
                ):
                    headers = [str(value).strip() if value is not None else "" for value in values]
                    header_row_index = row_index
                    continue
                if headers is None or header_row_index is None or row_index <= header_row_index:
                    continue
                if not any(value not in (None, "") for value in values):
                    continue
                mapped_values = {
                    headers[index]: values[index] if index < len(values) else None
                    for index in range(len(headers))
                    if headers[index]
                }
                rows.append(SourceRow(sheet_name=sheet_name, row_index=row_index, values=mapped_values))
        return rows

    def _load_source_headers(self, workbook_path: Path) -> list[SourceHeaderProfile]:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        headers_by_name: dict[str, SourceHeaderProfile] = {}
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() if value is not None else "" for value in row]
                if not any(HEADER_SENTINEL in value for value in values if value):
                    continue
                for value in values:
                    if not value:
                        continue
                    normalized = self._normalize_header(value)
                    entry = headers_by_name.get(normalized)
                    if entry is None:
                        headers_by_name[normalized] = SourceHeaderProfile(
                            display_name=value,
                            normalized_name=normalized,
                            sheets=(sheet_name,),
                        )
                    elif sheet_name not in entry.sheets:
                        headers_by_name[normalized] = SourceHeaderProfile(
                            display_name=entry.display_name,
                            normalized_name=entry.normalized_name,
                            sheets=tuple(sorted((*entry.sheets, sheet_name))),
                        )
                break
        return [entry for _, entry in sorted(headers_by_name.items())]

    def _build_source_field_coverage(
        self,
        headers: list[SourceHeaderProfile],
    ) -> list[SourceFieldCoverageEntry]:
        review_bundle = CanonicalReviewService().load_review_bundle()
        source_mappings: dict[str, list[tuple[str, list[str]]]] = {}
        for entity in review_bundle.entity_reviews:
            for field_review in entity.field_reviews:
                mapping = field_review.mapping
                schema_targets = [target.schema_path for target in mapping.schema_targets]
                for source_column in mapping.source_columns:
                    normalized = self._normalize_header(source_column)
                    source_mappings.setdefault(normalized, []).append((mapping.canonical_path, schema_targets))

        coverage_entries: list[SourceFieldCoverageEntry] = []
        for header in headers:
            mappings = source_mappings.get(header.normalized_name, [])
            canonical_targets = sorted({canonical_path for canonical_path, _ in mappings})
            schema_targets = sorted({target for _, targets in mappings for target in targets})
            status: SourceFieldCoverageStatus = "represented" if mappings else "not_yet_represented"
            notes = (
                "This source field is mapped into the active canonical review path."
                if mappings
                else "This source field does not yet have a documented canonical target in the active review artifact."
            )
            coverage_entries.append(
                SourceFieldCoverageEntry(
                    source_field=header.display_name,
                    source_sheets=list(header.sheets),
                    coverage_status=status,
                    canonical_targets=canonical_targets,
                    schema_targets=schema_targets,
                    notes=notes,
                )
            )
        return coverage_entries

    def _assemble_storage_condition_items(
        self,
        values: dict[str, object | None],
    ) -> list[StructuredListItemPreview]:
        if self._normalized_yes_no(values.get(STORAGE_APPLICABLE_HEADER)) != "Yes":
            return []
        items: list[StructuredListItemPreview] = []
        candidates = (
            (1, STORAGE_TYPE_HEADER, STORAGE_DESCRIPTION_ONE_HEADER),
            (2, STORAGE_TYPE_TWO_HEADER, STORAGE_DESCRIPTION_TWO_HEADER),
        )
        for sequence, type_header, description_header in candidates:
            item_type = self._string_value(values.get(type_header))
            description = self._string_value(values.get(description_header))
            if item_type is None and description is None:
                continue
            items.append(
                StructuredListItemPreview(
                    sequence=sequence,
                    item_type=item_type,
                    normalized_code=self._normalized_rule_value(type_header, item_type),
                    description=description,
                    source_fields=[type_header, description_header],
                )
            )
        return items

    def _assemble_critical_warning_items(
        self,
        values: dict[str, object | None],
    ) -> list[StructuredListItemPreview]:
        if self._normalized_yes_no(values.get(WARNING_APPLICABLE_HEADER)) != "Yes":
            return []
        warning_type = self._string_value(values.get(WARNING_TYPE_HEADER))
        if warning_type is None:
            return []
        return [
            StructuredListItemPreview(
                sequence=1,
                item_type=warning_type,
                normalized_code=self._normalized_rule_value(WARNING_TYPE_HEADER, warning_type),
                description=None,
                source_fields=[WARNING_TYPE_HEADER],
            )
        ]

    @staticmethod
    def _assemble_market_availability_items(
        reference_row: BasicUdiReferenceRow,
    ) -> list[MarketAvailabilityItemPreview]:
        return [
            MarketAvailabilityItemPreview(
                sequence=index,
                country=country,
                original_placed_on_market=country == reference_row.first_eu_market_country,
            )
            for index, country in enumerate(reference_row.available_market_countries, start=1)
        ]

    @staticmethod
    def _build_family_summaries(records: list[CanonicalValidationRecord]) -> list[FamilyValidationSummary]:
        grouped: dict[str, list[CanonicalValidationRecord]] = {}
        for record in records:
            grouped.setdefault(record.product_family, []).append(record)
        summaries: list[FamilyValidationSummary] = []
        for product_family, family_records in sorted(grouped.items()):
            summaries.append(
                FamilyValidationSummary(
                    product_family=product_family,
                    variant_count=len({record.product_variant for record in family_records}),
                    total_records=len(family_records),
                    ready_records=sum(1 for record in family_records if record.completeness.status == "complete"),
                    blocked_records=sum(1 for record in family_records if record.completeness.status == "incomplete"),
                    xml_ready_records=sum(1 for record in family_records if record.xml_readiness.status == "complete"),
                    xml_blocked_records=sum(1 for record in family_records if record.xml_readiness.status == "incomplete"),
                    post_records=sum(1 for record in family_records if record.submission_operation == "POST"),
                    patch_records=sum(1 for record in family_records if record.submission_operation == "PATCH"),
                )
            )
        return summaries

    @staticmethod
    def _build_variant_summaries(records: list[CanonicalValidationRecord]) -> list[VariantValidationSummary]:
        grouped: dict[tuple[str, str], list[CanonicalValidationRecord]] = {}
        for record in records:
            grouped.setdefault((record.product_family, record.product_variant), []).append(record)
        summaries: list[VariantValidationSummary] = []
        for (_, _), variant_records in sorted(grouped.items()):
            first = variant_records[0]
            blocker_counts: dict[str, int] = {}
            xml_blocker_counts: dict[str, int] = {}
            for record in variant_records:
                for blocker in record.blockers:
                    blocker_counts[blocker] = blocker_counts.get(blocker, 0) + 1
                for blocker in record.xml_blockers:
                    xml_blocker_counts[blocker] = xml_blocker_counts.get(blocker, 0) + 1
            summaries.append(
                VariantValidationSummary(
                    product_family=first.product_family,
                    product_variant=first.product_variant,
                    source_workbook=first.source_workbook,
                    source_sheet=first.source_sheet,
                    submission_operation=first.submission_operation,
                    total_records=len(variant_records),
                    ready_records=sum(1 for record in variant_records if record.completeness.status == "complete"),
                    blocked_records=sum(1 for record in variant_records if record.completeness.status == "incomplete"),
                    xml_ready_records=sum(1 for record in variant_records if record.xml_readiness.status == "complete"),
                    xml_blocked_records=sum(1 for record in variant_records if record.xml_readiness.status == "incomplete"),
                    missing_required_field_total=sum(
                        record.completeness.missing_required_fields for record in variant_records
                    ),
                    missing_xml_required_field_total=sum(
                        record.xml_readiness.missing_required_fields for record in variant_records
                    ),
                    common_blockers=[
                        blocker
                        for blocker, _ in sorted(blocker_counts.items(), key=lambda item: (-item[1], item[0]))[:3]
                    ],
                    common_xml_blockers=[
                        blocker
                        for blocker, _ in sorted(xml_blocker_counts.items(), key=lambda item: (-item[1], item[0]))[:3]
                    ],
                )
            )
        return summaries

    @staticmethod
    def _build_blocker_summaries(records: list[CanonicalValidationRecord]) -> list[ValidationBlockerSummary]:
        if not records:
            return []
        template_fields = records[0].fields
        summaries: list[ValidationBlockerSummary] = []
        for template_field in template_fields:
            missing_count = sum(
                1
                for record in records
                for field in record.fields
                if field.canonical_path == template_field.canonical_path and field.required and field.value is None
            )
            summaries.append(
                ValidationBlockerSummary(
                    canonical_path=template_field.canonical_path,
                    business_label=template_field.business_label,
                    missing_count=missing_count,
                )
            )
        return summaries

    @staticmethod
    def _build_source_field_coverage_summaries(
        entries: list[SourceFieldCoverageEntry],
    ) -> list[SourceFieldCoverageSummary]:
        labels = {
            "represented": "Represented",
            "partially_represented": "Partially represented",
            "not_yet_represented": "Not yet represented",
            "deferred_by_design": "Deferred by design",
        }
        ordered_statuses: tuple[SourceFieldCoverageStatus, ...] = (
            "represented",
            "partially_represented",
            "not_yet_represented",
            "deferred_by_design",
        )
        return [
            SourceFieldCoverageSummary(
                status=status,
                label=labels[status],
                field_count=sum(1 for entry in entries if entry.coverage_status == status),
            )
            for status in ordered_statuses
        ]

    @staticmethod
    def _build_sample_records(records: list[CanonicalValidationRecord]) -> list[CanonicalValidationRecord]:
        grouped: dict[tuple[str, str], CanonicalValidationRecord] = {}
        for record in records:
            grouped.setdefault((record.product_family, record.product_variant), record)
        return [grouped[key] for key in sorted(grouped)]

    @staticmethod
    def _merge_header_profiles(headers: list[SourceHeaderProfile]) -> list[SourceHeaderProfile]:
        merged: dict[str, SourceHeaderProfile] = {}
        for header in headers:
            existing = merged.get(header.normalized_name)
            if existing is None:
                merged[header.normalized_name] = header
            else:
                merged[header.normalized_name] = SourceHeaderProfile(
                    display_name=existing.display_name,
                    normalized_name=existing.normalized_name,
                    sheets=tuple(sorted(set(existing.sheets).union(header.sheets))),
                )
        return [merged[key] for key in sorted(merged)]

    def _normalized_rule_value(self, column: str, raw_value: str | None) -> str | None:
        if raw_value in (None, ""):
            return None
        return self.normalization_repository.rules_for_column(column).get(raw_value)

    @staticmethod
    def _field(
        *,
        canonical_path: str,
        business_label: str,
        value: str | None,
        source: ValueSourceType,
        source_detail: str,
        review_note: str | None = None,
        required: bool = True,
        xml_required: bool = False,
    ) -> CanonicalValidationFieldValue:
        return CanonicalValidationFieldValue(
            canonical_path=canonical_path,
            business_label=business_label,
            required=required,
            xml_required=xml_required,
            value=value,
            source=source,
            source_detail=source_detail,
            review_note=review_note,
        )

    @staticmethod
    def _completeness_snapshot(fields: list[CanonicalValidationFieldValue]) -> CompletenessSnapshot:
        values = [field.value for field in fields if field.required]
        mapped_required_fields = sum(1 for value in values if value is not None)
        total_required_fields = len(values)
        missing_required_fields = total_required_fields - mapped_required_fields
        return CompletenessSnapshot(
            mapped_required_fields=mapped_required_fields,
            total_required_fields=total_required_fields,
            missing_required_fields=missing_required_fields,
            status="complete" if missing_required_fields == 0 else "incomplete",
        )

    @staticmethod
    def _xml_readiness_snapshot(fields: list[CanonicalValidationFieldValue]) -> CompletenessSnapshot:
        values = [field.value for field in fields if field.xml_required]
        mapped_required_fields = sum(1 for value in values if value is not None)
        total_required_fields = len(values)
        missing_required_fields = total_required_fields - mapped_required_fields
        return CompletenessSnapshot(
            mapped_required_fields=mapped_required_fields,
            total_required_fields=total_required_fields,
            missing_required_fields=missing_required_fields,
            status="complete" if missing_required_fields == 0 else "incomplete",
        )

    @staticmethod
    def _catalogue_number(values: dict[str, object | None]) -> str | None:
        for header in CATALOGUE_HEADERS:
            value = values.get(header)
            if value not in (None, ""):
                return str(value).strip()
        return None

    @staticmethod
    def _first_string_value(
        values: dict[str, object | None],
        headers: tuple[str, ...],
    ) -> str | None:
        for header in headers:
            value = CanonicalValidationService._string_value(values.get(header))
            if value is not None:
                return value
        return None

    @staticmethod
    def _di_identifier(*, issuing_entity: str | None, di_code: str | None) -> str | None:
        if not issuing_entity or not di_code:
            return None
        return f"{issuing_entity}:{di_code}"

    @staticmethod
    def _issuing_entity_code(value: object | None) -> str | None:
        raw = CanonicalValidationService._string_value(value)
        if raw is None:
            return None
        mapping = {
            "gs1": "GS1",
            "hibcc": "HIBCC",
            "iccba": "ICCBBA",
            "ifi": "IFI",
        }
        return mapping.get(raw.lower(), raw.upper())

    @staticmethod
    def _status_code(value: object | None) -> str | None:
        raw = CanonicalValidationService._string_value(value)
        if raw is None:
            return None
        normalized = raw.strip().upper().replace(" ", "_")
        mapping = {
            "ON_THE_EU_MARKET": "ON_THE_MARKET",
            "ON_THE_MARKET": "ON_THE_MARKET",
        }
        return mapping.get(normalized, normalized)

    @staticmethod
    def _device_type(value: str | None) -> str | None:
        raw = CanonicalValidationService._string_value(value)
        if raw is None:
            return None
        token = raw.lower()
        if token == "no":
            return "DEVICE"
        if "procedure" in token:
            return "PROCEDURE_PACK"
        if "system" in token:
            return "SYSTEM"
        return raw

    @staticmethod
    def _regulation_code(value: str | None) -> str | None:
        raw = CanonicalValidationService._string_value(value)
        if raw is None:
            return None
        if raw.upper().startswith("MDR"):
            return "MDR"
        return raw

    @staticmethod
    def _number_of_reuses(values: dict[str, object | None]) -> str | None:
        applicable = CanonicalValidationService._string_value(values.get(MAX_REUSES_APPLICABLE_HEADER))
        single_use = CanonicalValidationService._string_value(values.get(SINGLE_USE_HEADER))
        if single_use and single_use.strip().lower() in {"yes", "true"}:
            return "0"
        if applicable and applicable.strip().lower() in {"no", "false"}:
            return "1"
        return None

    @staticmethod
    def _production_identifier(value: object | None) -> str | None:
        raw = CanonicalValidationService._string_value(value)
        if raw is None:
            return None
        normalized = raw.lower().replace(" / ", "/").replace("/ ", "/").replace(" /", "/")
        mapping = {
            "serial number/ manufacturing date": "SERIALISATION_NUMBER",
            "serial number/manufacturing date": "SERIALISATION_NUMBER",
        }
        return mapping.get(normalized, raw.upper().replace(" ", "_"))

    @staticmethod
    def _string_value(value: object | None) -> str | None:
        if value in (None, ""):
            return None
        return str(value).strip()

    @staticmethod
    def _normalized_boolean(value: object | None) -> str | None:
        if value in (None, ""):
            return None
        token = str(value).strip().lower()
        if token in {"yes", "true"}:
            return "true"
        if token in {"no", "false"}:
            return "false"
        return str(value).strip()

    @staticmethod
    def _normalized_yes_no(value: object | None) -> str | None:
        if value in (None, ""):
            return None
        token = str(value).strip().lower()
        if token in {"yes", "true"}:
            return "Yes"
        if token in {"no", "false"}:
            return "No"
        return str(value).strip()

    @staticmethod
    def _normalize_header(value: str) -> str:
        return " ".join(value.split()).strip().lower()
