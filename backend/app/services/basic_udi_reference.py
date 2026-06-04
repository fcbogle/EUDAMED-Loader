from __future__ import annotations

from dataclasses import dataclass

from openpyxl import load_workbook

from app.canonical_models import SubmissionOperation, VariantMappingSummary
from app.config import get_settings
from app.validation_models import ValueSourceType


@dataclass(frozen=True)
class BasicUdiReferenceRow:
    applicable_regulation: str | None
    issuing_entity: str | None
    device_model: str
    basic_udi_di: str
    manufacturer_srn: str | None
    manufacturer_srn_source: ValueSourceType | None
    authorised_representative_srn: str | None
    authorised_representative_srn_source: ValueSourceType | None
    device_type: str | None
    special_device_type: str | None
    risk_class: str | None
    implantable: str | None
    measuring_function: str | None
    reusable_surgical_instrument: str | None
    active_device: str | None
    administering_medicinal_product: str | None
    device_model_applicable: str | None
    additional_information_url: str | None
    submission_operation: SubmissionOperation | None
    source_version_marker: str | None
    first_eu_market_country: str | None
    available_market_countries: tuple[str, ...]


class BasicUdiReferenceService:
    def __init__(self) -> None:
        self.settings = get_settings()

    def list_variant_mappings(self) -> list[VariantMappingSummary]:
        reference_rows = self.rows_by_device_model()
        mappings: list[VariantMappingSummary] = []

        for workbook_path in sorted(self.settings.excel_dir.glob("*.xlsx")):
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
            workbook_excluded = workbook_path.name in self.settings.excluded_excel_workbook_names
            exclusion_notes = (
                ["Excluded from active variant-level mapping pending QMS clarification."]
                if workbook_excluded
                else []
            )

            for worksheet in workbook.worksheets:
                reference = reference_rows.get(worksheet.title)
                if workbook_excluded:
                    mappings.append(
                        VariantMappingSummary(
                            workbook=workbook_path.name,
                            sheet=worksheet.title,
                            device_model=None,
                            basic_udi_di=None,
                            submission_operation=None,
                            source_version_marker=None,
                            first_eu_market_country=None,
                            available_market_country_count=0,
                            match_status="excluded",
                            notes=exclusion_notes,
                        )
                    )
                    continue

                if reference is None:
                    mappings.append(
                        VariantMappingSummary(
                            workbook=workbook_path.name,
                            sheet=worksheet.title,
                            device_model=None,
                            basic_udi_di=None,
                            submission_operation=None,
                            source_version_marker=None,
                            first_eu_market_country=None,
                            available_market_country_count=0,
                            match_status="unmatched",
                            notes=[
                                "No exact Device Model match was found in the authoritative BasicUDIs.xlsx workbook."
                            ],
                        )
                    )
                    continue

                mappings.append(
                    VariantMappingSummary(
                        workbook=workbook_path.name,
                        sheet=worksheet.title,
                        device_model=reference.device_model,
                        basic_udi_di=reference.basic_udi_di,
                        submission_operation=reference.submission_operation,
                        source_version_marker=reference.source_version_marker,
                        first_eu_market_country=reference.first_eu_market_country,
                        available_market_country_count=len(reference.available_market_countries),
                        match_status="matched",
                        notes=[
                            "Exact sheet-to-Device Model match from the authoritative Basic UDI workbook.",
                            *(
                                [
                                    "Manufacturer and authorised representative SRN are supplemented from the legacy tracekey Basic UDI workbook."
                                ]
                                if reference.manufacturer_srn_source == "legacy_basic_udi_reference"
                                else []
                            ),
                        ],
                    )
                )

        return mappings

    def rows_by_device_model(self) -> dict[str, BasicUdiReferenceRow]:
        legacy_srn_rows = self._legacy_srn_rows_by_device_model()
        workbook = load_workbook(self.settings.basic_udi_reference_workbook, read_only=True, data_only=True)
        references: dict[str, BasicUdiReferenceRow] = {}
        if "BasicUDI" in workbook.sheetnames:
            self._load_legacy_basic_udi_sheet(workbook, legacy_srn_rows, references)
        else:
            self._load_current_basic_udi_sheets(workbook, legacy_srn_rows, references)
        return references

    def _legacy_srn_rows_by_device_model(self) -> dict[str, BasicUdiReferenceRow]:
        workbook = load_workbook(self.settings.legacy_basic_udi_reference_workbook, read_only=True, data_only=True)

        products_sheet = workbook["Products"]
        products_rows = list(products_sheet.iter_rows(values_only=True))
        products_headers = [self._stringify(cell) for cell in products_rows[0]]
        products_index = {header: i for i, header in enumerate(products_headers) if header}

        basic_sheet = workbook["MDR-Basic"]
        basic_rows = list(basic_sheet.iter_rows(values_only=True))
        basic_headers = [self._stringify(cell) for cell in basic_rows[0]]
        basic_index = {header: i for i, header in enumerate(basic_headers) if header}

        basic_by_material_number = {
            self._cell(row, basic_index, "Material Number"): row
            for row in basic_rows[1:]
            if self._cell(row, basic_index, "Material Number")
        }

        legacy_rows: dict[str, BasicUdiReferenceRow] = {}
        for row in products_rows[1:]:
            if self._cell(row, products_index, "Product Template") != "EUDAMED Basic UDI-DI":
                continue
            device_model = self._cell(row, products_index, "Name")
            material_number = self._cell(row, products_index, "Material Number")
            if not device_model or not material_number:
                continue
            basic_row = basic_by_material_number.get(material_number)
            if basic_row is None:
                continue
            legacy_rows[device_model] = BasicUdiReferenceRow(
                applicable_regulation=None,
                issuing_entity=None,
                device_model=device_model,
                basic_udi_di=self._cell(row, products_index, "Product Code (Value)"),
                manufacturer_srn=self._cell(basic_row, basic_index, "Manufacturer Code") or None,
                manufacturer_srn_source="legacy_basic_udi_reference",
                authorised_representative_srn=(
                    self._cell(basic_row, basic_index, "Authorised Representative Actor Code") or None
                ),
                authorised_representative_srn_source="legacy_basic_udi_reference",
                device_type=self._cell(basic_row, basic_index, "Type") or None,
                special_device_type=self._cell(basic_row, basic_index, "Special Device") or None,
                risk_class=self._cell(basic_row, basic_index, "Risk Class") or None,
                implantable=self._cell(basic_row, basic_index, "Implantable") or None,
                measuring_function=self._cell(basic_row, basic_index, "Measuring Function") or None,
                reusable_surgical_instrument=self._cell(basic_row, basic_index, "Reusable") or None,
                active_device=self._cell(basic_row, basic_index, "Active") or None,
                administering_medicinal_product=(
                    self._cell(basic_row, basic_index, "Administers or/and Removes Medicine") or None
                ),
                device_model_applicable=None,
                additional_information_url=None,
                submission_operation=None,
                source_version_marker=None,
                first_eu_market_country=None,
                available_market_countries=(),
            )
        return legacy_rows

    def _load_legacy_basic_udi_sheet(
        self,
        workbook,
        legacy_srn_rows: dict[str, BasicUdiReferenceRow],
        references: dict[str, BasicUdiReferenceRow],
    ) -> None:
        worksheet = workbook["BasicUDI"]
        rows = list(worksheet.iter_rows(values_only=True))
        headers = [self._stringify(cell) for cell in rows[0]]
        index = {header: i for i, header in enumerate(headers) if header}

        for row in rows[1:]:
            device_model = self._cell(row, index, "Device Model")
            if not device_model:
                continue
            self._add_reference_row(
                references=references,
                legacy_srn_rows=legacy_srn_rows,
                device_model=device_model,
                basic_udi_di=self._cell(row, index, "Basic UDI-DI code"),
                applicable_regulation=self._cell(row, index, "Applicable regulation") or None,
                issuing_entity=self._cell(row, index, "Issuing Entity") or None,
                device_type=self._cell(
                    row,
                    index,
                    "Is it a System or Procedure Pack which is a Device in itself?",
                )
                or None,
                special_device_type=self._cell(row, index, "Special device type") or None,
                risk_class=self._cell(row, index, "Risk class") or None,
                implantable=self._cell(row, index, "Implantable") or None,
                measuring_function=self._cell(row, index, "Measuring function") or None,
                reusable_surgical_instrument=self._cell(row, index, "Reusable surgical instrument") or None,
                active_device=self._cell(row, index, "Active device") or None,
                administering_medicinal_product=self._cell(
                    row,
                    index,
                    "Device intended to administer and/or remove medicinal product",
                )
                or None,
                device_model_applicable=self._cell(row, index, "Device model applicable") or None,
                additional_information_url=self._cell(
                    row,
                    index,
                    "URL for additional information (as electronic instructions for use):",
                )
                or None,
                submission_operation=self._submission_operation(self._cell(row, index, "Operation")),
                source_version_marker=self._cell(row, index, "Version") or None,
                first_eu_market_country=self._cell(
                    row,
                    index,
                    "Member State of the placing on the EU market of the Device:",
                )
                or None,
                available_market_countries=self._market_countries(
                    self._cell(
                        row,
                        index,
                        "Member States where device is or is to be made available on the market:",
                    )
                ),
            )

    def _load_current_basic_udi_sheets(
        self,
        workbook,
        legacy_srn_rows: dict[str, BasicUdiReferenceRow],
        references: dict[str, BasicUdiReferenceRow],
    ) -> None:
        sheet_definitions = (
            ("Upload(BasicUDI not registered)", "POST", "1"),
            ("Update(BasicUDI registered)", "PATCH", "2"),
        )
        for sheet_name, operation_token, version_marker in sheet_definitions:
            if sheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet_name]
            rows = list(worksheet.iter_rows(values_only=True))
            headers = [self._stringify(cell) for cell in rows[0]]
            index = {header: i for i, header in enumerate(headers) if header}
            for row in rows[1:]:
                device_model = self._cell(row, index, "Device Model")
                if not device_model:
                    continue
                self._add_reference_row(
                    references=references,
                    legacy_srn_rows=legacy_srn_rows,
                    device_model=device_model,
                    basic_udi_di=self._cell(row, index, "Basic UDI-DI code"),
                    applicable_regulation=self._cell(row, index, "Applicable regulation") or None,
                    issuing_entity=self._cell(row, index, "Issuing Entity") or None,
                    device_type=self._cell(
                        row,
                        index,
                        "Is it a System or Procedure Pack which is a Device in itself?",
                    )
                    or None,
                    special_device_type=self._cell(row, index, "Special device type") or None,
                    risk_class=self._cell(row, index, "Risk class") or None,
                    implantable=self._cell(row, index, "Implantable") or None,
                    measuring_function=self._cell(row, index, "Measuring function") or None,
                    reusable_surgical_instrument=self._cell(row, index, "Reusable surgical instrument") or None,
                    active_device=self._cell(row, index, "Active device") or None,
                    administering_medicinal_product=self._cell(
                        row,
                        index,
                        "Device intended to administer and/or remove medicinal product",
                    )
                    or None,
                    device_model_applicable=self._cell(row, index, "Device model applicable") or None,
                    additional_information_url=self._cell(
                        row,
                        index,
                        "URL for additional information (as electronic instructions for use):",
                    )
                    or None,
                    submission_operation=self._submission_operation(operation_token),
                    source_version_marker=version_marker,
                    first_eu_market_country=self._cell(
                        row,
                        index,
                        "Member State of the placing on the EU market of the Device:",
                    )
                    or None,
                    available_market_countries=self._market_countries(
                        self._cell(
                            row,
                            index,
                            "Member States where device is or is to be made available on the market:",
                        )
                    ),
                )

    def _add_reference_row(
        self,
        *,
        references: dict[str, BasicUdiReferenceRow],
        legacy_srn_rows: dict[str, BasicUdiReferenceRow],
        device_model: str,
        basic_udi_di: str,
        applicable_regulation: str | None,
        issuing_entity: str | None,
        device_type: str | None,
        special_device_type: str | None,
        risk_class: str | None,
        implantable: str | None,
        measuring_function: str | None,
        reusable_surgical_instrument: str | None,
        active_device: str | None,
        administering_medicinal_product: str | None,
        device_model_applicable: str | None,
        additional_information_url: str | None,
        submission_operation: SubmissionOperation | None,
        source_version_marker: str | None,
        first_eu_market_country: str | None,
        available_market_countries: tuple[str, ...],
    ) -> None:
        legacy_srn_row = self._matching_legacy_srn_row(device_model, legacy_srn_rows)
        references[device_model] = BasicUdiReferenceRow(
            applicable_regulation=applicable_regulation,
            issuing_entity=issuing_entity,
            device_model=device_model,
            basic_udi_di=basic_udi_di,
            manufacturer_srn=legacy_srn_row.manufacturer_srn if legacy_srn_row else None,
            manufacturer_srn_source=legacy_srn_row.manufacturer_srn_source if legacy_srn_row else None,
            authorised_representative_srn=(
                legacy_srn_row.authorised_representative_srn if legacy_srn_row else None
            ),
            authorised_representative_srn_source=(
                legacy_srn_row.authorised_representative_srn_source if legacy_srn_row else None
            ),
            device_type=device_type,
            special_device_type=special_device_type,
            risk_class=risk_class,
            implantable=implantable,
            measuring_function=measuring_function,
            reusable_surgical_instrument=reusable_surgical_instrument,
            active_device=active_device,
            administering_medicinal_product=administering_medicinal_product,
            device_model_applicable=device_model_applicable,
            additional_information_url=additional_information_url,
            submission_operation=submission_operation,
            source_version_marker=source_version_marker,
            first_eu_market_country=first_eu_market_country,
            available_market_countries=available_market_countries,
        )

    @staticmethod
    def _market_countries(value: str) -> tuple[str, ...]:
        return tuple(country for country in (item.strip() for item in value.split(";")) if country)

    @staticmethod
    def _matching_legacy_srn_row(
        device_model: str,
        legacy_rows: dict[str, BasicUdiReferenceRow],
    ) -> BasicUdiReferenceRow | None:
        exact = legacy_rows.get(device_model)
        if exact is not None:
            return exact
        for legacy_model, legacy_row in legacy_rows.items():
            if device_model.startswith(f"{legacy_model} "):
                return legacy_row
        # Current QMS confirmation: the legacy tracekey SRN values are valid for all products.
        if len(legacy_rows) == 1:
            return next(iter(legacy_rows.values()))
        return None

    @staticmethod
    def _cell(row: tuple[object | None, ...], index: dict[str, int], header: str) -> str:
        position = index.get(header)
        if position is None or position >= len(row):
            return ""
        return BasicUdiReferenceService._stringify(row[position])

    @staticmethod
    def _stringify(value: object | None) -> str:
        if value is None:
            return ""
        return " ".join(str(value).split())

    @staticmethod
    def _submission_operation(value: str) -> SubmissionOperation | None:
        token = value.strip().upper()
        if token == "POST":
            return "POST"
        if token == "PATCH":
            return "PATCH"
        if token == "PUT":
            return "PUT"
        if token == "GET":
            return "GET"
        return None
