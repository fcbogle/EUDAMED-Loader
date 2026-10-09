"""Prepare a reviewable production workbook without opening application SQLite.

Template values and exported acceptance are independent evidence. The companion
JSON retains every source occurrence, including conflicting duplicate rows.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime, UTC
import hashlib
import json
from pathlib import Path
import re
from typing import Any, cast

import lxml.etree as etree
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

BASE = "https://ec.europa.eu/tools/eudamed/dtx/"
NS = {
    "message": BASE + "servicemodel/Message/v1",
    "device": BASE + "datamodel/Entity/Device/v1",
    "entity": BASE + "datamodel/Entity/v1",
    "basic": BASE + "datamodel/Entity/Device/BasicUDI/v1",
    "common": BASE + "datamodel/Entity/Device/CommonDevice/v1",
    "udi": BASE + "datamodel/Entity/UDIDI/v1",
}
CONTROLS = [
    "Device Type", "Registration Status", "Review Required", "Review Reasons",
    "Issuing Entity", "UDI-DI", "Proposed Basic UDI-DI", "Proposed Parent Issuer",
    "Parent Reference Sources", "Template Sources", "Source Occurrence Count",
]
ACCEPTED = [
    "Accepted Basic UDI-DI", "Accepted Parent Issuer", "Parent State",
    "Parent Version", "Parent Version Date", "Device State", "Device Version",
    "Device Version Date", "Market Info State", "Market Info Version",
    "Market Info Version Date", "Export Sources", "Export Created At",
    "Encoding Overrides", "Extracted At",
]


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _serial(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return {"type": type(value).__name__, "value": value.isoformat()}
    return value


def _header_keys(headers: list[Any]) -> list[str]:
    counts: Counter[str] = Counter()
    result = []
    for index, value in enumerate(headers, 1):
        label = _text(value) or f"Unnamed column {index}"
        counts[label] += 1
        result.append(label if counts[label] == 1 else f"{label} [occurrence {counts[label]}]")
    return result


def _identity(issuer: Any, code: Any) -> tuple[str, str]:
    return (_text(issuer).upper(), _text(code))


def _device_type(path: Path, sheet: str) -> str:
    if sheet != "Sheet1":
        return sheet
    label = re.sub(r"^Template for\s+", "", path.stem)
    return re.split(r"\s+(?:on\s+)?EUDAMED\b", label)[0]


def _parent_reference(path: Path) -> tuple[dict, dict]:
    by_model: dict[str, list[dict]] = defaultdict(list)
    by_code: dict[tuple[str, str], list[dict]] = defaultdict(list)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            rows = iter(sheet.iter_rows(values_only=True))
            headers = list(next(rows, ()))
            if not {"Issuing Entity", "Basic UDI-DI code", "Device Model"}.issubset(headers):
                raise ValueError(f"Unsupported parent reference sheet: {sheet.title}")
            keys = _header_keys(headers)
            for row_number, values in enumerate(rows, 2):
                if not any(v is not None for v in values):
                    continue
                fields = dict(zip(keys, values))
                key = _identity(fields["Issuing Entity"], fields["Basic UDI-DI code"])
                if not all(key):
                    raise ValueError(f"Incomplete parent identity: {sheet.title}:{row_number}")
                entry = {"key": key, "fields": {k: _serial(v) for k, v in fields.items()},
                         "source": {"file": path.name, "sheet": sheet.title, "row": row_number}}
                by_code[key].append(entry)
                by_model[_text(fields["Device Model"]).casefold()].append(entry)
    finally:
        workbook.close()
    return dict(by_model), dict(by_code)


def _templates(directory: Path, models: dict, parents: dict, *, parent_reference: Path) -> tuple[list[dict], list[str], list[dict]]:
    occurrences, columns, files = [], [], []
    for path in sorted(directory.glob("*.xlsx")):
        if path.name.startswith("~$") or path.resolve() == parent_reference.resolve():
            continue
        workbook = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        files.append({"file": path.name, "sha256": _hash(path)})
        try:
            for sheet in workbook:
                rows = list(sheet.iter_rows())
                header_index = next((i for i, row in enumerate(rows[:15])
                                     if any(_text(c.value).startswith("Issuing Entity") for c in row)), None)
                if header_index is None:
                    raise ValueError(f"Cannot identify template headers: {path.name}:{sheet.title}")
                headers = [c.value for c in rows[header_index]]
                keys = _header_keys(headers)
                udi_index = next((i for i, h in enumerate(headers) if _text(h).startswith("UDI-DI code")), None)
                issuer_index = next(i for i, h in enumerate(headers) if _text(h).startswith("Issuing Entity"))
                basic_index = next((i for i, h in enumerate(headers) if _text(h) == "Basic UDI-DI"), None)
                if udi_index is None:
                    raise ValueError(f"Missing UDI-DI column: {path.name}:{sheet.title}")
                cached_rows = list(cached[sheet.title].iter_rows(values_only=True))
                for number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
                    if not any(c.value is not None and _text(c.value) for c in row):
                        continue
                    values = [c.value for c in row]
                    formulas = {keys[i]: c.value for i, c in enumerate(row) if c.data_type == "f"}
                    effective = [cached_rows[number - 1][i] if c.data_type == "f" else c.value
                                 for i, c in enumerate(row)]
                    fields = {k: _serial(v) for k, v in zip(keys, effective)}
                    for i, key in enumerate(keys):
                        if (headers[i] is not None or values[i] is not None) and key not in columns:
                            columns.append(key)
                    identity = _identity(effective[issuer_index], effective[udi_index])
                    reasons = []
                    if not identity[0] or not identity[1]:
                        reasons.append("Missing device identity")
                    elif identity[0] == "GS1" and not re.fullmatch(r"\d{14}", identity[1]):
                        reasons.append("GS1 UDI-DI must contain 14 digits; no padding was inferred")
                    if any(effective[i] is None for i, c in enumerate(row) if c.data_type == "f"):
                        reasons.append("Formula has no cached value")
                    kind = _device_type(path, sheet.title)
                    direct = _text(effective[basic_index]) if basic_index is not None else ""
                    references = parents.get((identity[0], direct), []) if direct else models.get(kind.casefold(), [])
                    references = [r for r in references if r["key"][0] == identity[0]]
                    choices = {tuple(r["key"]) for r in references}
                    parent = (identity[0], direct) if direct else next(iter(choices)) if len(choices) == 1 else None
                    if not direct and len(choices) > 1:
                        reasons.append("Ambiguous parent reference for device type")
                    if parent is None:
                        reasons.append("Missing Basic UDI-DI parent reference")
                    elif not references:
                        reasons.append("Parent properties absent from BasicUDI reference")
                    occurrences.append({"key": identity, "device_type": kind, "parent": parent,
                                        "fields": fields, "reasons": reasons, "parent_references": references,
                                        "source": {"file": path.name, "sheet": sheet.title, "row": number,
                                                   "header_row": header_index + 1},
                                        "original_headers": headers, "original_values": [_serial(v) for v in values],
                                        "cell_types": [c.data_type for c in row], "formulas": formulas})
        finally:
            workbook.close()
            cached.close()
    if not files:
        raise ValueError("No template workbooks found")
    return occurrences, columns, files


def _find(node: etree._Element, path: str) -> str:
    item = node.find(path, NS)
    return _text(item.text) if item is not None else ""


def _xml_value(node: etree._Element) -> Any:
    """Retain namespace-qualified structure, attributes and repeated children."""
    return {"tag": node.tag, "attributes": dict(node.attrib), "text": node.text,
            "children": [_xml_value(c) for c in node if isinstance(c.tag, str)]}


def _export_fields(node: etree._Element, label: str) -> dict:
    grouped: dict[str, list] = defaultdict(list)
    for child in node:
        if not isinstance(child.tag, str):
            continue
        name = etree.QName(child).localname
        value = _text(child.text) if len(child) == 0 and not child.attrib else _xml_value(child)
        grouped[f"Accepted {label}: {name}"].append(value)
    return {k: v[0] if len(v) == 1 else v for k, v in grouped.items()}


def _exports(directory: Path, overrides: dict[str, str]) -> tuple[list[dict], list[dict]]:
    entries, files = [], []
    paths = sorted(directory.glob("*.xml"))
    unused = set(overrides) - {p.name for p in paths}
    if unused:
        raise ValueError(f"Encoding overrides name absent files: {sorted(unused)}")
    for path in paths:
        encoding = overrides.get(path.name)
        parser = etree.XMLParser(encoding=encoding, resolve_entities=False, no_network=True)
        tree = etree.parse(str(path), parser)
        if tree.docinfo.doctype:
            raise ValueError(f"DOCTYPE is not supported: {path.name}")
        root = tree.getroot()
        if root.tag != f"{{{NS['message']}}}PullResponse":
            raise ValueError(f"Unsupported export envelope: {path.name}")
        devices = root.findall("message:payload/device:Device", NS)
        meta = {"file": path.name, "sha256": _hash(path), "encoding_override": encoding,
                "created_at": _find(root, "message:creationDateTime"),
                "page": _find(root, "message:pageNumber"),
                "number_of_pages": _find(root, "message:numberOfPages"),
                "page_size": _find(root, "message:pageSize"), "devices": len(devices)}
        files.append(meta)
        for index, device in enumerate(devices, 1):
            basic = device.find("device:MDRBasicUDI", NS)
            udi = device.find("device:MDRUDIDIData", NS)
            if basic is None or udi is None:
                raise ValueError(f"Unsupported or incomplete MDR device: {path.name}:{index}")
            key = _identity(_find(udi, "udi:identifier/common:issuingEntityCode"),
                            _find(udi, "udi:identifier/common:DICode"))
            parent = _identity(_find(basic, "basic:identifier/common:issuingEntityCode"),
                               _find(basic, "basic:identifier/common:DICode"))
            link = _identity(_find(udi, "udi:basicUDIIdentifier/common:issuingEntityCode"),
                             _find(udi, "udi:basicUDIIdentifier/common:DICode"))
            if not all(key) or not all(parent) or link != parent:
                raise ValueError(f"Missing or conflicting exported identity: {path.name}:{index}")
            if _find(udi, "entity:state") != "REGISTERED":
                raise ValueError(f"Exported device is not REGISTERED: {path.name}:{index}")
            market = udi.find("udi:marketInfos", NS)
            reasons = []
            for label, node in (("Parent", basic), ("Device", udi), ("Market Info", market)):
                if node is None or not _find(node, "entity:version") or not _find(node, "entity:state"):
                    reasons.append(f"Missing exported {label} state/version")
            fields = {"Accepted Basic UDI-DI": parent[1], "Accepted Parent Issuer": parent[0]}
            for label, node in (("Parent", basic), ("Device", udi), ("Market Info", market)):
                fields[label + " State"] = _find(node, "entity:state") if node is not None else ""
                fields[label + " Version"] = _find(node, "entity:version") if node is not None else ""
                fields[label + " Version Date"] = _find(node, "entity:versionDate") if node is not None else ""
            fields.update(_export_fields(basic, "Parent"))
            fields.update(_export_fields(udi, "Device"))
            entries.append({"key": key, "parent": parent,
                            "device_type": _find(basic, "basic:modelName/common:name"),
                            "fields": fields, "reasons": reasons, "source": {**meta, "device_index": index},
                            "xml": _xml_value(device)})
    if not files:
        raise ValueError("No exported XML files found")
    populated = [f for f in files if f["devices"]]
    declared = {int(f["number_of_pages"]) for f in populated}
    if len(declared) != 1 or Counter(int(f["page"]) for f in populated) != Counter(range(next(iter(declared)))):
        raise ValueError("Missing, duplicate or inconsistent populated export pages")
    return entries, files


def _cell(value: Any) -> Any:
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, str) and len(value) > 32767:
        raise ValueError("Source value exceeds Excel cell capacity; cannot preserve it in a workbook cell")
    if isinstance(value, str) and re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise ValueError("Source value contains unsupported Excel control characters")
    return value


def prepare_workbook(*, template_dir: Path, xml_dir: Path, parent_reference: Path,
                     output_dir: Path, encoding_overrides: dict[str, str] | None = None,
                     parent_market_overrides: dict[tuple[str, str], list[str]] | None = None,
                     previous_audit: Path | None = None) -> dict:
    """Write a dated review workbook and audit JSON; never change source files or SQLite."""
    models, parents = _parent_reference(parent_reference)
    templates, template_columns, template_files = _templates(template_dir, models, parents, parent_reference=parent_reference)
    exports, export_files = _exports(xml_dir, encoding_overrides or {})
    market_overrides = parent_market_overrides or {}
    for parent_key, countries in market_overrides.items():
        if parent_key not in parents or not countries or any(not _text(c) or "?" in c for c in countries):
            raise ValueError(f"Invalid parent market override: {parent_key}")
    previous = json.loads(previous_audit.read_text(encoding="utf-8")) if previous_audit else None
    proposed: dict[tuple, list[dict]] = defaultdict(list)
    accepted: dict[tuple, list[dict]] = defaultdict(list)
    for index, entry in enumerate(templates):
        key = tuple(entry["key"])
        proposed[key if all(key) else ("missing", str(index))].append(entry)
    for entry in exports:
        accepted[tuple(entry["key"])].append(entry)
    if any(len(v) != 1 for v in accepted.values()):
        raise ValueError("Duplicate exported device identities require review before preparation")
    stamp = datetime.now(UTC)
    extracted = stamp.isoformat()
    columns = CONTROLS + ["Proposed: " + k for k in template_columns]
    parent_columns = list(dict.fromkeys(k for entries in parents.values() for r in entries for k in r["fields"]))
    columns += ["Proposed Parent: " + k for k in parent_columns]
    accepted_columns = list(dict.fromkeys(ACCEPTED + [k for e in exports for k in e["fields"]]))
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    sheets: dict[str, Worksheet] = {
        name: workbook.create_sheet(name) for name in ("To Register", "Registered")
    }
    row_counts: dict[str, int] = {name: 1 for name in sheets}
    sheet_columns: dict[str, list[str]] = {
        "To Register": columns, "Registered": columns + accepted_columns
    }
    for name, sheet in sheets.items():
        sheet.append(sheet_columns[name])
        sheet.freeze_panes = "G2"
        sheet.row_dimensions[1].height = 75
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="244062")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, column in enumerate(sheet_columns[name], 1):
            sheet.column_dimensions[get_column_letter(i)].width = 45 if column in ("Review Reasons", "Template Sources") else 24
    rows, issues = [], []
    for key in sorted(set(proposed) | set(accepted)):
        sources, registered = proposed.get(key, []), accepted.get(key, [])
        reasons = list(dict.fromkeys(reason for s in sources + registered for reason in s["reasons"]))
        row: dict[str, Any] = {"Issuing Entity": (sources or registered)[0]["key"][0],
                               "UDI-DI": (sources or registered)[0]["key"][1],
                               "Device Type": "; ".join(dict.fromkeys(s["device_type"] for s in sources + registered)),
                               "Template Sources": [s["source"] for s in sources],
                               "Source Occurrence Count": len(sources)}
        for field in template_columns:
            values = [s["fields"][field] for s in sources if field in s["fields"] and s["fields"][field] is not None]
            distinct = list(dict.fromkeys(json.dumps(v, sort_keys=True, default=str) for v in values))
            if len(distinct) > 1:
                reasons.append("Conflicting template values: " + field)
            elif values:
                row["Proposed: " + field] = values[0]
        parent_keys = {tuple(s["parent"]) for s in sources if s["parent"]}
        if len(parent_keys) > 1:
            reasons.append("Conflicting proposed Basic UDI-DI parents")
        elif parent_keys:
            parent_key = next(iter(parent_keys))
            row.update({"Proposed Parent Issuer": parent_key[0], "Proposed Basic UDI-DI": parent_key[1]})
        refs = [r for s in sources for r in s["parent_references"]]
        row["Parent Reference Sources"] = list({json.dumps(r["source"], sort_keys=True): r["source"] for r in refs}.values())
        for field in parent_columns:
            values = [r["fields"].get(field) for r in refs if r["fields"].get(field) is not None]
            distinct = {json.dumps(v, sort_keys=True, default=str) for v in values}
            if len(distinct) > 1:
                reasons.append("Conflicting parent reference values: " + field)
            elif values:
                row["Proposed Parent: " + field] = values[0]
        if len(parent_keys) == 1 and (parent_key := next(iter(parent_keys))) in market_overrides:
            countries = market_overrides[parent_key]
            row["Proposed Parent: Member State of the placing on the EU market of the Device:"] = countries[0]
            row["Proposed Parent: Member States where device is or is to be made available on the market:"] = "; ".join(countries)
        for field in parent_columns:
            value = _text(row.get("Proposed Parent: " + field))
            if field.startswith("URL for additional information") and value.casefold() == "not available yet":
                reasons.append("Parent information URL unavailable; confirm optional-field handling")
            if field.startswith("Member States where device") and refs and (not value or "?" in value):
                reasons.append("Proposed parent market countries missing or tentative")
        if registered:
            export = registered[0]
            row.update(export["fields"])
            row.update({"Export Sources": [export["source"]], "Export Created At": export["source"]["created_at"],
                        "Encoding Overrides": export["source"]["encoding_override"], "Extracted At": extracted})
            if parent_keys and any(p != tuple(export["parent"]) for p in parent_keys):
                reasons.append("Proposed parent differs from exported accepted parent")
        reasons = list(dict.fromkeys(reasons))
        conflicting = any(r.startswith("Conflicting") for r in reasons)
        invalid_identity = any(r.startswith(("Missing device identity", "GS1 UDI-DI")) for r in reasons)
        status = "Registered" if registered else "Needs review" if conflicting or invalid_identity else "Not registered"
        row.update({"Registration Status": status, "Review Required": "Yes" if reasons else "No",
                    "Review Reasons": "; ".join(reasons)})
        tab = "Registered" if registered else "To Register"
        sheet = sheets[tab]
        sheet.append([_cell(row.get(column)) for column in sheet_columns[tab]])
        row_counts[tab] += 1
        row_number = row_counts[tab]
        # All strings must be literal text: source content beginning '=' is not a formula.
        for column_number in range(1, len(sheet_columns[tab]) + 1):
            c = sheet.cell(row_number, column_number)
            if isinstance(c.value, str):
                c.data_type = "s"
                c.number_format = "@"
            c.alignment = Alignment(vertical="top")
        colour = {"Registered": "E2EFDA", "Not registered": "FFF2CC", "Needs review": "F8CBAD"}[status]
        sheet.cell(row_number, 2).fill = PatternFill("solid", fgColor=colour)
        if reasons:
            for column_number in (3, 4):
                sheet.cell(row_number, column_number).fill = PatternFill("solid", fgColor="F8CBAD")
            issues.append({"key": row["Issuing Entity"] + ":" + row["UDI-DI"], "tab": tab,
                           "row": row_number, "device_type": row["Device Type"],
                           "catalogue": next((_text(v) for k, v in row.items() if k.startswith("Proposed: Reference/ Catalogue")), "")
                           or _text(row.get("Accepted Device: referenceNumber")),
                           "parent": row.get("Proposed Basic UDI-DI") or row.get("Accepted Basic UDI-DI"),
                           "reasons": reasons})
        rows.append({"key": key, "tab": tab, "status": status, "review_reasons": reasons})
    for sheet in sheets.values():
        sheet.auto_filter.ref = sheet.dimensions
    summary = {"template_workbooks": len(template_files), "template_rows": len(templates),
               "template_unique_identities": len(proposed), "export_files": len(export_files),
               "registered_devices": len(accepted), "matched_identities": len(set(proposed) & set(accepted)),
               "to_register_tab_rows": row_counts["To Register"] - 1,
               "registered_tab_rows": row_counts["Registered"] - 1,
               "rows_requiring_review": len(issues), "status_counts": dict(Counter(r["status"] for r in rows))}
    summary_entries = []
    def summary_entry(category: str, scope: str, count: int, details: str, action: str,
                      *, catalogue: str = "", identifier: str = "", parent: str = "", location: str = "") -> None:
        summary_entries.append({"Category": category, "Scope": scope, "Affected Rows": count,
                                "Catalogue": catalogue, "UDI-DI / Issuer": identifier,
                                "Basic UDI-DI": parent, "Data Location": location,
                                "Details": details, "Action / Status": action})
    for name, count in (("Source template rows", len(templates)), ("Distinct template identities", len(proposed)),
                        ("To Register", summary["to_register_tab_rows"]), ("Registered", len(accepted)),
                        ("Distinct output identities", len(rows)), ("Rows with review flags", len(issues))):
        summary_entry("Counts", name, count, "", "Informational")
    for parent_key, countries in market_overrides.items():
        count = sum(tuple(s["parent"] or ()) == parent_key for s in templates)
        summary_entry("Approved setting", parents[parent_key][0]["fields"]["Device Model"], count,
                      f"Markets: {'; '.join(countries)}. First placement: {countries[0]}.",
                      "Applied to proposed output; source workbook and accepted XML unchanged", parent=parent_key[1])
    grouped_optional: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for issue in issues:
        optional = [r for r in issue["reasons"] if r.startswith("Parent information URL")]
        if optional:
            grouped_optional[(issue["device_type"], issue["parent"])].append(issue)
        other = [r for r in issue["reasons"] if r not in optional]
        if other:
            details = "; ".join(other)
            issuer, code = issue["key"].split(":", 1)
            identity = (issuer, code)
            if "Proposed parent differs from exported accepted parent" in other:
                accepted_parent = accepted[identity][0]["parent"][1]
                details += f". Proposed parent: {issue['parent']}; accepted parent: {accepted_parent}."
            if "Conflicting proposed Basic UDI-DI parents" in other:
                options = sorted({s["parent"][1] for s in proposed[identity] if s["parent"]})
                emdn = sorted({_text(v) for s in proposed[identity] for k, v in s["fields"].items()
                               if k.startswith("Enter a nomenclature code") and v is not None})
                details += f". Parent alternatives: {'; '.join(options)}. EMDN alternatives: {'; '.join(emdn)}."
            summary_entry("Device exception", issue["device_type"], 1, details,
                          "QA/source review required; preserve existing acceptance", catalogue=issue["catalogue"],
                          identifier=issue["key"], parent=issue["parent"] or "",
                          location=f"{issue['tab']} row {issue['row']}")
    for (kind, parent), group in sorted(grouped_optional.items()):
        summary_entry("Optional-field exception", kind, len(group),
                      'Parent information URL contains "Not available yet"; this is not a URL.',
                      "Supply URL or agree omission of the optional field; not a registration-status failure", parent=parent)
    removed = []
    if previous:
        previous_groups: dict[tuple[str, str], list[dict]] = defaultdict(list)
        for occurrence in previous["template_occurrences"]:
            previous_groups[tuple(occurrence["key"])].append(occurrence)
        for old_key in sorted(set(previous_groups) - set(proposed)):
            sources = previous_groups[old_key]
            catalogue = next((_text(v) for s in sources for k, v in s["fields"].items()
                              if k.startswith("Reference/ Catalogue")), "")
            removed.append({"key": old_key, "catalogue": catalogue, "sources": [s["source"] for s in sources]})
            summary_entry("Removed from current templates", sources[0]["device_type"], 1,
                          "Present in previous preparation inputs; absent from the current templates.",
                          "Confirm intentional exclusion; absence does not delete accepted database state",
                          catalogue=catalogue, identifier=":".join(old_key),
                          parent=sources[0]["parent"][1] if sources[0]["parent"] else "")
    for meta in export_files:
        if meta["encoding_override"]:
            summary_entry("Recorded encoding override", meta["file"], meta["devices"], meta["encoding_override"],
                          "Applied explicitly; original XML retained")
    summary_entry("Later verification", "Production baseline and import", 0,
                  "Exports are dated snapshots. Parent-only registrations, export currency, actor configuration, "
                  "canonical mapping coverage and XML readiness need verification before database import/submission.",
                  "Deferred; this workbook does not initialize SQLite or certify XML readiness")
    summary_sheet: Worksheet = workbook.create_sheet("Summary")
    summary_headers = ["Category", "Scope", "Affected Rows", "Catalogue", "UDI-DI / Issuer", "Basic UDI-DI",
                       "Data Location", "Details", "Action / Status"]
    summary_sheet.append(summary_headers)
    for entry in summary_entries:
        summary_sheet.append([_cell(entry.get(h)) for h in summary_headers])
    summary_sheet.freeze_panes = "A2"
    summary_sheet.auto_filter.ref = summary_sheet.dimensions
    for index, header in enumerate(summary_headers, 1):
        summary_sheet.column_dimensions[get_column_letter(index)].width = 70 if header in ("Details", "Action / Status") else 30
    for index in range(1, len(summary_entries) + 2):
        for column in range(1, len(summary_headers) + 1):
            # This worksheet has no merged cells; each coordinate is a writable Cell.
            cell = cast(Cell, summary_sheet.cell(row=index, column=column))
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if index == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="244062")
            elif summary_entries[index - 2]["Category"] in ("Device exception", "Optional-field exception", "Removed from current templates"):
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
        summary_sheet.row_dimensions[index].height = 60 if index > 1 else 30
    summary["removed_template_identities"] = len(removed)
    audit = {"format_version": 1, "prepared_at": extracted, "purpose": "Workbook preparation; not database import",
             "inputs": {"templates": template_files, "exports": export_files,
                        "parent_reference": {"file": parent_reference.name, "sha256": _hash(parent_reference)}},
             "summary": summary, "issues": issues, "template_occurrences": templates,
             "export_occurrences": exports, "output_rows": rows,
             "parent_reference_entries": [r for entries in parents.values() for r in entries],
             "approved_parent_market_overrides": [{"parent": k, "countries": v, "first_placement": v[0]}
                                                  for k, v in market_overrides.items()],
             "summary_entries": summary_entries, "removed_template_identities": removed,
             "previous_audit": {"file": previous_audit.name, "sha256": _hash(previous_audit)} if previous_audit else None}
    output_dir.mkdir(parents=True, exist_ok=True)
    name = "production-import-" + stamp.strftime("%Y%m%d-%H%M%S-%f")
    destination, audit_path = output_dir / (name + ".xlsx"), output_dir / (name + ".audit.json")
    # Exclusive creation preserves any earlier output, including concurrent runs.
    created = []
    try:
        with destination.open("xb") as handle:
            created.append(destination)
            workbook.save(handle)
        audit["workbook_sha256"] = _hash(destination)
        with audit_path.open("x", encoding="utf-8") as handle:
            created.append(audit_path)
            json.dump(audit, handle, ensure_ascii=False, indent=2, default=str)
    except Exception:
        for path in created:
            path.unlink(missing_ok=True)
        raise
    finally:
        workbook.close()
    return {"workbook": str(destination.resolve()), "audit": str(audit_path.resolve()), **summary}
