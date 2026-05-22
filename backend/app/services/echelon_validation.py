from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.config import get_settings
from app.services.canonical_review import CanonicalReviewService
from app.validation_models import (
    BlockerSummary,
    CompletenessSnapshot,
    EchelonValidationBundle,
    EchelonValidationRecord,
    ExcludedSheetSummary,
    MatchStatus,
    SheetValidationSummary,
    SourceFieldCoverageEntry,
    SourceFieldCoverageStatus,
    SourceFieldCoverageSummary,
    ValidationFieldValue,
    ValueSourceType,
)

SOURCE_WORKBOOK_NAME = "Template for Echelon family EUDAMED.xlsx"
REFERENCE_WORKBOOK_NAME = "uat-eudamed_mdr_products_tracekey_sample_data.xlsx"
HEADER_SENTINEL = "UDI-DI code"

ISSUING_ENTITY_HEADER = "Issuing Entity e.g. GS1"
UDI_DI_HEADER = "UDI-DI code e.g. taken from 2nd page of DoC (Note: needs to be 14 digits long, add zero to front of code)"
SECONDARY_UDI_HEADER = "UDI-DI from another entity (secondary) applicable e.g. No"
EMDN_HEADER = "Enter a nomenclature code (EMDN code) e.g starts with Y06xxxx then click Find"
TRADE_NAME_HEADER = "Trade Name applicable e.g. YES (use  Decription detail from second page of DoC"
LANGUAGE_HEADER = "Select the language e.g English"
CATALOGUE_HEADERS = (
    "Reference/ Catalogue number e.g . Taken from Product code in second page of DoC",
    "Reference/ Catalogue number e.g . Taken from Product code on second page of DoC",
)
QUANTITY_HEADER = "Quantity of device e.g. 1"
UDI_PI_HEADER = "Type of UDI-PI e.g. select Serial number and Manufacturing Date"
STATUS_HEADER = "UDI-DI status e.g. On the EU market"
STERILE_HEADER = "Device labelled as sterile e.g. NO"
LATEX_HEADER = "Containing latex e.g. NO"
DIRECT_MARKING_HEADER = "Is the device directly marked? E.g. NO"
SINGLE_USE_HEADER = "Labelled as single use e.g.NO"
MAX_REUSES_APPLICABLE_HEADER = "Maximum number of reuses applicable e.g. NO"
STERILIZATION_HEADER = "Need for sterilsation before use e.g. NO"
REPROCESSED_HEADER = "Reprocessed single use device e.g. NO"
ANNEX_XVI_HEADER = "Intended purpose other than medical (Annex XVI) e.g. NO"
DESIGNED_BY_ANOTHER_HEADER = "Is the device designed and manufactured by another legal or natural person? E.g. NO"
CLINICAL_INVESTIGATION_HEADER = "Clinical Investigation e.g. NO"
HUMAN_TISSUE_HEADER = "Presence of human tissues or cells, or their derivatives e.g. NO"
ANIMAL_TISSUE_HEADER = "Presence of animal tissues or cells, or their derivatives e.g. NO"
MEDICINAL_SUBSTANCE_HEADER = "Presence of a substance which, if used separately, may be considered to be a medicinal product e.g. NO"
HUMAN_BLOOD_SUBSTANCE_HEADER = "Presence of a substance which, if used separately, may be considered to be a medicinal product derived from human blood or human plasma e.g. NO"
FIRST_EU_MARKET_HEADER = "Member state where first placed on the EU market e.g. Germany"

PRODUCT_TEMPLATE_HEADER = "Product Template"
PRODUCT_CODE_VALUE_HEADER = "Product Code (Value)"
PRODUCT_NAME_HEADER = "Name"
MATERIAL_NUMBER_HEADER = "Material Number"
MANUFACTURER_CODE_HEADER = "Manufacturer Code"
RISK_CLASS_HEADER = "Risk Class"
MODEL_TYPE_HEADER = "Model Type"
MODEL_NAME_HEADER = "Model Name"
MODEL_HEADER = "Model"
ANIMAL_TISSUES_HEADER = "Animal Tissues Cells"
AR_ACTOR_CODE_HEADER = "Authorised Representative Actor Code"
HUMAN_TISSUES_HEADER = "Human Tissues Cells"
HUMAN_PRODUCT_HEADER = "Human Product"
MEDICINAL_PRODUCT_HEADER = "Medicinal Product"
SPECIAL_DEVICE_HEADER = "Special Device"
TYPE_HEADER = "Type"
ACTIVE_HEADER = "Active"
ADMINISTERING_MEDICINE_HEADER = "Administers or/and Removes Medicine"
IMPLANTABLE_HEADER = "Implantable"
MEASURING_FUNCTION_HEADER = "Measuring Function"
REUSABLE_HEADER = "Reusable"
RISK_CLASS_IIB_IMPLANTABLE_HEADER = "Risk Class II B Implantable"
DEVICE_STATUS_HEADER = "Device Status"
EMDN_CODES_HEADER = "EMDN Codes"
PRODUCTION_IDENTIFIER_HEADER = "Production Identifier"
SECONDARY_IDENTIFIER_CODE_HEADER = "Secondary Identifier - Code"
SECONDARY_IDENTIFIER_ENTITY_HEADER = "Secondary Identifier - Issuing Entity"
NUMBER_OF_REUSES_HEADER = "Number Of Reuses"
BASE_QUANTITY_HEADER = "Base Quantity"
LATEX_REFERENCE_HEADER = "Latex"
REPROCESSED_REFERENCE_HEADER = "Reprocessed"

PARTIAL_SOURCE_HEADERS = {
    "storage/handling conditions, if applicable e.g. yes": (
        "Storage conditions are recognized, but the repeating list structure is still only partially modeled."
    ),
    "storage /handling conditions type e.g. lower limit of temp": (
        "Storage conditions are recognized, but the repeating list structure is still only partially modeled."
    ),
    "description e.g. taken from ifu technical data page storage temp range e.g. -15c": (
        "Storage condition descriptions are recognized, but the repeating list structure is still only partially modeled."
    ),
    "add another storage/handling condition e.g. upper limit of temp": (
        "Storage conditions are recognized, but the repeating list structure is still only partially modeled."
    ),
    "description e.g. taken from ifu technical data page storage temp range e.g. +50c": (
        "Storage condition descriptions are recognized, but the repeating list structure is still only partially modeled."
    ),
    "critical warnings or contra-indications, if applicable e.g. yes": (
        "Critical warnings are recognized, but the repeating warning structure is still only partially modeled."
    ),
    "critical warning type e.g. consult instructions for use": (
        "Critical warnings are recognized, but the repeating warning structure is still only partially modeled."
    ),
}

REPRESENTED_SCOPE_NOTES = {
    "clinical size applicable e.g. no": (
        "Represented for current Echelon scope because all reviewed rows indicate clinical size is not applicable."
    ),
    "labelled for presence of carcinogenic, mutagenic and toxic to reproduction (cmr) substances of category 1a or 1b e.g. no": (
        "Represented for current Echelon scope because all reviewed rows indicate CMR substance presence is not applicable."
    ),
    "labelled for presence of substances with endocrine-disrupting properties e.g. no": (
        "Represented for current Echelon scope because all reviewed rows indicate endocrine-disrupting substance presence is not applicable."
    ),
}


@dataclass(frozen=True)
class SourceHeaderProfile:
    display_name: str
    normalized_name: str
    sheets: tuple[str, ...]


@dataclass(frozen=True)
class SourceRow:
    sheet_name: str
    row_index: int
    values: dict[str, object | None]


class EchelonValidationService:
    def __init__(self) -> None:
        self.settings = get_settings()

    def build_validation_bundle(self) -> EchelonValidationBundle:
        source_workbook = self.settings.excel_dir / SOURCE_WORKBOOK_NAME
        reference_workbook = self.settings.basic_udi_reference_dir / REFERENCE_WORKBOOK_NAME

        source_rows = self._load_source_rows(source_workbook)
        source_headers = self._load_source_headers(source_workbook)
        source_field_coverage = self._build_source_field_coverage(source_headers)
        family_reference_context = self._load_family_reference_context(reference_workbook)

        included_records: list[EchelonValidationRecord] = []
        excluded_by_sheet: dict[str, int] = {}

        for row in source_rows:
            if not family_reference_context:
                excluded_by_sheet[row.sheet_name] = excluded_by_sheet.get(row.sheet_name, 0) + 1
                continue
            included_records.append(self._build_record(source_workbook.name, row, family_reference_context))

        excluded_records = sum(excluded_by_sheet.values())
        tracked_required_fields = (
            sum(1 for field in included_records[0].fields if field.required) if included_records else 0
        )

        return EchelonValidationBundle(
            family_scope="Echelon only",
            scope_note=(
                "All rows from the Echelon family workbook inherit the same Echelon Basic UDI-DI context. "
                "The validation subset therefore covers the full Echelon workbook population when that family-level "
                "Basic UDI reference is available."
            ),
            validation_note=(
                "Completeness is measured against the tracked canonical fields currently modeled for this "
                "preview. XML projection still remains a downstream step."
            ),
            source_workbook=source_workbook.name,
            total_source_records=len(source_rows),
            validation_subset_records=len(included_records),
            excluded_records=excluded_records,
            matched_reference_records=len(included_records),
            tracked_required_fields=tracked_required_fields,
            before_complete_records=sum(
                1 for record in included_records if record.before_completeness.status == "complete"
            ),
            after_complete_records=sum(
                1 for record in included_records if record.after_completeness.status == "complete"
            ),
            blocker_summaries=self._build_blocker_summaries(included_records),
            sheet_summaries=self._build_sheet_summaries(included_records),
            source_field_total=len(source_headers),
            source_field_coverage_summaries=self._build_source_field_coverage_summaries(source_field_coverage),
            source_field_coverage=source_field_coverage,
            sample_records=self._build_sample_records(included_records),
            excluded_sheet_summaries=[
                ExcludedSheetSummary(
                    sheet_name=sheet_name,
                    record_count=record_count,
                    reason="No family-level Basic UDI-DI reference coverage is currently available for Echelon.",
                )
                for sheet_name, record_count in sorted(excluded_by_sheet.items())
            ],
            records=included_records,
        )

    def _build_record(
        self,
        workbook_name: str,
        row: SourceRow,
        reference_match: dict[str, dict[str, object | None]],
    ) -> EchelonValidationRecord:
        source_values = row.values
        basic_product = reference_match["basic_product"]
        basic_details = reference_match["basic_details"]
        match_status: MatchStatus = "matched"
        issuing_entity = self._issuing_entity_code(source_values.get(ISSUING_ENTITY_HEADER))
        primary_udi_di = self._string_value(source_values.get(UDI_DI_HEADER))
        basic_identifier = self._di_identifier(
            issuing_entity=issuing_entity,
            di_code=self._string_value(basic_product.get(PRODUCT_CODE_VALUE_HEADER)),
        )
        udi_identifier = self._di_identifier(
            issuing_entity=issuing_entity,
            di_code=primary_udi_di,
        )
        secondary_identifier = self._di_identifier(
            issuing_entity=self._issuing_entity_code(source_values.get(SECONDARY_IDENTIFIER_ENTITY_HEADER)),
            di_code=self._string_value(source_values.get(SECONDARY_IDENTIFIER_CODE_HEADER)),
        )
        number_of_reuses = self._number_of_reuses(source_values)

        fields = [
            self._field(
                canonical_path="basic_device.risk_class",
                business_label="Basic Risk Class",
                before_value=None,
                after_value=self._string_value(basic_details.get(RISK_CLASS_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Risk Class",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.model_type",
                business_label="Basic Model Type",
                before_value=None,
                after_value=self._string_value(basic_details.get(MODEL_TYPE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Model Type",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.model_name",
                business_label="Basic Model Name",
                before_value=None,
                after_value=self._string_value(basic_details.get(MODEL_NAME_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Model Name",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.model",
                business_label="Basic Model",
                before_value=None,
                after_value=self._string_value(basic_details.get(MODEL_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Model",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.identifier",
                business_label="Basic Identifier",
                before_value=None,
                after_value=basic_identifier,
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="Products!Product Code (Value) + workbook issuing entity",
                update_reason="Composed as a DI identifier using the shared Basic UDI code and issuing entity.",
            ),
            self._field(
                canonical_path="manufacturer.issuing_entity",
                business_label="Issuing Entity",
                before_value=issuing_entity,
                after_value=issuing_entity,
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{ISSUING_ENTITY_HEADER}",
            ),
            self._field(
                canonical_path="manufacturer.manufacturer_srn",
                business_label="Manufacturer SRN",
                before_value=None,
                after_value=self._string_value(basic_details.get(MANUFACTURER_CODE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Manufacturer Code",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.basic_udi_di",
                business_label="Basic UDI-DI",
                before_value=None,
                after_value=self._string_value(basic_product.get(PRODUCT_CODE_VALUE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="Products!Product Code (Value)",
                update_reason="Added from the matched Basic UDI product row.",
            ),
            self._field(
                canonical_path="basic_device.authorised_representative_srn",
                business_label="Authorised Representative SRN",
                before_value=None,
                after_value=self._string_value(basic_details.get(AR_ACTOR_CODE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Authorised Representative Actor Code",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.regulation",
                business_label="Regulation",
                before_value="MDR",
                after_value="MDR",
                before_source="derived",
                after_source="derived",
                source_detail="QMS scope decision",
            ),
            self._field(
                canonical_path="basic_device.nomenclature_code",
                business_label="EMDN Code",
                before_value=self._string_value(source_values.get(EMDN_HEADER)),
                after_value=self._string_value(source_values.get(EMDN_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{EMDN_HEADER}",
            ),
            self._field(
                canonical_path="basic_device.animal_tissues_cells",
                business_label="Basic Animal Tissues Cells",
                before_value=self._normalized_boolean(source_values.get(ANIMAL_TISSUE_HEADER)),
                after_value=self._normalized_boolean(basic_details.get(ANIMAL_TISSUES_HEADER)),
                before_source="workbook",
                after_source="basic_udi_reference",
                source_detail="Workbook animal tissue field / MDR-Basic!Animal Tissues Cells",
                update_reason="Aligned to the Basic UDI reference source for XML-facing basic-device content.",
            ),
            self._field(
                canonical_path="basic_device.human_tissues_cells",
                business_label="Basic Human Tissues Cells",
                before_value=self._normalized_boolean(source_values.get(HUMAN_TISSUE_HEADER)),
                after_value=self._normalized_boolean(basic_details.get(HUMAN_TISSUES_HEADER)),
                before_source="workbook",
                after_source="basic_udi_reference",
                source_detail="Workbook human tissue field / MDR-Basic!Human Tissues Cells",
                update_reason="Aligned to the Basic UDI reference source for XML-facing basic-device content.",
            ),
            self._field(
                canonical_path="basic_device.human_product_check",
                business_label="Human Product Check",
                before_value=self._normalized_boolean(source_values.get(HUMAN_BLOOD_SUBSTANCE_HEADER)),
                after_value=self._normalized_boolean(basic_details.get(HUMAN_PRODUCT_HEADER)),
                before_source="workbook",
                after_source="basic_udi_reference",
                source_detail="Workbook human blood/plasma field / MDR-Basic!Human Product",
                update_reason="Aligned to the Basic UDI reference workbook for MDR basic-device output.",
            ),
            self._field(
                canonical_path="basic_device.medicinal_product_check",
                business_label="Medicinal Product Check",
                before_value=self._normalized_boolean(source_values.get(MEDICINAL_SUBSTANCE_HEADER)),
                after_value=self._normalized_boolean(basic_details.get(MEDICINAL_PRODUCT_HEADER)),
                before_source="workbook",
                after_source="basic_udi_reference",
                source_detail="Workbook medicinal substance field / MDR-Basic!Medicinal Product",
                update_reason="Aligned to the Basic UDI reference workbook for MDR basic-device output.",
            ),
            self._field(
                canonical_path="basic_device.special_device",
                business_label="Special Device",
                before_value=None,
                after_value=self._string_value(basic_details.get(SPECIAL_DEVICE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Special Device",
                required=False,
            ),
            self._field(
                canonical_path="basic_device.type",
                business_label="Basic Device Type",
                before_value=None,
                after_value=self._string_value(basic_details.get(TYPE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Type",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.active",
                business_label="Active Device",
                before_value=None,
                after_value=self._normalized_boolean(basic_details.get(ACTIVE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Active",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.administering_medicine",
                business_label="Administering Medicine",
                before_value=None,
                after_value=self._normalized_boolean(basic_details.get(ADMINISTERING_MEDICINE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Administers or/and Removes Medicine",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.implantable",
                business_label="Implantable",
                before_value=None,
                after_value=self._normalized_boolean(basic_details.get(IMPLANTABLE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Implantable",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.measuring_function",
                business_label="Measuring Function",
                before_value=None,
                after_value=self._normalized_boolean(basic_details.get(MEASURING_FUNCTION_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Measuring Function",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.reusable",
                business_label="Reusable",
                before_value=None,
                after_value=self._normalized_boolean(basic_details.get(REUSABLE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Reusable",
                update_reason="Resolved from the Basic UDI reference workbook.",
            ),
            self._field(
                canonical_path="basic_device.iib_implantable_exception",
                business_label="IIb Implantable Exception",
                before_value=None,
                after_value=self._normalized_boolean(basic_details.get(RISK_CLASS_IIB_IMPLANTABLE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="MDR-Basic!Risk Class II B Implantable",
                update_reason="Resolved from the Basic UDI reference workbook.",
                required=False,
            ),
            self._field(
                canonical_path="basic_device.medicinal_product_substances",
                business_label="Medicinal Product Substances",
                before_value=self._normalized_boolean(source_values.get(MEDICINAL_SUBSTANCE_HEADER)),
                after_value=self._normalized_boolean(basic_details.get(MEDICINAL_PRODUCT_HEADER)),
                before_source="workbook",
                after_source="basic_udi_reference",
                source_detail="Workbook medicinal substance field / MDR-Basic!Medicinal Product",
                update_reason="Aligned to the Basic UDI reference workbook for MDR applicable properties.",
            ),
            self._field(
                canonical_path="basic_device.human_product_substances",
                business_label="Human Product Substances",
                before_value=self._normalized_boolean(source_values.get(HUMAN_BLOOD_SUBSTANCE_HEADER)),
                after_value=self._normalized_boolean(basic_details.get(HUMAN_PRODUCT_HEADER)),
                before_source="workbook",
                after_source="basic_udi_reference",
                source_detail="Workbook human blood/plasma field / MDR-Basic!Human Product",
                update_reason="Aligned to the Basic UDI reference workbook for MDR applicable properties.",
            ),
            self._field(
                canonical_path="device_record.basic_device_ref",
                business_label="Basic Device Reference",
                before_value=None,
                after_value=self._string_value(basic_product.get(PRODUCT_CODE_VALUE_HEADER)),
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="Products!Product Code (Value)",
                update_reason="Linked by inheriting the shared Echelon Basic UDI-DI family context.",
            ),
            self._field(
                canonical_path="device_record.identifier",
                business_label="UDI-DI Identifier",
                before_value=udi_identifier,
                after_value=udi_identifier,
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{ISSUING_ENTITY_HEADER} + {UDI_DI_HEADER}",
            ),
            self._field(
                canonical_path="device_record.status",
                business_label="UDI-DI Status",
                before_value=self._status_code(source_values.get(STATUS_HEADER)),
                after_value=self._status_code(source_values.get(STATUS_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{STATUS_HEADER}",
            ),
            self._field(
                canonical_path="device_record.primary_udi_di",
                business_label="Primary UDI-DI",
                before_value=primary_udi_di,
                after_value=primary_udi_di,
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{UDI_DI_HEADER}",
            ),
            self._field(
                canonical_path="device_record.catalogue_number",
                business_label="Catalogue Number",
                before_value=self._catalogue_number(source_values),
                after_value=self._catalogue_number(source_values),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!Reference/ Catalogue number",
            ),
            self._field(
                canonical_path="device_record.trade_name",
                business_label="Trade Name",
                before_value=self._string_value(source_values.get(TRADE_NAME_HEADER)),
                after_value=self._string_value(source_values.get(TRADE_NAME_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{TRADE_NAME_HEADER}",
            ),
            self._field(
                canonical_path="device_record.language",
                business_label="Language",
                before_value=self._string_value(source_values.get(LANGUAGE_HEADER)),
                after_value=self._string_value(source_values.get(LANGUAGE_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{LANGUAGE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.basic_udi_identifier",
                business_label="Basic UDI Identifier",
                before_value=None,
                after_value=basic_identifier,
                before_source="missing",
                after_source="basic_udi_reference",
                source_detail="Products!Product Code (Value) + workbook issuing entity",
                update_reason="Added from the shared Basic UDI-DI family context.",
            ),
            self._field(
                canonical_path="device_record.production_identifier",
                business_label="Production Identifier",
                before_value=self._production_identifier(source_values.get(UDI_PI_HEADER)),
                after_value=self._production_identifier(source_values.get(UDI_PI_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{UDI_PI_HEADER}",
            ),
            self._field(
                canonical_path="device_record.quantity",
                business_label="Quantity",
                before_value=self._string_value(source_values.get(QUANTITY_HEADER)),
                after_value=self._string_value(source_values.get(QUANTITY_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{QUANTITY_HEADER}",
            ),
            self._field(
                canonical_path="device_record.udi_pi_type",
                business_label="UDI-PI Type",
                before_value=self._string_value(source_values.get(UDI_PI_HEADER)),
                after_value=self._string_value(source_values.get(UDI_PI_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{UDI_PI_HEADER}",
            ),
            self._field(
                canonical_path="device_record.number_of_reuses",
                business_label="Number Of Reuses",
                before_value=number_of_reuses,
                after_value=number_of_reuses,
                before_source="derived",
                after_source="derived",
                source_detail=f"{row.sheet_name}!{MAX_REUSES_APPLICABLE_HEADER}",
                update_reason="Derived from the workbook applicability field for XML-facing UDI-DI data.",
            ),
            self._field(
                canonical_path="device_record.secondary_udi_di_applicable",
                business_label="Secondary UDI-DI Applicable",
                before_value=self._normalized_yes_no(source_values.get(SECONDARY_UDI_HEADER)),
                after_value=self._normalized_yes_no(source_values.get(SECONDARY_UDI_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{SECONDARY_UDI_HEADER}",
            ),
            self._field(
                canonical_path="device_record.secondary_identifier",
                business_label="Secondary Identifier",
                before_value=secondary_identifier,
                after_value=secondary_identifier,
                before_source="missing" if secondary_identifier is None else "workbook",
                after_source="missing" if secondary_identifier is None else "workbook",
                source_detail="MDR-UDI secondary identifier columns are not currently populated for Echelon.",
                required=False,
            ),
            self._field(
                canonical_path="device_record.sterile",
                business_label="Sterile",
                before_value=self._normalized_boolean(source_values.get(STERILE_HEADER)),
                after_value=self._normalized_boolean(source_values.get(STERILE_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{STERILE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.sterilization",
                business_label="Sterilization Before Use",
                before_value=self._normalized_boolean(source_values.get(STERILIZATION_HEADER)),
                after_value=self._normalized_boolean(source_values.get(STERILIZATION_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{STERILIZATION_HEADER}",
            ),
            self._field(
                canonical_path="device_record.contains_latex",
                business_label="Contains Latex",
                before_value=self._normalized_boolean(source_values.get(LATEX_HEADER)),
                after_value=self._normalized_boolean(source_values.get(LATEX_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{LATEX_HEADER}",
            ),
            self._field(
                canonical_path="device_record.reprocessed",
                business_label="Reprocessed",
                before_value=self._normalized_boolean(source_values.get(REPROCESSED_HEADER)),
                after_value=self._normalized_boolean(source_values.get(REPROCESSED_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{REPROCESSED_HEADER}",
            ),
            self._field(
                canonical_path="device_record.direct_marking",
                business_label="Direct Marking",
                before_value=self._normalized_boolean(source_values.get(DIRECT_MARKING_HEADER)),
                after_value=self._normalized_boolean(source_values.get(DIRECT_MARKING_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{DIRECT_MARKING_HEADER}",
            ),
            self._field(
                canonical_path="device_record.single_use",
                business_label="Single Use",
                before_value=self._normalized_boolean(source_values.get(SINGLE_USE_HEADER)),
                after_value=self._normalized_boolean(source_values.get(SINGLE_USE_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{SINGLE_USE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.annex_xvi_applicable",
                business_label="Annex XVI Applicable",
                before_value=self._normalized_boolean(source_values.get(ANNEX_XVI_HEADER)),
                after_value=self._normalized_boolean(source_values.get(ANNEX_XVI_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{ANNEX_XVI_HEADER}",
            ),
            self._field(
                canonical_path="device_record.designed_by_another_legal_entity",
                business_label="Designed By Another Legal Entity",
                before_value=self._normalized_boolean(source_values.get(DESIGNED_BY_ANOTHER_HEADER)),
                after_value=self._normalized_boolean(source_values.get(DESIGNED_BY_ANOTHER_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{DESIGNED_BY_ANOTHER_HEADER}",
            ),
            self._field(
                canonical_path="device_record.clinical_investigation",
                business_label="Clinical Investigation",
                before_value=self._normalized_boolean(source_values.get(CLINICAL_INVESTIGATION_HEADER)),
                after_value=self._normalized_boolean(source_values.get(CLINICAL_INVESTIGATION_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{CLINICAL_INVESTIGATION_HEADER}",
            ),
            self._field(
                canonical_path="device_record.human_tissues_present",
                business_label="Human Tissues Present",
                before_value=self._normalized_boolean(source_values.get(HUMAN_TISSUE_HEADER)),
                after_value=self._normalized_boolean(source_values.get(HUMAN_TISSUE_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{HUMAN_TISSUE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.animal_tissues_present",
                business_label="Animal Tissues Present",
                before_value=self._normalized_boolean(source_values.get(ANIMAL_TISSUE_HEADER)),
                after_value=self._normalized_boolean(source_values.get(ANIMAL_TISSUE_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{ANIMAL_TISSUE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.medicinal_substance_present",
                business_label="Medicinal Substance Present",
                before_value=self._normalized_boolean(source_values.get(MEDICINAL_SUBSTANCE_HEADER)),
                after_value=self._normalized_boolean(source_values.get(MEDICINAL_SUBSTANCE_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{MEDICINAL_SUBSTANCE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.human_blood_substance_present",
                business_label="Human Blood/Plasma Substance Present",
                before_value=self._normalized_boolean(source_values.get(HUMAN_BLOOD_SUBSTANCE_HEADER)),
                after_value=self._normalized_boolean(source_values.get(HUMAN_BLOOD_SUBSTANCE_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{HUMAN_BLOOD_SUBSTANCE_HEADER}",
            ),
            self._field(
                canonical_path="device_record.market_availability.market_status",
                business_label="Market Status",
                before_value=self._status_code(source_values.get(STATUS_HEADER)),
                after_value=self._status_code(source_values.get(STATUS_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{STATUS_HEADER}",
            ),
            self._field(
                canonical_path="device_record.market_availability.first_eu_market_country",
                business_label="First EU Market Country",
                before_value=self._string_value(source_values.get(FIRST_EU_MARKET_HEADER)),
                after_value=self._string_value(source_values.get(FIRST_EU_MARKET_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{FIRST_EU_MARKET_HEADER}",
            ),
            self._field(
                canonical_path="device_record.base_quantity",
                business_label="Base Quantity",
                before_value=self._string_value(source_values.get(QUANTITY_HEADER)),
                after_value=self._string_value(source_values.get(QUANTITY_HEADER)),
                before_source="workbook",
                after_source="workbook",
                source_detail=f"{row.sheet_name}!{QUANTITY_HEADER}",
            ),
        ]

        before_blockers = [
            f"{field.business_label} is not populated before Basic UDI enrichment."
            for field in fields
            if field.required and field.before_value is None
        ]
        after_blockers = [
            f"{field.business_label} remains missing after Basic UDI enrichment."
            for field in fields
            if field.required and field.after_value is None
        ]

        return EchelonValidationRecord(
            source_workbook=workbook_name,
            source_sheet=row.sheet_name,
            source_row_index=row.row_index,
            trade_name=self._string_value(source_values.get(TRADE_NAME_HEADER)),
            primary_udi_di=self._string_value(source_values.get(UDI_DI_HEADER)),
            catalogue_number=self._catalogue_number(source_values),
            issuing_entity=self._string_value(source_values.get(ISSUING_ENTITY_HEADER)),
            reference_match_status=match_status,
            basic_reference_material_number=self._string_value(basic_product.get(MATERIAL_NUMBER_HEADER)),
            basic_reference_name=self._string_value(basic_product.get(PRODUCT_NAME_HEADER)),
            before_completeness=self._completeness_snapshot(fields, state="before"),
            after_completeness=self._completeness_snapshot(fields, state="after"),
            before_blockers=before_blockers,
            after_blockers=after_blockers,
            fields=fields,
        )

    def _load_source_rows(self, workbook_path: Path) -> list[SourceRow]:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        rows: list[SourceRow] = []
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            header_row_index: int | None = None
            headers: list[str] | None = None
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = list(row)
                if headers is None and any(
                    isinstance(value, str) and HEADER_SENTINEL in value for value in values if value
                ):
                    headers = [str(value).strip() if value is not None else "" for value in values]
                    header_row_index = row_index
                    continue
                if headers is None or header_row_index is None or row_index <= header_row_index:
                    continue
                if not any(value not in (None, "") for value in values):
                    continue
                mapped_values: dict[str, object | None] = {
                    headers[index]: values[index] if index < len(values) else None
                    for index in range(len(headers))
                    if headers[index]
                }
                rows.append(SourceRow(sheet_name=sheet_name, row_index=row_index, values=mapped_values))
        return rows

    def _load_source_headers(self, workbook_path: Path) -> list[SourceHeaderProfile]:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        headers_by_name: dict[str, SourceHeaderProfile] = {}
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            for row in worksheet.iter_rows(values_only=True):
                values = [str(value).strip() if value is not None else "" for value in row]
                if not any(HEADER_SENTINEL in value for value in values if value):
                    continue
                for value in values:
                    if not value:
                        continue
                    normalized = self._normalize_header(value)
                    entry = headers_by_name.get(normalized)
                    if entry is None:
                        headers_by_name[normalized] = SourceHeaderProfile(
                            display_name=value,
                            normalized_name=normalized,
                            sheets=(sheet_name,),
                        )
                        continue
                    if sheet_name in entry.sheets:
                        continue
                    headers_by_name[normalized] = SourceHeaderProfile(
                        display_name=entry.display_name,
                        normalized_name=entry.normalized_name,
                        sheets=tuple(sorted((*entry.sheets, sheet_name))),
                    )
                break
        return [entry for _, entry in sorted(headers_by_name.items())]

    def _load_family_reference_context(
        self, workbook_path: Path
    ) -> dict[str, dict[str, object | None]] | None:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        product_rows = self._sheet_rows(workbook["Products"])
        basic_rows = self._sheet_rows(workbook["MDR-Basic"])

        basic_product = next(
            (
                row
                for row in product_rows
                if row.get(PRODUCT_TEMPLATE_HEADER) == "EUDAMED Basic UDI-DI"
                and self._string_value(row.get(PRODUCT_NAME_HEADER)) == "Echelon"
            ),
            None,
        )
        if not basic_product:
            return None

        basic_material_number = self._string_value(basic_product.get(MATERIAL_NUMBER_HEADER))
        if not basic_material_number:
            return None

        basic_detail_row = next(
            (
                row
                for row in basic_rows
                if self._string_value(row.get(MATERIAL_NUMBER_HEADER)) == basic_material_number
            ),
            None,
        )
        if not basic_detail_row:
            return None

        return {
            "basic_product": basic_product,
            "basic_details": basic_detail_row,
        }

    @staticmethod
    def _sheet_rows(worksheet: Worksheet) -> list[dict[str, object | None]]:
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        output: list[dict[str, object | None]] = []
        for row in rows[1:]:
            if not any(value not in (None, "") for value in row):
                continue
            output.append(
                {
                    headers[index]: row[index] if index < len(row) else None
                    for index in range(len(headers))
                    if headers[index]
                }
            )
        return output

    @staticmethod
    def _build_blocker_summaries(
        records: list[EchelonValidationRecord],
    ) -> list[BlockerSummary]:
        if not records:
            return []

        ordered_fields = records[0].fields
        summaries: list[BlockerSummary] = []
        for template_field in ordered_fields:
            before_missing_count = sum(
                1
                for record in records
                for field in record.fields
                if field.canonical_path == template_field.canonical_path and field.before_value is None
            )
            after_missing_count = sum(
                1
                for record in records
                for field in record.fields
                if field.canonical_path == template_field.canonical_path and field.after_value is None
            )
            summaries.append(
                BlockerSummary(
                    canonical_path=template_field.canonical_path,
                    business_label=template_field.business_label,
                    before_missing_count=before_missing_count,
                    after_missing_count=after_missing_count,
                )
            )
        return summaries

    @staticmethod
    def _build_sheet_summaries(
        records: list[EchelonValidationRecord],
    ) -> list[SheetValidationSummary]:
        grouped: dict[str, list[EchelonValidationRecord]] = {}
        for record in records:
            grouped.setdefault(record.source_sheet, []).append(record)

        summaries: list[SheetValidationSummary] = []
        for sheet_name, sheet_records in sorted(grouped.items()):
            summaries.append(
                SheetValidationSummary(
                    sheet_name=sheet_name,
                    record_count=len(sheet_records),
                    before_complete_records=sum(
                        1 for record in sheet_records if record.before_completeness.status == "complete"
                    ),
                    after_complete_records=sum(
                        1 for record in sheet_records if record.after_completeness.status == "complete"
                    ),
                    before_missing_field_total=sum(
                        record.before_completeness.missing_required_fields for record in sheet_records
                    ),
                    after_missing_field_total=sum(
                        record.after_completeness.missing_required_fields for record in sheet_records
                    ),
                )
            )
        return summaries

    def _build_source_field_coverage(
        self,
        headers: list[SourceHeaderProfile],
    ) -> list[SourceFieldCoverageEntry]:
        review_bundle = CanonicalReviewService().load_review_bundle()
        source_mappings: dict[str, list[tuple[str, list[str]]]] = {}
        for entity in review_bundle.entity_reviews:
            for field_review in entity.field_reviews:
                mapping = field_review.mapping
                schema_targets = [target.schema_path for target in mapping.schema_targets]
                for source_column in mapping.source_columns:
                    normalized = self._normalize_header(source_column)
                    source_mappings.setdefault(normalized, []).append((mapping.canonical_path, schema_targets))

        coverage_entries: list[SourceFieldCoverageEntry] = []
        for header in headers:
            mappings = source_mappings.get(header.normalized_name, [])
            canonical_targets = sorted({canonical_path for canonical_path, _ in mappings})
            schema_targets = sorted({target for _, targets in mappings for target in targets})
            if header.normalized_name in PARTIAL_SOURCE_HEADERS:
                status: SourceFieldCoverageStatus = "partially_represented"
                notes = PARTIAL_SOURCE_HEADERS[header.normalized_name]
            elif mappings:
                status = "represented"
                notes = REPRESENTED_SCOPE_NOTES.get(
                    header.normalized_name,
                    "This source field is mapped into the current Echelon canonical/XML-facing review path.",
                )
            else:
                status = "not_yet_represented"
                notes = "This source field does not yet have a documented canonical target in the current review artifact."
            coverage_entries.append(
                SourceFieldCoverageEntry(
                    source_field=header.display_name,
                    source_sheets=list(header.sheets),
                    coverage_status=status,
                    canonical_targets=canonical_targets,
                    schema_targets=schema_targets,
                    notes=notes,
                )
            )
        return coverage_entries

    @staticmethod
    def _build_source_field_coverage_summaries(
        entries: list[SourceFieldCoverageEntry],
    ) -> list[SourceFieldCoverageSummary]:
        labels = {
            "represented": "Represented",
            "partially_represented": "Partially represented",
            "not_yet_represented": "Not yet represented",
            "deferred_by_design": "Deferred by design",
        }
        ordered_statuses: tuple[SourceFieldCoverageStatus, ...] = (
            "represented",
            "partially_represented",
            "not_yet_represented",
            "deferred_by_design",
        )
        return [
            SourceFieldCoverageSummary(
                status=status,
                label=labels[status],
                field_count=sum(1 for entry in entries if entry.coverage_status == status),
            )
            for status in ordered_statuses
        ]

    @staticmethod
    def _build_sample_records(
        records: list[EchelonValidationRecord],
    ) -> list[EchelonValidationRecord]:
        grouped: dict[str, EchelonValidationRecord] = {}
        for record in records:
            grouped.setdefault(record.source_sheet, record)
        return [grouped[sheet_name] for sheet_name in sorted(grouped)]

    @staticmethod
    def _field(
        *,
        canonical_path: str,
        business_label: str,
        before_value: str | None,
        after_value: str | None,
        before_source: ValueSourceType,
        after_source: ValueSourceType,
        source_detail: str,
        update_reason: str | None = None,
        required: bool = True,
    ) -> ValidationFieldValue:
        return ValidationFieldValue(
            canonical_path=canonical_path,
            business_label=business_label,
            required=required,
            before_value=before_value,
            after_value=after_value,
            before_source=before_source,
            after_source=after_source,
            source_detail=source_detail,
            update_reason=update_reason,
        )

    @staticmethod
    def _completeness_snapshot(
        fields: list[ValidationFieldValue], *, state: Literal["before", "after"]
    ) -> CompletenessSnapshot:
        values = [field.before_value if state == "before" else field.after_value for field in fields if field.required]
        mapped_required_fields = sum(1 for value in values if value is not None)
        total_required_fields = len(values)
        missing_required_fields = total_required_fields - mapped_required_fields
        return CompletenessSnapshot(
            mapped_required_fields=mapped_required_fields,
            total_required_fields=total_required_fields,
            missing_required_fields=missing_required_fields,
            status="complete" if missing_required_fields == 0 else "incomplete",
        )

    @staticmethod
    def _catalogue_number(values: dict[str, object | None]) -> str | None:
        for header in CATALOGUE_HEADERS:
            value = values.get(header)
            if value not in (None, ""):
                return str(value).strip()
        return None

    @staticmethod
    def _di_identifier(*, issuing_entity: str | None, di_code: str | None) -> str | None:
        if not issuing_entity or not di_code:
            return None
        return f"{issuing_entity}:{di_code}"

    @staticmethod
    def _issuing_entity_code(value: object | None) -> str | None:
        raw = EchelonValidationService._string_value(value)
        if raw is None:
            return None
        mapping = {
            "gs1": "GS1",
            "hibcc": "HIBCC",
            "iccba": "ICCBBA",
            "ifi": "IFI",
        }
        return mapping.get(raw.lower(), raw.upper())

    @staticmethod
    def _status_code(value: object | None) -> str | None:
        raw = EchelonValidationService._string_value(value)
        if raw is None:
            return None
        normalized = raw.strip().upper().replace(" ", "_")
        mapping = {
            "ON_THE_EU_MARKET": "ON_THE_MARKET",
            "ON_THE_MARKET": "ON_THE_MARKET",
        }
        return mapping.get(normalized, normalized)

    @staticmethod
    def _production_identifier(value: object | None) -> str | None:
        raw = EchelonValidationService._string_value(value)
        if raw is None:
            return None
        mapping = {
            "serial number/ manufacturing date": "SERIALISATION_NUMBER",
            "serial number/manufacturing date": "SERIALISATION_NUMBER",
        }
        return mapping.get(raw.lower(), raw.upper().replace(" ", "_"))

    @staticmethod
    def _number_of_reuses(values: dict[str, object | None]) -> str | None:
        applicable = EchelonValidationService._string_value(values.get(MAX_REUSES_APPLICABLE_HEADER))
        single_use = EchelonValidationService._string_value(values.get(SINGLE_USE_HEADER))
        if single_use and single_use.strip().lower() in {"yes", "true"}:
            return "0"
        if applicable and applicable.strip().lower() in {"no", "false"}:
            return "1"
        return None

    @staticmethod
    def _string_value(value: object | None) -> str | None:
        if value in (None, ""):
            return None
        return str(value).strip()

    @staticmethod
    def _normalized_boolean(value: object | None) -> str | None:
        if value in (None, ""):
            return None
        token = str(value).strip().lower()
        if token in {"yes", "true"}:
            return "true"
        if token in {"no", "false"}:
            return "false"
        return str(value).strip()

    @staticmethod
    def _normalized_yes_no(value: object | None) -> str | None:
        if value in (None, ""):
            return None
        token = str(value).strip().lower()
        if token in {"yes", "true"}:
            return "Yes"
        if token in {"no", "false"}:
            return "No"
        return str(value).strip()

    @staticmethod
    def _normalize_header(value: str) -> str:
        return " ".join(value.split()).strip().lower()
