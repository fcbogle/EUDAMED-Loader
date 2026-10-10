from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
import hashlib
import json
import shutil
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.config import get_settings
from app.models import (
    DeviceIdentityIssueDetail,
    DeviceIdentityIssueSummary,
    DeviceSubjectDetail,
    DeviceSubjectSummary,
    DatabaseColumnSummary,
    DatabaseForeignKeySummary,
    DatabaseHealthIssue,
    DatabaseHealthSummary,
    DatabaseIndexSummary,
    DatabaseSchemaSummary,
    DatabaseTableHealthSummary,
    DatabaseTableSchemaSummary,
    ImportedWorkbookSummary,
    SourceRowDetail,
    SourceRowSummary,
    WorkbookImportDiffSummary,
    WorkbookImportBatchSummary,
    WorkbookImportWorkbookDiff,
    WorkbookImportDuplicateGroup,
    WorkbookImportOperationCount,
    WorkbookImportRunResponse,
    WorkbookImportSnapshotSummary,
    WorkbookImportTableCount,
)
from app.services.canonical_validation import (
    FAMILY_LABELS,
    HEADER_SENTINEL,
    CanonicalValidationService,
)
from app.validation_models import CanonicalValidationBundle, CanonicalValidationRecord


@dataclass(frozen=True)
class ImportedSourceRow:
    sheet_name: str
    row_index: int
    values: dict[str, object | None]


@dataclass(frozen=True)
class DeviceSubjectRecord:
    id: int
    subject_key: str
    product_family: str | None
    product_variant: str | None
    catalogue_number: str | None
    primary_udi_di: str | None
    basic_udi_di: str | None
    current_source_row_id: int | None
    created_at: str
    updated_at: str


@dataclass(frozen=True)
class DeviceSubjectMatchDecision:
    created: bool
    issue_code: str | None = None


class WorkbookImportService:
    def __init__(self) -> None:
        self.settings = get_settings()
        self.validation_service = CanonicalValidationService()
        self._ensure_schema()

    @property
    def db_path(self) -> Path:
        return self.settings.testing_state_db_path

    @staticmethod
    def _require_lastrowid(cursor: sqlite3.Cursor, *, entity_name: str) -> int:
        row_id = cursor.lastrowid
        if row_id is None:
            raise RuntimeError(f"SQLite did not return a row id for inserted {entity_name}.")
        return int(row_id)

    def run_import(
        self,
        *,
        imported_by: str | None = None,
        label: str | None = None,
        notes: str | None = None,
    ) -> WorkbookImportRunResponse:
        imported_at = datetime.now(UTC).replace(microsecond=0).isoformat()
        resolved_label = label or f"Workbook import {imported_at}"
        source_type = "source_excel"
        promoted_lookup = self._promotion_lookup()

        with self._connect() as connection:
            cursor = connection.execute(
                """
                INSERT INTO import_batch (source_type, label, imported_at, imported_by, notes)
                VALUES (?, ?, ?, ?, ?)
                """,
                (source_type, resolved_label, imported_at, imported_by, notes),
            )
            import_batch_id = self._require_lastrowid(cursor, entity_name="import batch")
            workbook_count = 0
            source_row_count = 0
            device_subject_count = 0

            for workbook_path in sorted(self.settings.excel_dir.glob("*.xlsx")):
                if workbook_path.name in self.settings.excluded_excel_workbook_names:
                    continue
                workbook_count += 1
                workbook_cursor = connection.execute(
                    """
                    INSERT INTO source_workbook (import_batch_id, workbook_name, file_path, file_hash, loaded_at)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        import_batch_id,
                        workbook_path.name,
                        str(workbook_path),
                        self._file_hash(workbook_path),
                        imported_at,
                    ),
                )
                source_workbook_id = self._require_lastrowid(workbook_cursor, entity_name="source workbook")
                workbook_rows = self._load_source_rows(workbook_path)
                for row in workbook_rows:
                    source_row_count += 1
                    promoted = promoted_lookup.get((workbook_path.name, row.sheet_name, row.row_index), {})
                    product_family = self._optional_string(promoted.get("product_family")) or FAMILY_LABELS.get(workbook_path.name)
                    product_variant = self._optional_string(promoted.get("product_variant")) or row.sheet_name
                    catalogue_number = self._optional_string(promoted.get("catalogue_number"))
                    primary_udi_di = self._optional_string(promoted.get("primary_udi_di"))
                    submission_operation = self._optional_string(promoted.get("submission_operation"))
                    canonical_status = self._optional_string(promoted.get("canonical_status")) or "not_yet_mapped"
                    basic_udi_di = self._optional_string(promoted.get("basic_udi_di"))
                    raw_payload_json = json.dumps(row.values, default=str, sort_keys=True)
                    row_cursor = connection.execute(
                        """
                        INSERT INTO source_row (
                            source_workbook_id,
                            sheet_name,
                            row_index,
                            product_family,
                            product_variant,
                            catalogue_number,
                            primary_udi_di,
                            submission_operation,
                            raw_payload_json,
                            canonical_status,
                            created_at
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            source_workbook_id,
                            row.sheet_name,
                            row.row_index,
                            product_family,
                            product_variant,
                            catalogue_number,
                            primary_udi_di,
                            submission_operation,
                            raw_payload_json,
                            canonical_status,
                            imported_at,
                        ),
                    )
                    source_row_id = self._require_lastrowid(row_cursor, entity_name="source row")
                    decision = self._match_device_subject(
                        connection=connection,
                        product_family=product_family,
                        product_variant=product_variant,
                        catalogue_number=catalogue_number,
                        primary_udi_di=primary_udi_di,
                        basic_udi_di=basic_udi_di,
                        current_source_row_id=source_row_id,
                        updated_at=imported_at,
                    )
                    if decision.created:
                        device_subject_count += 1

        from app.services.testing_state_store import TestingStateStore

        TestingStateStore().refresh_device_subject_links()
        self.rebuild_canonical_projection(import_batch_id=import_batch_id)
        self._create_import_backup(import_batch_id=import_batch_id)

        return WorkbookImportRunResponse(
            import_batch_id=import_batch_id,
            source_type=source_type,
            label=resolved_label,
            imported_at=imported_at,
            workbook_count=workbook_count,
            source_row_count=source_row_count,
            device_subject_count=device_subject_count,
        )

    def rebuild_canonical_projection(self, *, import_batch_id: int | None = None) -> int:
        target_batch_id = import_batch_id
        if target_batch_id is None:
            latest_batch = self.latest_import_batch()
            if latest_batch is None:
                return 0
            target_batch_id = latest_batch.import_batch_id

        bundle = self._validation_bundle()
        with self._connect() as connection:
            connection.execute("DELETE FROM canonical_projection_snapshot")
            connection.execute("DELETE FROM canonical_field_value")
            connection.execute("DELETE FROM canonical_device_record")
            source_row_lookup = self._source_row_lookup_for_batch(connection, import_batch_id=target_batch_id)
            persisted_count = 0
            for record in bundle.records:
                source_row = source_row_lookup.get((record.source_workbook, record.source_sheet, record.source_row_index))
                if source_row is None:
                    continue
                device_subject_id = self._device_subject_id_for_source_row(connection, source_row_id=source_row["id"])
                if device_subject_id is None:
                    continue
                canonical_cursor = connection.execute(
                    """
                    INSERT INTO canonical_device_record (
                        device_subject_id,
                        source_row_id,
                        source_import_batch_id,
                        canonical_version,
                        canonical_status,
                        completeness_status,
                        xml_readiness_status,
                        xml_ready,
                        record_json,
                        completeness_json,
                        blockers_json
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        device_subject_id,
                        int(source_row["id"]),
                        target_batch_id,
                        1,
                        "xml_ready" if record.xml_readiness.status == "complete" else "xml_blocked",
                        record.completeness.status,
                        record.xml_readiness.status,
                        int(record.xml_readiness.status == "complete"),
                        json.dumps(record.model_dump(mode="json"), sort_keys=True),
                        json.dumps(
                            {
                                "completeness": record.completeness.model_dump(mode="json"),
                                "xml_readiness": record.xml_readiness.model_dump(mode="json"),
                            },
                            sort_keys=True,
                        ),
                        json.dumps(
                            {
                                "blockers": record.blockers,
                                "xml_blockers": record.xml_blockers,
                            },
                            sort_keys=True,
                        ),
                    ),
                )
                canonical_device_record_id = self._require_lastrowid(
                    canonical_cursor,
                    entity_name="canonical device record",
                )
                for field in record.fields:
                    connection.execute(
                        """
                        INSERT INTO canonical_field_value (
                            canonical_device_record_id,
                            canonical_path,
                            field_status,
                            required,
                            value_json,
                            source_headers_json,
                            review_note
                        ) VALUES (?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            canonical_device_record_id,
                            field.canonical_path,
                            "populated" if field.value is not None else "missing",
                            int(field.required),
                            json.dumps(field.value) if field.value is not None else None,
                            json.dumps([field.source_detail], sort_keys=True)
                            if field.source_detail
                            else json.dumps([], sort_keys=True),
                            field.review_note,
                        ),
                    )
                persisted_count += 1
            connection.execute(
                """
                INSERT INTO canonical_projection_snapshot (
                    source_import_batch_id,
                    family_scope,
                    scope_note,
                    validation_note,
                    total_source_records,
                    validation_subset_records,
                    excluded_records,
                    matched_reference_records,
                    tracked_required_fields,
                    tracked_xml_required_fields,
                    ready_records,
                    blocked_records,
                    xml_ready_records,
                    xml_blocked_records,
                    source_field_total,
                    source_field_coverage_summaries_json,
                    source_field_coverage_json,
                    deferred_scope_summaries_json
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    target_batch_id,
                    bundle.family_scope,
                    bundle.scope_note,
                    bundle.validation_note,
                    bundle.total_source_records,
                    bundle.validation_subset_records,
                    bundle.excluded_records,
                    bundle.matched_reference_records,
                    bundle.tracked_required_fields,
                    bundle.tracked_xml_required_fields,
                    bundle.ready_records,
                    bundle.blocked_records,
                    bundle.xml_ready_records,
                    bundle.xml_blocked_records,
                    bundle.source_field_total,
                    json.dumps(
                        [summary.model_dump(mode="json") for summary in bundle.source_field_coverage_summaries],
                        sort_keys=True,
                    ),
                    json.dumps(
                        [entry.model_dump(mode="json") for entry in bundle.source_field_coverage],
                        sort_keys=True,
                    ),
                    json.dumps(
                        [summary.model_dump(mode="json") for summary in bundle.deferred_scope_summaries],
                        sort_keys=True,
                    ),
                ),
            )
        return persisted_count

    def canonical_validation_bundle_from_sqlite(
        self,
        *,
        import_batch_id: int | None = None,
    ) -> CanonicalValidationBundle | None:
        target_batch_id = import_batch_id
        if target_batch_id is None:
            latest_batch = self.latest_import_batch()
            if latest_batch is None:
                return None
            target_batch_id = latest_batch.import_batch_id

        with self._connect() as connection:
            snapshot_row = connection.execute(
                """
                SELECT *
                FROM canonical_projection_snapshot
                WHERE source_import_batch_id = ?
                LIMIT 1
                """,
                (target_batch_id,),
            ).fetchone()
            if snapshot_row is None:
                return None
            record_rows = connection.execute(
                """
                SELECT record_json
                FROM canonical_device_record
                WHERE source_import_batch_id = ?
                ORDER BY id
                """,
                (target_batch_id,),
            ).fetchall()

        records = [
            CanonicalValidationRecord.model_validate(json.loads(str(row["record_json"])))
            for row in record_rows
        ]
        if self.settings.environment == "prod":
            from app.services.production_completeness import populate_export_market_status
            for record in records:
                if populate_export_market_status(record):
                    record.completeness = self.validation_service._completeness_snapshot(record.fields)
                    record.blockers = [
                        f"{field.business_label} is not populated."
                        for field in record.fields if field.required and field.value is None
                    ]
        ready_records = (sum(record.completeness.status == "complete" for record in records)
                         if self.settings.environment == "prod" else int(snapshot_row["ready_records"]))
        return CanonicalValidationBundle(
            family_scope=str(snapshot_row["family_scope"]),
            scope_note=str(snapshot_row["scope_note"]),
            validation_note=str(snapshot_row["validation_note"]),
            total_source_records=int(snapshot_row["total_source_records"]),
            validation_subset_records=int(snapshot_row["validation_subset_records"]),
            excluded_records=int(snapshot_row["excluded_records"]),
            matched_reference_records=int(snapshot_row["matched_reference_records"]),
            tracked_required_fields=int(snapshot_row["tracked_required_fields"]),
            tracked_xml_required_fields=int(snapshot_row["tracked_xml_required_fields"]),
            ready_records=ready_records,
            blocked_records=(len(records) - ready_records if self.settings.environment == "prod"
                             else int(snapshot_row["blocked_records"])),
            xml_ready_records=int(snapshot_row["xml_ready_records"]),
            xml_blocked_records=int(snapshot_row["xml_blocked_records"]),
            family_summaries=self.validation_service._build_family_summaries(records),
            variant_summaries=self.validation_service._build_variant_summaries(records),
            blocker_summaries=self.validation_service._build_blocker_summaries(records),
            source_field_total=int(snapshot_row["source_field_total"]),
            source_field_coverage_summaries=json.loads(str(snapshot_row["source_field_coverage_summaries_json"]) or "[]"),
            source_field_coverage=json.loads(str(snapshot_row["source_field_coverage_json"]) or "[]"),
            sample_records=self.validation_service._build_sample_records(records),
            deferred_scope_summaries=json.loads(str(snapshot_row["deferred_scope_summaries_json"]) or "[]"),
            records=records,
        )

    def import_batch_count(self) -> int:
        with self._connect() as connection:
            row = connection.execute("SELECT COUNT(*) FROM import_batch").fetchone()
        return int(row[0]) if row is not None else 0

    def latest_import_batch(self) -> WorkbookImportBatchSummary | None:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT
                    batch.id,
                    batch.source_type,
                    batch.label,
                    batch.imported_at,
                    batch.imported_by,
                    batch.notes,
                    COUNT(DISTINCT workbook.id) AS workbook_count,
                    COUNT(DISTINCT source_row.id) AS source_row_count,
                    COUNT(DISTINCT subject.id) AS device_subject_count
                FROM import_batch batch
                LEFT JOIN source_workbook workbook ON workbook.import_batch_id = batch.id
                LEFT JOIN source_row ON source_row.source_workbook_id = workbook.id
                LEFT JOIN device_subject subject ON subject.current_source_row_id = source_row.id
                GROUP BY batch.id
                ORDER BY batch.id DESC
                LIMIT 1
                """
            ).fetchone()
        if row is None:
            return None
        return self._batch_summary_from_row(row)

    def imported_workbooks(self, *, import_batch_id: int) -> list[ImportedWorkbookSummary]:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT
                    workbook.id,
                    workbook.import_batch_id,
                    workbook.workbook_name,
                    workbook.file_path,
                    workbook.file_hash,
                    workbook.loaded_at,
                    COUNT(source_row.id) AS row_count
                FROM source_workbook workbook
                LEFT JOIN source_row ON source_row.source_workbook_id = workbook.id
                WHERE workbook.import_batch_id = ?
                GROUP BY workbook.id
                ORDER BY workbook.id
                """,
                (import_batch_id,),
            ).fetchall()
        return [
            ImportedWorkbookSummary(
                source_workbook_id=int(row["id"]),
                import_batch_id=int(row["import_batch_id"]),
                workbook_name=str(row["workbook_name"]),
                file_path=str(row["file_path"]),
                file_hash=str(row["file_hash"]),
                loaded_at=str(row["loaded_at"]),
                row_count=int(row["row_count"]),
            )
            for row in rows
        ]

    def latest_import_snapshot_summary(self) -> WorkbookImportSnapshotSummary | None:
        latest_batch = self.latest_import_batch()
        if latest_batch is None:
            return None

        imported_workbooks = self.imported_workbooks(import_batch_id=latest_batch.import_batch_id)
        with self._connect() as connection:
            available_tables = set(self._table_names(connection))
            latest_batch_identity_issue_count = int(
                connection.execute(
                    """
                    SELECT COUNT(*)
                    FROM device_identity_issue di
                    JOIN source_row sr ON sr.id = di.source_row_id
                    JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                    WHERE sw.import_batch_id = ?
                    """,
                    (latest_batch.import_batch_id,),
                ).fetchone()[0]
            )
            table_counts = [
                WorkbookImportTableCount(
                    table_name=table_name,
                    row_count=row_count,
                    summary_label=summary_label,
                )
                for table_name, summary_label, row_count in (
                    ("import_batch", "Import batches", 1),
                    ("source_workbook", "Workbook snapshots", latest_batch.workbook_count),
                    ("source_row", "Imported workbook rows", latest_batch.source_row_count),
                    ("device_subject", "Stable device subjects", latest_batch.device_subject_count),
                    ("device_identity_issue", "Identity issues", latest_batch_identity_issue_count),
                    ("testing_subjects", "Tracked testing subjects", self._safe_row_count(connection, "testing_subjects")),
                    ("testing_events", "Tracked testing events", self._safe_row_count(connection, "testing_events")),
                    ("generated_packages", "Generated ZIP packages", self._safe_row_count(connection, "generated_packages")),
                    ("reviewed_post_baselines", "Reviewed POST baselines", self._safe_row_count(connection, "reviewed_post_baselines")),
                )
                if table_name in available_tables
            ]
            operation_rows = connection.execute(
                """
                SELECT COALESCE(sr.submission_operation, 'UNCLASSIFIED') AS submission_operation, COUNT(*) AS device_subject_count
                FROM device_subject ds
                JOIN source_row sr ON sr.id = ds.current_source_row_id
                JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                WHERE sw.import_batch_id = ?
                GROUP BY COALESCE(sr.submission_operation, 'UNCLASSIFIED')
                ORDER BY submission_operation
                """,
                (latest_batch.import_batch_id,),
            ).fetchall()
            duplicate_subject_rows = connection.execute(
                """
                SELECT
                    ds.subject_key,
                    COUNT(sr.id) AS source_row_count,
                    MIN(ds.product_family) AS product_family,
                    MIN(ds.product_variant) AS product_variant,
                    MIN(ds.catalogue_number) AS catalogue_number,
                    MIN(ds.primary_udi_di) AS primary_udi_di,
                    GROUP_CONCAT(DISTINCT sw.workbook_name) AS workbook_names,
                    GROUP_CONCAT(DISTINCT sr.sheet_name) AS sheet_names,
                    GROUP_CONCAT(sr.row_index) AS row_indexes
                FROM device_subject ds
                JOIN source_row current_row ON current_row.id = ds.current_source_row_id
                JOIN source_workbook current_workbook ON current_workbook.id = current_row.source_workbook_id
                JOIN source_row sr
                  ON COALESCE(sr.catalogue_number, '') = COALESCE(ds.catalogue_number, '')
                 AND COALESCE(sr.primary_udi_di, '') = COALESCE(ds.primary_udi_di, '')
                 AND COALESCE(sr.product_family, '') = COALESCE(ds.product_family, '')
                 AND COALESCE(sr.product_variant, '') = COALESCE(ds.product_variant, '')
                JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                WHERE current_workbook.import_batch_id = ?
                  AND sw.import_batch_id = current_workbook.import_batch_id
                GROUP BY ds.subject_key
                HAVING COUNT(sr.id) > 1
                ORDER BY source_row_count DESC, ds.subject_key
                LIMIT 8
                """,
                (latest_batch.import_batch_id,),
            ).fetchall()
            duplicate_subject_count = int(
                connection.execute(
                    """
                    SELECT COUNT(*)
                    FROM (
                        SELECT ds.subject_key
                        FROM device_subject ds
                        JOIN source_row current_row ON current_row.id = ds.current_source_row_id
                        JOIN source_workbook current_workbook ON current_workbook.id = current_row.source_workbook_id
                        JOIN source_row sr
                         ON COALESCE(sr.catalogue_number, '') = COALESCE(ds.catalogue_number, '')
                         AND COALESCE(sr.primary_udi_di, '') = COALESCE(ds.primary_udi_di, '')
                         AND COALESCE(sr.product_family, '') = COALESCE(ds.product_family, '')
                         AND COALESCE(sr.product_variant, '') = COALESCE(ds.product_variant, '')
                        JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                        WHERE current_workbook.import_batch_id = ?
                          AND sw.import_batch_id = current_workbook.import_batch_id
                        GROUP BY ds.subject_key
                        HAVING COUNT(sr.id) > 1
                    )
                    """,
                    (latest_batch.import_batch_id,),
                ).fetchone()[0]
            )
            workbook_duplicate_row_count = int(
                connection.execute(
                    """
                    SELECT COALESCE(SUM(duplicate_group.row_count), 0)
                    FROM (
                        SELECT COUNT(*) AS row_count
                        FROM source_row sr
                        JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                        WHERE sw.import_batch_id = ?
                          AND COALESCE(TRIM(sr.product_family), '') <> ''
                          AND COALESCE(TRIM(sr.product_variant), '') <> ''
                          AND COALESCE(TRIM(sr.catalogue_number), '') <> ''
                        GROUP BY
                            COALESCE(TRIM(sr.product_family), ''),
                            COALESCE(TRIM(sr.product_variant), ''),
                            COALESCE(TRIM(sr.catalogue_number), '')
                        HAVING COUNT(*) > 1
                    ) duplicate_group
                    """,
                    (latest_batch.import_batch_id,),
                ).fetchone()[0]
            )
            workbook_duplicate_group_count = int(
                connection.execute(
                    """
                    SELECT COUNT(*)
                    FROM (
                        SELECT 1
                        FROM source_row sr
                        JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                        WHERE sw.import_batch_id = ?
                          AND COALESCE(TRIM(sr.product_family), '') <> ''
                          AND COALESCE(TRIM(sr.product_variant), '') <> ''
                          AND COALESCE(TRIM(sr.catalogue_number), '') <> ''
                        GROUP BY
                            COALESCE(TRIM(sr.product_family), ''),
                            COALESCE(TRIM(sr.product_variant), ''),
                            COALESCE(TRIM(sr.catalogue_number), '')
                        HAVING COUNT(*) > 1
                    ) duplicate_group
                    """,
                    (latest_batch.import_batch_id,),
                ).fetchone()[0]
            )
            merged_source_row_count = max(
                latest_batch.source_row_count - latest_batch.device_subject_count - latest_batch_identity_issue_count,
                0,
            )
            snapshot_row = (
                connection.execute(
                    """
                    SELECT source_import_batch_id, total_source_records
                    FROM canonical_projection_snapshot
                    WHERE source_import_batch_id = ?
                    LIMIT 1
                    """,
                    (latest_batch.import_batch_id,),
                ).fetchone()
                if "canonical_projection_snapshot" in available_tables
                else None
            )
            canonical_projection_import_batch_id = (
                int(snapshot_row["source_import_batch_id"]) if snapshot_row is not None else None
            )
            if snapshot_row is None:
                canonical_projection_status = "missing"
            elif int(snapshot_row["total_source_records"]) == latest_batch.source_row_count:
                canonical_projection_status = "ready"
            else:
                canonical_projection_status = "stale"

        operation_counts = [
            WorkbookImportOperationCount(
                submission_operation=str(row["submission_operation"]),
                device_subject_count=int(row["device_subject_count"]),
            )
            for row in operation_rows
        ]
        duplicate_groups = [
            WorkbookImportDuplicateGroup(
                subject_key=str(row["subject_key"]),
                source_row_count=int(row["source_row_count"]),
                product_family=self._optional_string(row["product_family"]),
                product_variant=self._optional_string(row["product_variant"]),
                catalogue_number=self._optional_string(row["catalogue_number"]),
                primary_udi_di=self._optional_string(row["primary_udi_di"]),
                workbook_names=self._split_csv_values(row["workbook_names"]),
                sheet_names=self._split_csv_values(row["sheet_names"]),
                row_indexes=self._split_csv_ints(row["row_indexes"]),
            )
            for row in duplicate_subject_rows
        ]
        unresolved_identity_row_count = self._unresolved_identity_row_count_for_batch(
            latest_batch.import_batch_id,
        )

        return WorkbookImportSnapshotSummary(
                import_batch=latest_batch,
                imported_workbooks=imported_workbooks,
                table_counts=table_counts,
                operation_counts=operation_counts,
                canonical_projection_status=canonical_projection_status,
                canonical_projection_import_batch_id=canonical_projection_import_batch_id,
                duplicate_source_row_delta=merged_source_row_count,
                merged_source_row_count=merged_source_row_count,
                duplicate_subject_count=duplicate_subject_count,
                workbook_duplicate_row_count=workbook_duplicate_row_count,
                workbook_duplicate_group_count=workbook_duplicate_group_count,
                unresolved_identity_row_count=unresolved_identity_row_count,
                top_duplicate_groups=duplicate_groups,
        )
        

    def schema_summary(self) -> DatabaseSchemaSummary:
        with self._connect() as connection:
            table_summaries = [
                DatabaseTableSchemaSummary(
                    table_name=table_name,
                    row_count=self._safe_row_count(connection, table_name),
                    columns=self._columns_for_table(connection, table_name),
                    foreign_keys=self._foreign_keys_for_table(connection, table_name),
                    indexes=self._indexes_for_table(connection, table_name),
                )
                for table_name in self._table_names(connection)
            ]
        return DatabaseSchemaSummary(
            db_path=str(self.db_path),
            table_count=len(table_summaries),
            tables=table_summaries,
        )

    def database_health_summary(self) -> DatabaseHealthSummary:
        generated_at = datetime.now(UTC).replace(microsecond=0).isoformat()
        with self._connect() as connection:
            available_tables = set(self._table_names(connection))
            issues: list[DatabaseHealthIssue] = []
            expected_tables = {
                "import_batch",
                "source_workbook",
                "source_row",
                "device_subject",
                "device_identity_issue",
                "testing_subjects",
                "testing_events",
                "generated_packages",
                "reviewed_post_baselines",
            }
            for table_name in sorted(expected_tables.difference(available_tables)):
                issues.append(
                    DatabaseHealthIssue(
                        level="warning",
                        code="missing_table",
                        message=f"Expected table {table_name} is not present in the current SQLite file.",
                        table_name=table_name,
                    )
                )

            table_summaries = [
                DatabaseTableHealthSummary(
                    table_name="import_batch",
                    row_count=self._safe_row_count(connection, "import_batch"),
                    orphan_count=0,
                    identity_gap_count=0,
                ),
                DatabaseTableHealthSummary(
                    table_name="source_workbook",
                    row_count=self._safe_row_count(connection, "source_workbook"),
                    orphan_count=self._safe_scalar(
                        connection,
                        """
                        SELECT COUNT(*)
                        FROM source_workbook workbook
                        LEFT JOIN import_batch batch ON batch.id = workbook.import_batch_id
                        WHERE batch.id IS NULL
                        """,
                        "source_workbook",
                    ),
                    identity_gap_count=0,
                ),
                DatabaseTableHealthSummary(
                    table_name="source_row",
                    row_count=self._safe_row_count(connection, "source_row"),
                    orphan_count=self._safe_scalar(
                        connection,
                        """
                        SELECT COUNT(*)
                        FROM source_row row
                        LEFT JOIN source_workbook workbook ON workbook.id = row.source_workbook_id
                        WHERE workbook.id IS NULL
                        """,
                        "source_row",
                    ),
                    identity_gap_count=self._safe_scalar(
                        connection,
                        """
                        SELECT COUNT(*)
                        FROM source_row
                        WHERE COALESCE(TRIM(product_family), '') = ''
                           OR COALESCE(TRIM(product_variant), '') = ''
                           OR (
                               COALESCE(TRIM(catalogue_number), '') = ''
                               AND COALESCE(TRIM(primary_udi_di), '') = ''
                           )
                        """,
                        "source_row",
                    ),
                ),
                DatabaseTableHealthSummary(
                    table_name="device_subject",
                    row_count=self._safe_row_count(connection, "device_subject"),
                    orphan_count=self._safe_scalar(
                        connection,
                        """
                        SELECT COUNT(*)
                        FROM device_subject subject
                        LEFT JOIN source_row row ON row.id = subject.current_source_row_id
                        WHERE subject.current_source_row_id IS NOT NULL
                          AND row.id IS NULL
                        """,
                        "device_subject",
                    ),
                    identity_gap_count=self._safe_scalar(
                        connection,
                        """
                        SELECT COUNT(*)
                        FROM device_subject
                        WHERE COALESCE(TRIM(product_family), '') = ''
                           OR COALESCE(TRIM(product_variant), '') = ''
                           OR (
                               COALESCE(TRIM(catalogue_number), '') = ''
                               AND COALESCE(TRIM(primary_udi_di), '') = ''
                           )
                        """,
                        "device_subject",
                    ),
                ),
                DatabaseTableHealthSummary(
                    table_name="device_identity_issue",
                    row_count=self._safe_row_count(connection, "device_identity_issue"),
                    orphan_count=self._safe_scalar(
                        connection,
                        """
                        SELECT COUNT(*)
                        FROM device_identity_issue issue
                        LEFT JOIN source_row row ON row.id = issue.source_row_id
                        LEFT JOIN device_subject subject ON subject.id = issue.device_subject_id
                        WHERE row.id IS NULL
                           OR (
                               issue.device_subject_id IS NOT NULL
                               AND subject.id IS NULL
                           )
                        """,
                        "device_identity_issue",
                    ),
                    identity_gap_count=0,
                ),
                DatabaseTableHealthSummary(
                    table_name="testing_subjects",
                    row_count=self._safe_row_count(connection, "testing_subjects"),
                    orphan_count=0,
                    identity_gap_count=self._safe_scalar(
                        connection,
                        """
                        SELECT COUNT(*)
                        FROM testing_subjects
                        WHERE COALESCE(TRIM(product_family), '') = ''
                           OR COALESCE(TRIM(product_variant), '') = ''
                           OR (
                               COALESCE(TRIM(catalogue_number), '') = ''
                               AND COALESCE(TRIM(primary_udi_di), '') = ''
                           )
                        """,
                        "testing_subjects",
                    ),
                ),
                DatabaseTableHealthSummary(
                    table_name="testing_events",
                    row_count=self._safe_row_count(connection, "testing_events"),
                    orphan_count=self._safe_scalar(
                        connection,
                        """
                        SELECT COUNT(*)
                        FROM testing_events event
                        LEFT JOIN testing_subjects subject ON subject.id = event.subject_id
                        WHERE subject.id IS NULL
                        """,
                        "testing_events",
                    ),
                    identity_gap_count=0,
                ),
                DatabaseTableHealthSummary(
                    table_name="generated_packages",
                    row_count=self._safe_row_count(connection, "generated_packages"),
                    orphan_count=0,
                    identity_gap_count=0,
                ),
                DatabaseTableHealthSummary(
                    table_name="reviewed_post_baselines",
                    row_count=self._safe_row_count(connection, "reviewed_post_baselines"),
                    orphan_count=0,
                    identity_gap_count=self._safe_scalar(
                        connection,
                        """
                        SELECT COUNT(*)
                        FROM reviewed_post_baselines
                        WHERE COALESCE(TRIM(product_family), '') = ''
                           OR COALESCE(TRIM(product_variant), '') = ''
                           OR COALESCE(TRIM(catalogue_number), '') = ''
                        """,
                        "reviewed_post_baselines",
                    ),
                ),
            ]

        for summary in table_summaries:
            if summary.orphan_count:
                issues.append(
                    DatabaseHealthIssue(
                        level="error",
                        code="orphan_rows",
                        message=f"{summary.orphan_count} orphaned row(s) detected in {summary.table_name}.",
                        table_name=summary.table_name,
                    )
                )
            if summary.identity_gap_count:
                issues.append(
                    DatabaseHealthIssue(
                        level="warning",
                        code="identity_gap",
                        message=f"{summary.identity_gap_count} row(s) in {summary.table_name} are missing core identity fields.",
                        table_name=summary.table_name,
                    )
                )

        return DatabaseHealthSummary(
            db_path=str(self.db_path),
            generated_at=generated_at,
            table_summaries=[summary for summary in table_summaries if summary.row_count or summary.table_name in available_tables],
            issues=issues,
        )

    def latest_import_diff_summary(self) -> WorkbookImportDiffSummary | None:
        with self._connect() as connection:
            rows = connection.execute(
                """
                SELECT
                    batch.id,
                    batch.label,
                    COUNT(DISTINCT workbook.id) AS workbook_count,
                    COUNT(DISTINCT source_row.id) AS source_row_count
                FROM import_batch batch
                LEFT JOIN source_workbook workbook ON workbook.import_batch_id = batch.id
                LEFT JOIN source_row ON source_row.source_workbook_id = workbook.id
                GROUP BY batch.id
                ORDER BY batch.id DESC
                LIMIT 2
                """
            ).fetchall()
            if not rows:
                return None
            current_row = rows[0]
            previous_row = rows[1] if len(rows) > 1 else None
            current_device_subject_count = self._device_subject_count_for_batch(connection, int(current_row["id"]))
            previous_device_subject_count = (
                self._device_subject_count_for_batch(connection, int(previous_row["id"])) if previous_row is not None else 0
            )
            changed_workbooks = self._workbook_diffs(
                connection,
                current_import_batch_id=int(current_row["id"]),
                previous_import_batch_id=int(previous_row["id"]) if previous_row is not None else None,
            )
        return WorkbookImportDiffSummary(
            current_import_batch_id=int(current_row["id"]),
            previous_import_batch_id=int(previous_row["id"]) if previous_row is not None else None,
            current_label=str(current_row["label"]),
            previous_label=self._optional_string(previous_row["label"]) if previous_row is not None else None,
            source_row_delta=int(current_row["source_row_count"]) - int(previous_row["source_row_count"]) if previous_row is not None else int(current_row["source_row_count"]),
            device_subject_delta=current_device_subject_count - previous_device_subject_count if previous_row is not None else current_device_subject_count,
            workbook_count_delta=int(current_row["workbook_count"]) - int(previous_row["workbook_count"]) if previous_row is not None else int(current_row["workbook_count"]),
            changed_workbooks=changed_workbooks,
        )

    def list_device_subjects(
        self,
        *,
        product_family: str | None = None,
        product_variant: str | None = None,
        catalogue_number: str | None = None,
        import_batch_id: int | None = None,
        limit: int = 200,
    ) -> list[DeviceSubjectSummary]:
        clauses: list[str] = []
        params: list[Any] = []
        if product_family:
            clauses.append("COALESCE(ds.product_family, '') = ?")
            params.append(product_family)
        if product_variant:
            clauses.append("COALESCE(ds.product_variant, '') = ?")
            params.append(product_variant)
        if catalogue_number:
            clauses.append("COALESCE(ds.catalogue_number, '') = ?")
            params.append(catalogue_number)
        if import_batch_id is not None:
            clauses.append("sw.import_batch_id = ?")
            params.append(import_batch_id)
        where_clause = f"WHERE {' AND '.join(clauses)}" if clauses else ""
        with self._connect() as connection:
            rows = connection.execute(
                f"""
                SELECT
                    ds.id,
                    ds.subject_key,
                    ds.product_family,
                    ds.product_variant,
                    ds.catalogue_number,
                    ds.primary_udi_di,
                    ds.basic_udi_di,
                    ds.current_source_row_id,
                    sw.import_batch_id AS current_import_batch_id,
                    ds.created_at,
                    ds.updated_at
                FROM device_subject ds
                LEFT JOIN source_row sr ON sr.id = ds.current_source_row_id
                LEFT JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                {where_clause}
                ORDER BY ds.product_family, ds.product_variant, ds.catalogue_number, ds.id
                LIMIT ?
                """,
                (*params, limit),
            ).fetchall()
        return [
            DeviceSubjectSummary(
                id=int(row["id"]),
                subject_key=str(row["subject_key"]),
                product_family=self._optional_string(row["product_family"]),
                product_variant=self._optional_string(row["product_variant"]),
                catalogue_number=self._optional_string(row["catalogue_number"]),
                primary_udi_di=self._optional_string(row["primary_udi_di"]),
                basic_udi_di=self._optional_string(row["basic_udi_di"]),
                current_source_row_id=int(row["current_source_row_id"]) if row["current_source_row_id"] is not None else None,
                current_import_batch_id=int(row["current_import_batch_id"]) if row["current_import_batch_id"] is not None else None,
                created_at=str(row["created_at"]),
                updated_at=str(row["updated_at"]),
            )
            for row in rows
        ]

    def get_device_subject(self, subject_id: int) -> DeviceSubjectDetail | None:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT
                    ds.id,
                    ds.subject_key,
                    ds.product_family,
                    ds.product_variant,
                    ds.catalogue_number,
                    ds.primary_udi_di,
                    ds.basic_udi_di,
                    ds.current_source_row_id,
                    sw.import_batch_id AS current_import_batch_id,
                    ds.created_at,
                    ds.updated_at,
                    sw.workbook_name AS current_source_workbook_name,
                    sr.sheet_name AS current_source_sheet_name,
                    sr.row_index AS current_source_row_index
                FROM device_subject ds
                LEFT JOIN source_row sr ON sr.id = ds.current_source_row_id
                LEFT JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                WHERE ds.id = ?
                """,
                (subject_id,),
            ).fetchone()
        if row is None:
            return None
        return DeviceSubjectDetail(
            id=int(row["id"]),
            subject_key=str(row["subject_key"]),
            product_family=self._optional_string(row["product_family"]),
            product_variant=self._optional_string(row["product_variant"]),
            catalogue_number=self._optional_string(row["catalogue_number"]),
            primary_udi_di=self._optional_string(row["primary_udi_di"]),
            basic_udi_di=self._optional_string(row["basic_udi_di"]),
            current_source_row_id=int(row["current_source_row_id"]) if row["current_source_row_id"] is not None else None,
            current_import_batch_id=int(row["current_import_batch_id"]) if row["current_import_batch_id"] is not None else None,
            created_at=str(row["created_at"]),
            updated_at=str(row["updated_at"]),
            current_source_workbook_name=self._optional_string(row["current_source_workbook_name"]),
            current_source_sheet_name=self._optional_string(row["current_source_sheet_name"]),
            current_source_row_index=int(row["current_source_row_index"]) if row["current_source_row_index"] is not None else None,
        )

    def list_source_rows(
        self,
        *,
        product_family: str | None = None,
        product_variant: str | None = None,
        catalogue_number: str | None = None,
        submission_operation: str | None = None,
        import_batch_id: int | None = None,
        limit: int = 200,
    ) -> list[SourceRowSummary]:
        clauses: list[str] = []
        params: list[Any] = []
        if product_family:
            clauses.append("COALESCE(sr.product_family, '') = ?")
            params.append(product_family)
        if product_variant:
            clauses.append("COALESCE(sr.product_variant, '') = ?")
            params.append(product_variant)
        if catalogue_number:
            clauses.append("COALESCE(sr.catalogue_number, '') = ?")
            params.append(catalogue_number)
        if submission_operation:
            clauses.append("COALESCE(sr.submission_operation, '') = ?")
            params.append(submission_operation)
        if import_batch_id is not None:
            clauses.append("sw.import_batch_id = ?")
            params.append(import_batch_id)
        where_clause = f"WHERE {' AND '.join(clauses)}" if clauses else ""
        with self._connect() as connection:
            rows = connection.execute(
                f"""
                SELECT
                    sr.id,
                    sr.source_workbook_id,
                    sw.import_batch_id,
                    sw.workbook_name,
                    sr.sheet_name,
                    sr.row_index,
                    sr.product_family,
                    sr.product_variant,
                    sr.catalogue_number,
                    sr.primary_udi_di,
                    sr.submission_operation,
                    sr.canonical_status,
                    sr.created_at
                FROM source_row sr
                JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                {where_clause}
                ORDER BY sw.import_batch_id DESC, sw.workbook_name, sr.sheet_name, sr.row_index
                LIMIT ?
                """,
                (*params, limit),
            ).fetchall()
        return [
            SourceRowSummary(
                id=int(row["id"]),
                source_workbook_id=int(row["source_workbook_id"]),
                import_batch_id=int(row["import_batch_id"]),
                workbook_name=str(row["workbook_name"]),
                sheet_name=str(row["sheet_name"]),
                row_index=int(row["row_index"]),
                product_family=self._optional_string(row["product_family"]),
                product_variant=self._optional_string(row["product_variant"]),
                catalogue_number=self._optional_string(row["catalogue_number"]),
                primary_udi_di=self._optional_string(row["primary_udi_di"]),
                submission_operation=self._optional_string(row["submission_operation"]),
                canonical_status=self._optional_string(row["canonical_status"]),
                created_at=str(row["created_at"]),
            )
            for row in rows
        ]

    def get_source_row(self, source_row_id: int) -> SourceRowDetail | None:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT
                    sr.id,
                    sr.source_workbook_id,
                    sw.import_batch_id,
                    sw.workbook_name,
                    sr.sheet_name,
                    sr.row_index,
                    sr.product_family,
                    sr.product_variant,
                    sr.catalogue_number,
                    sr.primary_udi_di,
                    sr.submission_operation,
                    sr.canonical_status,
                    sr.created_at,
                    sr.raw_payload_json,
                    ds.id AS linked_device_subject_id,
                    ds.subject_key AS linked_device_subject_key
                FROM source_row sr
                JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                LEFT JOIN device_subject ds ON ds.current_source_row_id = sr.id
                WHERE sr.id = ?
                """,
                (source_row_id,),
            ).fetchone()
        if row is None:
            return None
        return SourceRowDetail(
            id=int(row["id"]),
            source_workbook_id=int(row["source_workbook_id"]),
            import_batch_id=int(row["import_batch_id"]),
            workbook_name=str(row["workbook_name"]),
            sheet_name=str(row["sheet_name"]),
            row_index=int(row["row_index"]),
            product_family=self._optional_string(row["product_family"]),
            product_variant=self._optional_string(row["product_variant"]),
            catalogue_number=self._optional_string(row["catalogue_number"]),
            primary_udi_di=self._optional_string(row["primary_udi_di"]),
            submission_operation=self._optional_string(row["submission_operation"]),
            canonical_status=self._optional_string(row["canonical_status"]),
            created_at=str(row["created_at"]),
            raw_payload_json=str(row["raw_payload_json"]),
            linked_device_subject_id=int(row["linked_device_subject_id"]) if row["linked_device_subject_id"] is not None else None,
            linked_device_subject_key=self._optional_string(row["linked_device_subject_key"]),
        )

    def list_device_identity_issues(
        self,
        *,
        issue_code: str | None = None,
        product_family: str | None = None,
        product_variant: str | None = None,
        catalogue_number: str | None = None,
        import_batch_id: int | None = None,
        limit: int = 200,
    ) -> list[DeviceIdentityIssueSummary]:
        clauses: list[str] = []
        params: list[Any] = []
        if issue_code:
            clauses.append("issue.issue_code = ?")
            params.append(issue_code)
        if product_family:
            clauses.append("COALESCE(sr.product_family, '') = ?")
            params.append(product_family)
        if product_variant:
            clauses.append("COALESCE(sr.product_variant, '') = ?")
            params.append(product_variant)
        if catalogue_number:
            clauses.append("COALESCE(sr.catalogue_number, '') = ?")
            params.append(catalogue_number)
        if import_batch_id is not None:
            clauses.append("sw.import_batch_id = ?")
            params.append(import_batch_id)
        where_clause = f"WHERE {' AND '.join(clauses)}" if clauses else ""
        with self._connect() as connection:
            rows = connection.execute(
                f"""
                SELECT
                    issue.id,
                    issue.source_row_id,
                    issue.device_subject_id,
                    issue.issue_code,
                    issue.severity,
                    issue.created_at,
                    sr.product_family,
                    sr.product_variant,
                    sr.catalogue_number,
                    sr.primary_udi_di,
                    sw.import_batch_id
                FROM device_identity_issue issue
                JOIN source_row sr ON sr.id = issue.source_row_id
                JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                {where_clause}
                ORDER BY issue.created_at DESC, issue.id DESC
                LIMIT ?
                """,
                (*params, limit),
            ).fetchall()
        return [
            DeviceIdentityIssueSummary(
                id=int(row["id"]),
                source_row_id=int(row["source_row_id"]),
                device_subject_id=int(row["device_subject_id"]) if row["device_subject_id"] is not None else None,
                issue_code=str(row["issue_code"]),
                severity=str(row["severity"]),
                created_at=str(row["created_at"]),
                product_family=self._optional_string(row["product_family"]),
                product_variant=self._optional_string(row["product_variant"]),
                catalogue_number=self._optional_string(row["catalogue_number"]),
                primary_udi_di=self._optional_string(row["primary_udi_di"]),
                import_batch_id=int(row["import_batch_id"]) if row["import_batch_id"] is not None else None,
            )
            for row in rows
        ]

    def get_device_identity_issue(self, issue_id: int) -> DeviceIdentityIssueDetail | None:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT
                    issue.id,
                    issue.source_row_id,
                    issue.device_subject_id,
                    issue.issue_code,
                    issue.severity,
                    issue.details_json,
                    issue.created_at,
                    issue.resolved_at,
                    issue.resolution_note,
                    sr.product_family,
                    sr.product_variant,
                    sr.catalogue_number,
                    sr.primary_udi_di,
                    sw.import_batch_id
                FROM device_identity_issue issue
                JOIN source_row sr ON sr.id = issue.source_row_id
                JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                WHERE issue.id = ?
                """,
                (issue_id,),
            ).fetchone()
        if row is None:
            return None
        return DeviceIdentityIssueDetail(
            id=int(row["id"]),
            source_row_id=int(row["source_row_id"]),
            device_subject_id=int(row["device_subject_id"]) if row["device_subject_id"] is not None else None,
            issue_code=str(row["issue_code"]),
            severity=str(row["severity"]),
            created_at=str(row["created_at"]),
            product_family=self._optional_string(row["product_family"]),
            product_variant=self._optional_string(row["product_variant"]),
            catalogue_number=self._optional_string(row["catalogue_number"]),
            primary_udi_di=self._optional_string(row["primary_udi_di"]),
            import_batch_id=int(row["import_batch_id"]) if row["import_batch_id"] is not None else None,
            details_json=self._json_object(row["details_json"]),
            resolved_at=self._optional_string(row["resolved_at"]),
            resolution_note=self._optional_string(row["resolution_note"]),
        )

    def _ensure_schema(self) -> None:
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        with self._connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS import_batch (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_type TEXT NOT NULL,
                    label TEXT NOT NULL,
                    imported_at TEXT NOT NULL,
                    imported_by TEXT,
                    notes TEXT
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS source_workbook (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    import_batch_id INTEGER NOT NULL,
                    workbook_name TEXT NOT NULL,
                    file_path TEXT NOT NULL,
                    file_hash TEXT NOT NULL,
                    loaded_at TEXT NOT NULL,
                    FOREIGN KEY(import_batch_id) REFERENCES import_batch(id) ON DELETE CASCADE
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS source_row (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_workbook_id INTEGER NOT NULL,
                    sheet_name TEXT NOT NULL,
                    row_index INTEGER NOT NULL,
                    product_family TEXT,
                    product_variant TEXT,
                    catalogue_number TEXT,
                    primary_udi_di TEXT,
                    submission_operation TEXT,
                    raw_payload_json TEXT NOT NULL,
                    canonical_status TEXT,
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(source_workbook_id) REFERENCES source_workbook(id) ON DELETE CASCADE
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS device_subject (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    subject_key TEXT NOT NULL UNIQUE,
                    product_family TEXT,
                    product_variant TEXT,
                    catalogue_number TEXT,
                    primary_udi_di TEXT,
                    basic_udi_di TEXT,
                    current_source_row_id INTEGER,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    FOREIGN KEY(current_source_row_id) REFERENCES source_row(id) ON DELETE SET NULL
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS device_identity_issue (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_row_id INTEGER NOT NULL,
                    device_subject_id INTEGER,
                    issue_code TEXT NOT NULL,
                    severity TEXT NOT NULL,
                    details_json TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    resolved_at TEXT,
                    resolution_note TEXT,
                    FOREIGN KEY(source_row_id) REFERENCES source_row(id) ON DELETE CASCADE,
                    FOREIGN KEY(device_subject_id) REFERENCES device_subject(id) ON DELETE SET NULL
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS canonical_device_record (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    device_subject_id INTEGER NOT NULL UNIQUE,
                    source_row_id INTEGER,
                    source_import_batch_id INTEGER,
                    canonical_version INTEGER NOT NULL DEFAULT 1,
                    canonical_status TEXT NOT NULL DEFAULT 'draft',
                    completeness_status TEXT,
                    xml_readiness_status TEXT,
                    xml_ready INTEGER NOT NULL DEFAULT 0,
                    record_json TEXT,
                    completeness_json TEXT,
                    blockers_json TEXT,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY(device_subject_id) REFERENCES device_subject(id) ON DELETE CASCADE,
                    FOREIGN KEY(source_row_id) REFERENCES source_row(id) ON DELETE SET NULL,
                    FOREIGN KEY(source_import_batch_id) REFERENCES import_batch(id) ON DELETE SET NULL
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS canonical_field_value (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    canonical_device_record_id INTEGER NOT NULL,
                    canonical_path TEXT NOT NULL,
                    field_status TEXT NOT NULL DEFAULT 'pending',
                    required INTEGER NOT NULL DEFAULT 0,
                    value_json TEXT,
                    source_headers_json TEXT,
                    review_note TEXT,
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY(canonical_device_record_id) REFERENCES canonical_device_record(id) ON DELETE CASCADE,
                    UNIQUE(canonical_device_record_id, canonical_path)
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS canonical_projection_snapshot (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_import_batch_id INTEGER NOT NULL UNIQUE,
                    family_scope TEXT NOT NULL,
                    scope_note TEXT NOT NULL,
                    validation_note TEXT NOT NULL,
                    total_source_records INTEGER NOT NULL,
                    validation_subset_records INTEGER NOT NULL,
                    excluded_records INTEGER NOT NULL,
                    matched_reference_records INTEGER NOT NULL,
                    tracked_required_fields INTEGER NOT NULL,
                    tracked_xml_required_fields INTEGER NOT NULL,
                    ready_records INTEGER NOT NULL,
                    blocked_records INTEGER NOT NULL,
                    xml_ready_records INTEGER NOT NULL,
                    xml_blocked_records INTEGER NOT NULL,
                    source_field_total INTEGER NOT NULL DEFAULT 0,
                    source_field_coverage_summaries_json TEXT NOT NULL DEFAULT '[]',
                    source_field_coverage_json TEXT NOT NULL DEFAULT '[]',
                    deferred_scope_summaries_json TEXT NOT NULL DEFAULT '[]',
                    created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                    FOREIGN KEY(source_import_batch_id) REFERENCES import_batch(id) ON DELETE CASCADE
                )
                """
            )

    def _promotion_lookup(self) -> dict[tuple[str, str, int], dict[str, str | None]]:
        bundle = self._validation_bundle()
        lookup: dict[tuple[str, str, int], dict[str, str | None]] = {}
        for record in bundle.records:
            basic_udi_di = next(
                (
                    field.value
                    for field in record.fields
                    if field.canonical_path in {"basic_device.basic_udi_di", "device_record.basic_udi_identifier"}
                    and field.value
                ),
                None,
            )
            canonical_status = "xml_ready" if record.xml_readiness.status == "complete" else "xml_blocked"
            lookup[(record.source_workbook, record.source_sheet, record.source_row_index)] = {
                "product_family": record.product_family,
                "product_variant": record.product_variant,
                "catalogue_number": record.catalogue_number,
                "primary_udi_di": record.primary_udi_di,
                "submission_operation": record.submission_operation,
                "basic_udi_di": basic_udi_di,
                "canonical_status": canonical_status,
            }
        return lookup

    def _validation_bundle(self) -> CanonicalValidationBundle:
        return self.validation_service.build_validation_bundle()

    def _load_source_rows(self, workbook_path: Path) -> list[ImportedSourceRow]:
        workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        rows: list[ImportedSourceRow] = []
        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            header_row_index: int | None = None
            headers: list[str] | None = None
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = list(row)
                if headers is None and any(
                    isinstance(value, str) and HEADER_SENTINEL in value for value in values if value
                ):
                    headers = [str(value).strip() if value is not None else "" for value in values]
                    header_row_index = row_index
                    continue
                if headers is None or header_row_index is None or row_index <= header_row_index:
                    continue
                if not any(value not in (None, "") for value in values):
                    continue
                mapped_values: dict[str, object | None] = {
                    headers[index]: values[index] if index < len(values) else None
                    for index in range(len(headers))
                    if headers[index]
                }
                rows.append(ImportedSourceRow(sheet_name=sheet_name, row_index=row_index, values=mapped_values))
        return rows

    @staticmethod
    def _device_subject_id_for_source_row(connection: sqlite3.Connection, *, source_row_id: int) -> int | None:
        row = connection.execute(
            """
            SELECT id
            FROM device_subject
            WHERE current_source_row_id = ?
            LIMIT 1
            """,
            (source_row_id,),
        ).fetchone()
        return int(row["id"]) if row is not None else None

    @staticmethod
    def _source_row_lookup_for_batch(
        connection: sqlite3.Connection,
        *,
        import_batch_id: int,
    ) -> dict[tuple[str, str, int], sqlite3.Row]:
        rows = connection.execute(
            """
            SELECT sr.id, sw.workbook_name, sr.sheet_name, sr.row_index
            FROM source_row sr
            JOIN source_workbook sw ON sw.id = sr.source_workbook_id
            WHERE sw.import_batch_id = ?
            """,
            (import_batch_id,),
        ).fetchall()
        return {
            (str(row["workbook_name"]), str(row["sheet_name"]), int(row["row_index"])): row
            for row in rows
        }

    def _match_device_subject(
        self,
        *,
        connection: sqlite3.Connection,
        product_family: str | None,
        product_variant: str | None,
        catalogue_number: str | None,
        primary_udi_di: str | None,
        basic_udi_di: str | None,
        current_source_row_id: int,
        updated_at: str,
    ) -> DeviceSubjectMatchDecision:
        normalized_primary = self._normalize_identity(primary_udi_di)
        fallback_reliable = self._fallback_identity_is_complete(
            product_family=product_family,
            product_variant=product_variant,
            catalogue_number=catalogue_number,
        )
        primary_matches = self._find_device_subjects_by_primary(connection, primary_udi_di=primary_udi_di)
        fallback_matches = self._find_device_subjects_by_fallback(
            connection,
            product_family=product_family,
            product_variant=product_variant,
            catalogue_number=catalogue_number,
        )

        if normalized_primary:
            if len(primary_matches) == 1:
                primary_match = primary_matches[0]
                conflicting_fallbacks = [subject for subject in fallback_matches if subject.id != primary_match.id]
                if conflicting_fallbacks:
                    self._record_identity_issue(
                        connection,
                        source_row_id=current_source_row_id,
                        device_subject_id=primary_match.id,
                        issue_code="identifier_conflict",
                        details={
                            "match_strategy": "primary_udi_di",
                            "conflicting_fallback_subject_ids": [subject.id for subject in conflicting_fallbacks],
                            "primary_udi_di": primary_udi_di,
                            "product_family": product_family,
                            "product_variant": product_variant,
                            "catalogue_number": catalogue_number,
                        },
                        created_at=updated_at,
                    )
                    return DeviceSubjectMatchDecision(created=False, issue_code="identifier_conflict")
                self._save_device_subject(
                    connection,
                    existing_subject=primary_match,
                    product_family=product_family,
                    product_variant=product_variant,
                    catalogue_number=catalogue_number,
                    primary_udi_di=primary_udi_di,
                    basic_udi_di=basic_udi_di,
                    current_source_row_id=current_source_row_id,
                    updated_at=updated_at,
                )
                if self._has_identity_label_drift(
                    primary_match,
                    product_family=product_family,
                    product_variant=product_variant,
                    catalogue_number=catalogue_number,
                ):
                    self._record_identity_issue(
                        connection,
                        source_row_id=current_source_row_id,
                        device_subject_id=primary_match.id,
                        issue_code="identity_label_drift",
                        details={
                            "match_strategy": "primary_udi_di",
                            "previous_product_family": primary_match.product_family,
                            "previous_product_variant": primary_match.product_variant,
                            "previous_catalogue_number": primary_match.catalogue_number,
                            "next_product_family": product_family,
                            "next_product_variant": product_variant,
                            "next_catalogue_number": catalogue_number,
                            "primary_udi_di": primary_udi_di,
                        },
                        created_at=updated_at,
                    )
                    return DeviceSubjectMatchDecision(created=False, issue_code="identity_label_drift")
                return DeviceSubjectMatchDecision(created=False)
            if len(primary_matches) > 1:
                self._record_identity_issue(
                    connection,
                    source_row_id=current_source_row_id,
                    device_subject_id=None,
                    issue_code="identifier_conflict",
                    details={
                        "reason": "multiple_subjects_share_primary_udi_di",
                        "primary_udi_di": primary_udi_di,
                        "matching_subject_ids": [subject.id for subject in primary_matches],
                    },
                    created_at=updated_at,
                )
                return DeviceSubjectMatchDecision(created=False, issue_code="identifier_conflict")
            if len(fallback_matches) == 1:
                fallback_match = fallback_matches[0]
                existing_primary = self._normalize_identity(fallback_match.primary_udi_di)
                if existing_primary and existing_primary != normalized_primary:
                    self._record_identity_issue(
                        connection,
                        source_row_id=current_source_row_id,
                        device_subject_id=fallback_match.id,
                        issue_code="identifier_conflict",
                        details={
                            "match_strategy": "fallback_tuple",
                            "existing_primary_udi_di": fallback_match.primary_udi_di,
                            "incoming_primary_udi_di": primary_udi_di,
                            "product_family": product_family,
                            "product_variant": product_variant,
                            "catalogue_number": catalogue_number,
                        },
                        created_at=updated_at,
                    )
                    return DeviceSubjectMatchDecision(created=False, issue_code="identifier_conflict")
                self._save_device_subject(
                    connection,
                    existing_subject=fallback_match,
                    product_family=product_family,
                    product_variant=product_variant,
                    catalogue_number=catalogue_number,
                    primary_udi_di=primary_udi_di,
                    basic_udi_di=basic_udi_di,
                    current_source_row_id=current_source_row_id,
                    updated_at=updated_at,
                )
                return DeviceSubjectMatchDecision(created=False)
            if len(fallback_matches) > 1:
                self._record_identity_issue(
                    connection,
                    source_row_id=current_source_row_id,
                    device_subject_id=None,
                    issue_code="ambiguous_fallback_match",
                    details={
                        "primary_udi_di": primary_udi_di,
                        "product_family": product_family,
                        "product_variant": product_variant,
                        "catalogue_number": catalogue_number,
                        "matching_subject_ids": [subject.id for subject in fallback_matches],
                    },
                    created_at=updated_at,
                )
                return DeviceSubjectMatchDecision(created=False, issue_code="ambiguous_fallback_match")
            self._save_device_subject(
                connection,
                existing_subject=None,
                product_family=product_family,
                product_variant=product_variant,
                catalogue_number=catalogue_number,
                primary_udi_di=primary_udi_di,
                basic_udi_di=basic_udi_di,
                current_source_row_id=current_source_row_id,
                updated_at=updated_at,
            )
            return DeviceSubjectMatchDecision(created=True)

        if fallback_reliable:
            if len(fallback_matches) == 1:
                self._save_device_subject(
                    connection,
                    existing_subject=fallback_matches[0],
                    product_family=product_family,
                    product_variant=product_variant,
                    catalogue_number=catalogue_number,
                    primary_udi_di=primary_udi_di,
                    basic_udi_di=basic_udi_di,
                    current_source_row_id=current_source_row_id,
                    updated_at=updated_at,
                )
                return DeviceSubjectMatchDecision(created=False)
            if len(fallback_matches) > 1:
                self._record_identity_issue(
                    connection,
                    source_row_id=current_source_row_id,
                    device_subject_id=None,
                    issue_code="ambiguous_fallback_match",
                    details={
                        "product_family": product_family,
                        "product_variant": product_variant,
                        "catalogue_number": catalogue_number,
                        "matching_subject_ids": [subject.id for subject in fallback_matches],
                    },
                    created_at=updated_at,
                )
                return DeviceSubjectMatchDecision(created=False, issue_code="ambiguous_fallback_match")
            self._save_device_subject(
                connection,
                existing_subject=None,
                product_family=product_family,
                product_variant=product_variant,
                catalogue_number=catalogue_number,
                primary_udi_di=primary_udi_di,
                basic_udi_di=basic_udi_di,
                current_source_row_id=current_source_row_id,
                updated_at=updated_at,
            )
            return DeviceSubjectMatchDecision(created=True)

        self._record_identity_issue(
            connection,
            source_row_id=current_source_row_id,
            device_subject_id=None,
            issue_code="insufficient_identity_data",
            details={
                "product_family": product_family,
                "product_variant": product_variant,
                "catalogue_number": catalogue_number,
                "primary_udi_di": primary_udi_di,
            },
            created_at=updated_at,
        )
        return DeviceSubjectMatchDecision(created=False, issue_code="insufficient_identity_data")

    def _save_device_subject(
        self,
        connection: sqlite3.Connection,
        *,
        existing_subject: DeviceSubjectRecord | None,
        product_family: str | None,
        product_variant: str | None,
        catalogue_number: str | None,
        primary_udi_di: str | None,
        basic_udi_di: str | None,
        current_source_row_id: int,
        updated_at: str,
    ) -> None:
        subject_key = self._subject_key(
            product_family=product_family,
            product_variant=product_variant,
            catalogue_number=catalogue_number,
            primary_udi_di=primary_udi_di,
        ) or (existing_subject.subject_key if existing_subject is not None else None)
        if subject_key is None:
            return
        if existing_subject is None:
            connection.execute(
                """
                INSERT INTO device_subject (
                    subject_key,
                    product_family,
                    product_variant,
                    catalogue_number,
                    primary_udi_di,
                    basic_udi_di,
                    current_source_row_id,
                    created_at,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    subject_key,
                    product_family,
                    product_variant,
                    catalogue_number,
                    primary_udi_di,
                    basic_udi_di,
                    current_source_row_id,
                    updated_at,
                    updated_at,
                ),
            )
            return
        connection.execute(
            """
            UPDATE device_subject
            SET subject_key = ?,
                product_family = ?,
                product_variant = ?,
                catalogue_number = ?,
                primary_udi_di = ?,
                basic_udi_di = ?,
                current_source_row_id = ?,
                updated_at = ?
            WHERE id = ?
            """,
            (
                subject_key,
                product_family,
                product_variant,
                catalogue_number,
                primary_udi_di,
                basic_udi_di,
                current_source_row_id,
                updated_at,
                existing_subject.id,
            ),
        )

    def _find_device_subjects_by_primary(
        self,
        connection: sqlite3.Connection,
        *,
        primary_udi_di: str | None,
    ) -> list[DeviceSubjectRecord]:
        normalized_primary = self._normalize_identity(primary_udi_di)
        if not normalized_primary:
            return []
        rows = connection.execute(
            """
            SELECT
                id,
                subject_key,
                product_family,
                product_variant,
                catalogue_number,
                primary_udi_di,
                basic_udi_di,
                current_source_row_id,
                created_at,
                updated_at
            FROM device_subject
            WHERE COALESCE(TRIM(primary_udi_di), '') <> ''
            ORDER BY id
            """
        ).fetchall()
        return [
            self._device_subject_from_row(row)
            for row in rows
            if self._normalize_identity(row["primary_udi_di"]) == normalized_primary
        ]

    def _find_device_subjects_by_fallback(
        self,
        connection: sqlite3.Connection,
        *,
        product_family: str | None,
        product_variant: str | None,
        catalogue_number: str | None,
    ) -> list[DeviceSubjectRecord]:
        if not self._fallback_identity_is_complete(
            product_family=product_family,
            product_variant=product_variant,
            catalogue_number=catalogue_number,
        ):
            return []
        normalized_family = self._normalize_identity(product_family)
        normalized_variant = self._normalize_identity(product_variant)
        normalized_catalogue = self._normalize_identity(catalogue_number)
        rows = connection.execute(
            """
            SELECT
                id,
                subject_key,
                product_family,
                product_variant,
                catalogue_number,
                primary_udi_di,
                basic_udi_di,
                current_source_row_id,
                created_at,
                updated_at
            FROM device_subject
            WHERE COALESCE(TRIM(product_family), '') <> ''
              AND COALESCE(TRIM(product_variant), '') <> ''
              AND COALESCE(TRIM(catalogue_number), '') <> ''
            ORDER BY id
            """
        ).fetchall()
        return [
            self._device_subject_from_row(row)
            for row in rows
            if self._normalize_identity(row["product_family"]) == normalized_family
            and self._normalize_identity(row["product_variant"]) == normalized_variant
            and self._normalize_identity(row["catalogue_number"]) == normalized_catalogue
        ]

    def _record_identity_issue(
        self,
        connection: sqlite3.Connection,
        *,
        source_row_id: int,
        device_subject_id: int | None,
        issue_code: str,
        details: dict[str, Any],
        created_at: str,
    ) -> None:
        connection.execute(
            """
            INSERT INTO device_identity_issue (
                source_row_id,
                device_subject_id,
                issue_code,
                severity,
                details_json,
                created_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                source_row_id,
                device_subject_id,
                issue_code,
                "warning",
                json.dumps(details, default=str, sort_keys=True),
                created_at,
            ),
        )

    @classmethod
    def _has_identity_label_drift(
        cls,
        existing_subject: DeviceSubjectRecord,
        *,
        product_family: str | None,
        product_variant: str | None,
        catalogue_number: str | None,
    ) -> bool:
        return any(
            cls._normalize_identity(current) != cls._normalize_identity(next_value)
            for current, next_value in (
                (existing_subject.product_family, product_family),
                (existing_subject.product_variant, product_variant),
                (existing_subject.catalogue_number, catalogue_number),
            )
        )

    @classmethod
    def _fallback_identity_is_complete(
        cls,
        *,
        product_family: str | None,
        product_variant: str | None,
        catalogue_number: str | None,
    ) -> bool:
        return all(
            cls._normalize_identity(value)
            for value in (product_family, product_variant, catalogue_number)
        )

    @staticmethod
    def _file_hash(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    @staticmethod
    def _optional_string(value: object) -> str | None:
        if value is None:
            return None
        if not isinstance(value, str):
            return str(value)
        normalized = value.strip()
        return normalized or None

    @classmethod
    def _normalize_identity(cls, value: object) -> str:
        text = cls._optional_string(value)
        if not text:
            return ""
        return "".join(text.casefold().split())

    @classmethod
    def _subject_key(
        cls,
        *,
        product_family: str | None,
        product_variant: str | None,
        catalogue_number: str | None,
        primary_udi_di: str | None,
    ) -> str | None:
        primary = cls._normalize_identity(primary_udi_di)
        if primary:
            return f"primary:{primary}"
        family = cls._normalize_identity(product_family)
        variant = cls._normalize_identity(product_variant)
        catalogue = cls._normalize_identity(catalogue_number)
        if not family or not variant or not catalogue:
            return None
        return f"fallback:{family}|{variant}|{catalogue}"

    @classmethod
    def _device_subject_from_row(cls, row: sqlite3.Row) -> DeviceSubjectRecord:
        return DeviceSubjectRecord(
            id=int(row["id"]),
            subject_key=str(row["subject_key"]),
            product_family=cls._optional_string(row["product_family"]),
            product_variant=cls._optional_string(row["product_variant"]),
            catalogue_number=cls._optional_string(row["catalogue_number"]),
            primary_udi_di=cls._optional_string(row["primary_udi_di"]),
            basic_udi_di=cls._optional_string(row["basic_udi_di"]),
            current_source_row_id=int(row["current_source_row_id"]) if row["current_source_row_id"] is not None else None,
            created_at=str(row["created_at"]),
            updated_at=str(row["updated_at"]),
        )

    @staticmethod
    def _batch_summary_from_row(row: sqlite3.Row) -> WorkbookImportBatchSummary:
        return WorkbookImportBatchSummary(
            import_batch_id=int(row["id"]),
            source_type=str(row["source_type"]),
            label=str(row["label"]),
            imported_at=str(row["imported_at"]),
            imported_by=WorkbookImportService._optional_string(row["imported_by"]),
            notes=WorkbookImportService._optional_string(row["notes"]),
            workbook_count=int(row["workbook_count"]),
            source_row_count=int(row["source_row_count"]),
            device_subject_count=int(row["device_subject_count"]),
        )

    @staticmethod
    def _split_csv_values(value: object) -> list[str]:
        if value is None:
            return []
        return [item for item in (str(value).split(",") if value else []) if item]

    @staticmethod
    def _split_csv_ints(value: object) -> list[int]:
        if value is None:
            return []
        items: list[int] = []
        for part in str(value).split(","):
            part = part.strip()
            if not part:
                continue
            try:
                items.append(int(part))
            except ValueError:
                continue
        return items

    @staticmethod
    def _json_object(value: object) -> dict[str, Any]:
        if value is None:
            return {}
        if isinstance(value, dict):
            return value
        try:
            parsed = json.loads(str(value))
        except (TypeError, ValueError):
            return {}
        return parsed if isinstance(parsed, dict) else {}

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.db_path)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")
        return connection

    def _create_import_backup(self, *, import_batch_id: int) -> Path:
        backup_dir = self.settings.testing_state_backup_dir
        backup_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now(UTC).strftime("%Y%m%dT%H%M%SZ")
        backup_path = backup_dir / f"testing-state-batch-{import_batch_id}-{timestamp}.sqlite3"
        shutil.copy2(self.db_path, backup_path)
        self._prune_import_backups(backup_dir)
        return backup_path

    def _prune_import_backups(self, backup_dir: Path) -> None:
        keep_count = max(self.settings.testing_state_backup_keep_count, 1)
        backups = sorted(
            backup_dir.glob("testing-state-batch-*.sqlite3"),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )
        for stale_backup in backups[keep_count:]:
            stale_backup.unlink(missing_ok=True)

    @staticmethod
    def _table_names(connection: sqlite3.Connection) -> list[str]:
        rows = connection.execute(
            """
            SELECT name
            FROM sqlite_master
            WHERE type = 'table'
              AND name NOT LIKE 'sqlite_%'
            ORDER BY name
            """
        ).fetchall()
        return [str(row["name"]) for row in rows]

    @classmethod
    def _safe_row_count(cls, connection: sqlite3.Connection, table_name: str) -> int:
        return cls._safe_scalar(connection, f"SELECT COUNT(*) FROM {table_name}", table_name)

    @staticmethod
    def _safe_scalar(connection: sqlite3.Connection, query: str, table_name: str) -> int:
        try:
            row = connection.execute(query).fetchone()
        except sqlite3.OperationalError:
            return 0
        if row is None:
            return 0
        return int(row[0])

    def _columns_for_table(self, connection: sqlite3.Connection, table_name: str) -> list[DatabaseColumnSummary]:
        rows = connection.execute(f"PRAGMA table_info('{table_name}')").fetchall()
        return [
            DatabaseColumnSummary(
                name=str(row["name"]),
                data_type=str(row["type"] or ""),
                nullable=not bool(row["notnull"]),
                primary_key_position=int(row["pk"]),
            )
            for row in rows
        ]

    def _foreign_keys_for_table(self, connection: sqlite3.Connection, table_name: str) -> list[DatabaseForeignKeySummary]:
        rows = connection.execute(f"PRAGMA foreign_key_list('{table_name}')").fetchall()
        return [
            DatabaseForeignKeySummary(
                from_column=str(row["from"]),
                target_table=str(row["table"]),
                target_column=str(row["to"]),
                on_delete=str(row["on_delete"]),
            )
            for row in rows
        ]

    def _indexes_for_table(self, connection: sqlite3.Connection, table_name: str) -> list[DatabaseIndexSummary]:
        rows = connection.execute(f"PRAGMA index_list('{table_name}')").fetchall()
        indexes: list[DatabaseIndexSummary] = []
        for row in rows:
            index_name = str(row["name"])
            column_rows = connection.execute(f"PRAGMA index_info('{index_name}')").fetchall()
            indexes.append(
                DatabaseIndexSummary(
                    name=index_name,
                    unique=bool(row["unique"]),
                    columns=[str(column_row["name"]) for column_row in column_rows if column_row["name"]],
                )
            )
        return indexes

    def _workbook_diffs(
        self,
        connection: sqlite3.Connection,
        *,
        current_import_batch_id: int,
        previous_import_batch_id: int | None,
    ) -> list[WorkbookImportWorkbookDiff]:
        current_rows = connection.execute(
            """
            SELECT workbook_name, file_hash, COUNT(source_row.id) AS row_count
            FROM source_workbook workbook
            LEFT JOIN source_row ON source_row.source_workbook_id = workbook.id
            WHERE workbook.import_batch_id = ?
            GROUP BY workbook.id
            """,
            (current_import_batch_id,),
        ).fetchall()
        previous_lookup: dict[str, sqlite3.Row] = {}
        if previous_import_batch_id is not None:
            previous_rows = connection.execute(
                """
                SELECT workbook_name, file_hash, COUNT(source_row.id) AS row_count
                FROM source_workbook workbook
                LEFT JOIN source_row ON source_row.source_workbook_id = workbook.id
                WHERE workbook.import_batch_id = ?
                GROUP BY workbook.id
                """,
                (previous_import_batch_id,),
            ).fetchall()
            previous_lookup = {str(row["workbook_name"]): row for row in previous_rows}

        diffs: list[WorkbookImportWorkbookDiff] = []
        current_names = {str(row["workbook_name"]) for row in current_rows}
        for row in current_rows:
            workbook_name = str(row["workbook_name"])
            previous_row = previous_lookup.pop(workbook_name, None)
            current_hash = self._optional_string(row["file_hash"])
            previous_hash = self._optional_string(previous_row["file_hash"]) if previous_row is not None else None
            current_row_count = int(row["row_count"])
            previous_row_count = int(previous_row["row_count"]) if previous_row is not None else None
            change_type = "unchanged"
            if previous_row is None:
                change_type = "added"
            elif current_hash != previous_hash or current_row_count != previous_row_count:
                change_type = "changed"
            diffs.append(
                WorkbookImportWorkbookDiff(
                    workbook_name=workbook_name,
                    change_type=change_type,
                    previous_row_count=previous_row_count,
                    current_row_count=current_row_count,
                    previous_hash=previous_hash,
                    current_hash=current_hash,
                )
            )
        for workbook_name, previous_row in previous_lookup.items():
            diffs.append(
                WorkbookImportWorkbookDiff(
                    workbook_name=workbook_name,
                    change_type="removed",
                    previous_row_count=int(previous_row["row_count"]),
                    current_row_count=None,
                    previous_hash=self._optional_string(previous_row["file_hash"]),
                    current_hash=None,
                )
            )
        return sorted(
            diffs,
            key=lambda item: (
                {"changed": 0, "added": 1, "removed": 2, "unchanged": 3}.get(item.change_type, 9),
                item.workbook_name,
            ),
        )

    def _device_subject_count_for_batch(self, connection: sqlite3.Connection, import_batch_id: int) -> int:
        row = connection.execute(
            """
            SELECT COUNT(*)
            FROM (
                SELECT ds.subject_key
                FROM device_subject ds
                JOIN source_row sr
                  ON COALESCE(sr.catalogue_number, '') = COALESCE(ds.catalogue_number, '')
                 AND COALESCE(sr.primary_udi_di, '') = COALESCE(ds.primary_udi_di, '')
                 AND COALESCE(sr.product_family, '') = COALESCE(ds.product_family, '')
                 AND COALESCE(sr.product_variant, '') = COALESCE(ds.product_variant, '')
                JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                WHERE sw.import_batch_id = ?
                GROUP BY ds.subject_key
            )
            """,
            (import_batch_id,),
        ).fetchone()
        return int(row[0]) if row is not None else 0

    def _unresolved_identity_row_count_for_batch(self, import_batch_id: int) -> int:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT COUNT(DISTINCT sr.id)
                FROM source_row sr
                JOIN source_workbook sw ON sw.id = sr.source_workbook_id
                JOIN device_identity_issue issue ON issue.source_row_id = sr.id
                LEFT JOIN device_subject ds ON ds.current_source_row_id = sr.id
                WHERE sw.import_batch_id = ?
                  AND ds.id IS NULL
                """,
                (import_batch_id,),
            ).fetchone()
        return int(row[0]) if row is not None else 0
