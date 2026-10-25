"""Read-only comparison of Dev accepted evidence with a fixed Playground export.

No application store constructors, migrations, baseline writes or review receipts.
Unknown local fields never fall back to proposed Template/canonical values.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from contextlib import closing
from datetime import datetime, UTC
import hashlib
import json
from pathlib import Path
import sqlite3
import re
import tempfile
from uuid import uuid4

from lxml import etree
from fastapi.encoders import jsonable_encoder
from openpyxl import Workbook
from openpyxl.styles import Font

from app.config import Settings
from app.services.accepted_state import accepted_device_state, accepted_market_state, accepted_post_state
from app.services.dev_batch_preparation import DevBatchPreparation
from app.services.dev_input_batches import fingerprint
from app.services.production_workbook_adapter import accepted_snapshot
from app.services.production_workbook_preparation import _cell, _exports

PARENT_FIELDS = {
    "basic_identifier_code", "basic_identifier_issuing_entity", "risk_class", "model_name",
    "manufacturer_srn", "authorised_representative_srn", "human_tissues_cells", "animal_tissues_cells",
    "human_product_check", "medicinal_product_check", "basic_device_type", "active",
    "administering_medicine", "implantable", "measuring_function", "reusable",
}
SKIP_FIELDS = {"product_family", "product_variant", "submission_operation", "source_version_marker", "post_payload_mode", "service_id_override", "patch_version_override", "include_market_infos_in_patch",
               "market_countries", "storage_conditions", "critical_warnings", "reference_number"}
VERSION_FIELDS = {"device_version", "parent_version", "market_info_version"}


def normalized(field: str, value):
    if isinstance(value, str):
        value = value.strip()
        if field in VERSION_FIELDS or field in {"base_quantity", "number_of_reuses"}:
            return str(int(value)) if value.isdigit() else value
        return value or None
    if isinstance(value, dict):
        return {key: normalized(key, entry) for key, entry in sorted(value.items()) if entry is not None and entry != ""}
    if isinstance(value, list):
        # Order does not change a country/code set; preserve duplicate occurrences.
        return sorted((normalized(field, entry) for entry in value), key=lambda entry: json.dumps(entry, sort_keys=True))
    if field in VERSION_FIELDS | {"base_quantity", "number_of_reuses"} and value is not None:
        return str(value)
    return value


def timestamp(value):
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed if parsed.tzinfo else None
    except (ValueError, TypeError):
        return None


def positive_version(value):
    try:
        number = int(str(value))
        return str(number) if number > 0 else None
    except (ValueError, TypeError):
        return None


def field_review(field, local, exported, *, known: bool, evidence: str, group="Device"):
    state = "Missing local evidence" if not known else "Matches" if normalized(field, local) == normalized(field, exported) else "Difference found"
    return {"field": field, "group": group, "local": local if known else None,
            "exported": exported, "comparison": state, "local_evidence": evidence if known else "Not available"}


class PlaygroundReconciliation:
    def __init__(self, settings: Settings):
        self.batches = DevBatchPreparation(settings)  # Dev guard is enforced before any DB access.
        self.settings = settings

    def assess(self, batch_id: str) -> dict:
        batch, manifest, _, input_token = self.batches._inputs(batch_id)
        exports, files = _exports(batch / "eudamed-xml", {})
        if not exports:
            raise ValueError("No exported devices to assess")
        if any(count > 1 for count in Counter(tuple(entry["key"]) for entry in exports).values()):
            raise ValueError("Duplicate exported device identity requires source review")
        for path in (batch / "eudamed-xml").glob("*.xml"):
            tree = etree.parse(str(path), etree.XMLParser(resolve_entities=False, no_network=True))
            if tree.getroot().get("version") != "3.0.32" or set(tree.xpath('//*[local-name()="MFActorCode"]/text()')) != {manifest["manufacturer_srn"]}:
                raise ValueError("Export schema/manufacturer does not match the selected Playground batch")
        database = self.settings.testing_state_db_path.resolve()
        if any(database.parts[i:i+2] == ("data", "prod") for i in range(len(database.parts)-1)):
            raise PermissionError("Playground assessment cannot read the Production database")
        if not database.is_file():
            raise ValueError("The existing Dev database is absent; assessment never creates one")
        codes = sorted({entry["key"][1] for entry in exports})
        database_digest = hashlib.sha256()
        records = []
        with closing(sqlite3.connect(database.as_uri() + "?mode=ro", uri=True)) as connection:
            connection.row_factory = sqlite3.Row
            connection.execute("PRAGMA query_only=ON")
            connection.execute("BEGIN")
            placeholders = ",".join("?" for _ in codes)
            identities = defaultdict(list)
            by_subject = {}
            # Proposed canonical data is used ONLY to resolve issuer-qualified identity.
            # It is never used as accepted field evidence.
            canonical_rows = connection.execute(f'''SELECT d.id,d.primary_udi_di,d.basic_udi_di,c.record_json
                FROM device_subject d LEFT JOIN canonical_device_record c ON c.id=(
                    SELECT id FROM canonical_device_record WHERE device_subject_id=d.id
                    ORDER BY canonical_version DESC,id DESC LIMIT 1)
                WHERE d.primary_udi_di IN ({placeholders}) ORDER BY d.id''', codes).fetchall()
            for row in canonical_rows:
                database_digest.update(json.dumps(dict(row), sort_keys=True).encode())
                raw = json.loads(row["record_json"] or "{}")
                fields = {field["canonical_path"]: field.get("value") for field in raw.get("fields", [])}
                issuer = str(fields.get("manufacturer.issuing_entity") or "").strip().upper()
                code = str(fields.get("device_record.primary_udi_di") or "").strip()
                if issuer and code:
                    identity = {"key": (issuer, code), "device_subject_id": row["id"], "basic_udi_di": row["basic_udi_di"]}
                    identities[identity["key"]].append(identity)
                    by_subject[row["id"]] = identity
            local_subjects = connection.execute("SELECT * FROM testing_subjects ORDER BY id").fetchall()
            subjects_by_key = defaultdict(list)
            for subject in local_subjects:
                database_digest.update(json.dumps(dict(subject), sort_keys=True).encode())
                linked = by_subject.get(subject["device_subject_id"])
                if linked is not None:
                    subjects_by_key[linked["key"]].append(subject)
                elif not subject["device_subject_id"]:
                    possible = [key for key in identities if key[1] == str(subject["primary_udi_di"] or "").strip()]
                    if len(possible) == 1:
                        subjects_by_key[possible[0]].append(subject)
            events = defaultdict(list)
            for row in connection.execute("SELECT id,subject_id,message_type,status,version,tested_at FROM testing_events WHERE status='SUCCESS' ORDER BY id"):
                database_digest.update(json.dumps(dict(row), sort_keys=True).encode())
                events[row["subject_id"]].append(dict(row))
            review_count = connection.execute("SELECT count(*) FROM reviewed_post_baselines").fetchone()[0]
            for entry in exports:
                key = tuple(entry["key"])
                subject_options = subjects_by_key.get(key, [])
                canonical_options = identities.get(key, [])
                subject = subject_options[0] if len(subject_options) == 1 and len(canonical_options) == 1 else None
                result = self._compare(connection, entry, subject, canonical_options,
                                       events.get(subject["id"], []) if subject else [])
                records.append(result)
                database_digest.update(json.dumps(result, sort_keys=True).encode())
            connection.rollback()
        if self.batches._inputs(batch_id)[3] != input_token:
            raise ValueError("Batch inputs changed during reconciliation assessment; assess again")
        outcomes = Counter(row["outcome"] for row in records)
        field_counts = Counter(field["comparison"] for row in records for field in row["fields"])
        return {"report_version": 1, "environment": "dev", "batch_id": batch_id,
            "manufacturer_srn": manifest["manufacturer_srn"], "assessed_at": datetime.now(UTC).isoformat(),
            "read_only": True, "database_changed": False, "apply_available": False,
            "input_token": input_token, "database_snapshot_sha256": database_digest.hexdigest(),
            "database_snapshot_scope": "Compared identities, testing subjects, successful event metadata and projected comparison evidence; not a whole-file hash.",
            "export_files": files, "records": records,
            "summary": {"devices": len(records), "parents": len({tuple(row["parent"]) for row in exports}),
                "outcomes": dict(outcomes), "field_comparisons": dict(field_counts),
                "full_local_snapshots": sum(row["local_snapshot"] == "Full" for row in records),
                "partial_local_snapshots": sum(row["local_snapshot"] == "Partial" for row in records),
                "missing_local_snapshots": sum(row["local_snapshot"] == "Missing" for row in records)},
            "review_markers": {"current_count": review_count, "recreated": False,
                "note": "Exports cannot establish operator review. Cleared legacy review entries require fresh review."},
            "limitations": ["Comparison covers the existing accepted projection; it is not full raw-XML equality or XSD validation.",
                "Missing local evidence is not a proven changed value. Proposed Templates are never substituted for accepted state.",
                "These dated exports cannot rule out later changes made externally in Playground.",
                "No baseline, historical submission, source import, proposed edit or review approval has been changed."]}

    def create_report(self, batch_id: str) -> dict:
        report = self.assess(batch_id)
        batch = self.batches._batch(batch_id)
        reports = batch / "reconciliation"
        reports.mkdir(exist_ok=True)
        if reports.is_symlink():
            raise ValueError("Reconciliation output must not be a symlink")
        report_id = uuid4().hex
        report["report_id"] = report_id
        with tempfile.TemporaryDirectory(prefix=".assessing-", dir=reports) as temporary:
            output = Path(temporary) / "output"
            output.mkdir()
            json_path = output / "playground-baseline-assessment.json"
            excel_path = output / "playground-baseline-assessment.xlsx"
            json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
            write_workbook(report, excel_path)
            metadata = {"report_id": report_id, "batch_id": batch_id,
                        "sha256": {"json": fingerprint(json_path), "xlsx": fingerprint(excel_path)}}
            (output / "metadata.json").write_text(json.dumps(metadata, indent=2) + "\n")
            output.rename(reports / report_id)
        # Keep raw XML and proposed snapshots in the audit download, outside the UI response.
        return {**report, "records": [{key: value for key, value in row.items() if key != "proposed_snapshot"} for row in report["records"]]}

    def artifact(self, batch_id: str, report_id: str, kind: str) -> Path:
        if not re.fullmatch(r"[a-f0-9]{32}", report_id) or kind not in {"json", "xlsx"}:
            raise ValueError("Invalid reconciliation download")
        batch = self.batches._batch(batch_id)
        root = batch / "reconciliation"
        directory = root / report_id
        if root.is_symlink() or directory.is_symlink() or directory.resolve().parent != root.resolve():
            raise ValueError("Reconciliation download path escapes Dev storage")
        metadata_path = directory / "metadata.json"
        path = directory / f"playground-baseline-assessment.{kind}"
        if path.is_symlink() or metadata_path.is_symlink():
            raise ValueError("Reconciliation evidence must not be a symlink")
        metadata = json.loads(metadata_path.read_text())
        if metadata["report_id"] != report_id or metadata["batch_id"] != batch_id or fingerprint(path) != metadata["sha256"][kind]:
            raise ValueError("Reconciliation report fingerprint does not match")
        return path

    def _compare(self, connection, entry, subject, canonical, events):
        record = {"issuer": entry["key"][0], "udi_di": entry["key"][1], "parent": list(entry["parent"]),
            "model": entry["device_type"], "catalogue": "", "device_subject_id": canonical[0]["device_subject_id"] if len(canonical) == 1 else None,
            "testing_subject_id": subject["id"] if subject else None, "export_source": entry["source"],
            "outcome": "Update proposed", "warnings": [], "fields": [], "local_snapshot": "Missing"}
        try:
            exported = jsonable_encoder(accepted_snapshot(entry, self.settings))
        except ValueError as exc:
            return {**record, "outcome": "Unable to compare", "warnings": [str(exc)]}
        record["catalogue"] = exported["device_record"]["catalogue_number"]
        record["proposed_snapshot"] = exported
        warnings = record["warnings"]
        requires_review = False
        if subject is None:
            warnings.append("No unique local issuer-qualified device/testing identity; resolve identity before any baseline update.")
            requires_review = True
        elif not subject["device_subject_id"]:
            warnings.append("Legacy testing subject resolved through its unique canonical identity; relational link needs review.")
            requires_review = True
        elif subject["primary_udi_di"] and subject["primary_udi_di"] != entry["key"][1]:
            warnings.append("Linked testing identity conflicts with the canonical Device UDI-DI.")
            requires_review = True
        local = accepted_device_state(connection, subject) if subject else None
        post = accepted_post_state(connection, subject) if subject else None
        market = accepted_market_state(connection, subject) if subject else None
        record["local_snapshot"] = "Full" if local and local.get("device_record") else "Partial" if local else "Missing"
        if exported["unsupported_projection"]:
            warnings.extend(exported["unsupported_projection"]); requires_review = True
        latest = timestamp(entry["source"]["created_at"])
        if latest is None:
            warnings.append("Export timestamp lacks a usable timezone; freshness requires review."); requires_review = True
        for event in events:
            when = timestamp(event["tested_at"])
            if when is None or latest is None or when > latest:
                warnings.append(f"Successful event #{event['id']} is newer than the export or has unresolved timing; do not overwrite it.")
                requires_review = True
        known_versions = []
        if subject:
            known_versions = [positive_version(subject[name]) for name in ("latest_successful_post_version", "latest_successful_patch_version")]
        known_versions += [positive_version(event["version"]) for event in events if event["message_type"] in {"DEVICE.POST", "UDI_DI.POST", "UDI_DI.PATCH"}]
        known_versions = [value for value in known_versions if value is not None]
        device_version = max(known_versions, key=int) if known_versions else None
        export_baseline = bool(local and local.get('accepted_state_source') == 'playground_export')
        if export_baseline:
            device_version = positive_version(local.get('version'))
        version = field_review("device_version", device_version, exported["version"], known=device_version is not None, evidence="Reviewed Playground export baseline" if export_baseline else "Qualified POST/PATCH version evidence")
        record["fields"].append(version)
        if device_version and int(device_version) > int(exported["version"]):
            warnings.append("Known local device version exceeds the export; no automatic downgrade is proposed."); requires_review = True
        if not device_version and subject and subject["latest_successful_version"]:
            record["legacy_version_observation"] = subject["latest_successful_version"]
            warnings.append(f"Legacy generic version {subject['latest_successful_version']} is not qualified as a device version; review its meaning.")
            requires_review = True
        elif not export_baseline and device_version and subject and subject["latest_successful_version"] and str(subject["latest_successful_version"]) != device_version:
            record["legacy_version_observation"] = subject["latest_successful_version"]
            warnings.append("Generic summary version differs from qualified device evidence; review its meaning.")
            requires_review = True
        parent_version = local.get("parent_version") if local and "parent_version" in local else post.get("parent_version") if post else None
        record["fields"].append(field_review("parent_version", parent_version, exported["parent_version"], known=parent_version is not None, evidence="Accepted parent snapshot", group="Parent"))
        market_version = subject["latest_successful_market_info_version"] if subject else None
        market_source = "Dedicated accepted Market Info version"
        if not market_version and market:
            market_version = market.get("version"); market_source = "Existing accepted-state reader (may derive initial version from POST)"
        record["fields"].append(field_review("market_info_version", market_version, exported["market_info_version"], known=market_version is not None, evidence=market_source, group="Market Information"))
        if market_version and positive_version(market_version) and int(market_version) > int(exported["market_info_version"]):
            warnings.append("Known local Market Info version exceeds the export; review before updating."); requires_review = True
        if subject and positive_version(subject["latest_observed_market_info_version"]) and int(subject["latest_observed_market_info_version"]) > int(exported["market_info_version"]):
            warnings.append("Observed local Market Info version floor exceeds the export."); requires_review = True
        full = local.get("device_record", {}) if local else {}
        for name, value in exported["device_record"].items():
            if name in SKIP_FIELDS:
                continue
            known = name in full or bool(local and name in local)
            before = full[name] if name in full else local.get(name) if local else None
            record["fields"].append(field_review(name, before, value, known=known,
                evidence="Accepted full device snapshot" if name in full else "Accepted partial snapshot",
                group="Parent" if name in PARENT_FIELDS else "Device"))
        for name in ("storage_conditions", "critical_warnings"):
            known = bool(local and name in local) or name in full
            before = local[name] if local and name in local else full.get(name)
            record["fields"].append(field_review(name, before, exported[name], known=known, evidence="Accepted device snapshot"))
        record["fields"].append(field_review("market_countries", market.get("market_countries") if market else None,
            exported["market_countries"], known=bool(market and "market_countries" in market), evidence="Existing accepted Market Info snapshot", group="Market Information"))
        counts = Counter(field["comparison"] for field in record["fields"])
        record["counts"] = dict(counts)
        record["outcome"] = "Review required" if requires_review else "Update proposed" if counts["Difference found"] or counts["Missing local evidence"] else "Matches"
        return record


def write_workbook(report: dict, path: Path) -> None:
    """Business review, with all strings kept literal (including Excel formulas)."""
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    sheets = {
        "Summary": [["Playground baseline assessment", "Read only — no database changes"],
            ["Batch", report["batch_id"]], ["Assessed at", report["assessed_at"]],
            ["Manufacturer", report["manufacturer_srn"]], ["Devices", report["summary"]["devices"]],
            ["Parents", report["summary"]["parents"]],
            *[[key, value] for key, value in report["summary"]["outcomes"].items()],
            *[[key, value] for key, value in report["summary"]["field_comparisons"].items()],
            ["Operator review", report["review_markers"]["note"]],
            *[["Limitation", value] for value in report["limitations"]]],
        "Device Review": [["Issuer", "UDI-DI", "Catalogue", "Model", "Outcome", "Local snapshot", "Differences", "Unknown fields", "Review notes", "Export file", "Export time"]],
        "Field Review": [["Issuer", "UDI-DI", "Catalogue", "Group", "Field", "Comparison", "Local accepted evidence", "Playground export", "Local evidence source"]],
        "Export Evidence": [["File", "SHA256", "Created at", "Page", "Number of pages", "Devices"]],
    }
    for row in report["records"]:
        sheets["Device Review"].append([row["issuer"], row["udi_di"], row["catalogue"], row["model"], row["outcome"], row["local_snapshot"],
            row.get("counts", {}).get("Difference found", 0), row.get("counts", {}).get("Missing local evidence", 0),
            "\n".join(row["warnings"]), row["export_source"]["file"], row["export_source"]["created_at"]])
        for field in row["fields"]:
            sheets["Field Review"].append([row["issuer"], row["udi_di"], row["catalogue"], field["group"], field["field"], field["comparison"],
                display_value(field["local"]), display_value(field["exported"]), field["local_evidence"]])
    for entry in report["export_files"]:
        sheets["Export Evidence"].append([entry[key] for key in ("file", "sha256", "created_at", "page", "number_of_pages", "devices")])
    for name, rows in sheets.items():
        sheet = summary if name == "Summary" else book.create_sheet(name)
        for row in rows:
            sheet.append([_cell(value) for value in row])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(65, max(18, max(len(str(c.value or "")) for c in column) + 2))
    book.save(path)
    book.close()


def display_value(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True) if isinstance(value, (dict, list)) else value
