from __future__ import annotations

from collections import Counter
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

from app.config import get_settings
from app.models import (
    ColumnProfile,
    DistinctValueItem,
    DistinctValueProfile,
    FileInventoryItem,
    SheetProfile,
    SheetSummary,
    WorkbookSummary,
)
from app.services.normalization import NormalizationRepository


class ExcelProfiler:
    def __init__(self) -> None:
        self.settings = get_settings()

    def list_workbooks(self) -> list[WorkbookSummary]:
        items: list[WorkbookSummary] = []
        for path in self._workbook_paths():
            workbook = self._load_workbook(path)
            sheet_summaries = self._sheet_summaries(path, workbook)
            items.append(
                WorkbookSummary(
                    workbook=path.name,
                    sheet_count=len(workbook.sheetnames),
                    total_rows=sum(sheet.data_rows for sheet in sheet_summaries),
                    total_columns=max((sheet.max_columns for sheet in sheet_summaries), default=0),
                    sheets=workbook.sheetnames,
                )
            )
        return items

    def list_sheets(self) -> list[SheetSummary]:
        summaries: list[SheetSummary] = []
        for path in self._workbook_paths():
            workbook = self._load_workbook(path)
            summaries.extend(self._sheet_summaries(path, workbook))
        return summaries

    def get_sheet_profile(self, workbook_name: str, sheet_name: str) -> SheetProfile | None:
        path = self.settings.excel_dir / workbook_name
        if not path.exists():
            return None
        workbook = self._load_workbook(path)
        if sheet_name not in workbook.sheetnames:
            return None
        ws = workbook[sheet_name]
        header_row_index, rows = self._extract_rows(ws)
        headers = self._build_headers(
            rows[header_row_index - 1] if header_row_index else [],
            sheet_name=sheet_name,
        )
        data_rows = rows[header_row_index:] if header_row_index else []
        columns: list[ColumnProfile] = []
        for idx, header in enumerate(headers, start=1):
            values = [self._stringify(row[idx - 1]) for row in data_rows if idx - 1 < len(row)]
            non_null_values = [value for value in values if value]
            columns.append(
                ColumnProfile(
                    index=idx,
                    header=header,
                    non_null_count=len(non_null_values),
                    null_count=len(data_rows) - len(non_null_values),
                    distinct_count=len(set(non_null_values)),
                    sample_values=sorted(set(non_null_values))[:5],
                )
            )
        return SheetProfile(
            workbook=workbook_name,
            sheet=sheet_name,
            header_row=header_row_index,
            data_rows=len(data_rows),
            columns=columns,
        )

    def get_distinct_values(
        self,
        *,
        column: str,
        workbook: str | None,
        sheet: str | None,
    ) -> DistinctValueProfile:
        counter: Counter[str] = Counter()
        for workbook_name, sheet_name, headers, data_rows in self._iter_sheet_data():
            if workbook and workbook_name != workbook:
                continue
            if sheet and sheet_name != sheet:
                continue
            if column not in headers:
                continue
            index = headers.index(column)
            for row in data_rows:
                if index < len(row):
                    value = self._stringify(row[index])
                    if value:
                        counter[value] += 1

        normalization_repository = NormalizationRepository()
        rule_lookup = normalization_repository.rules_for_column(column)
        accepted_values = normalization_repository.accepted_values_for_column(column)
        values = [
            DistinctValueItem(
                raw_value=raw,
                count=count,
                normalized_value=rule_lookup.get(raw, raw if raw in accepted_values else None),
                status="mapped" if raw in rule_lookup or raw in accepted_values else "unmapped",
            )
            for raw, count in counter.most_common()
        ]
        return DistinctValueProfile(workbook=workbook, sheet=sheet, column=column, values=values)

    def _sheet_summaries(self, path: Path, workbook) -> list[SheetSummary]:
        summaries: list[SheetSummary] = []
        for ws in workbook.worksheets:
            header_row, rows = self._extract_rows(ws)
            headers = self._build_headers(
                rows[header_row - 1] if header_row else [],
                sheet_name=ws.title,
            )
            data_rows = rows[header_row:] if header_row else []
            populated_columns = sum(1 for header in headers if header)
            summaries.append(
                SheetSummary(
                    workbook=path.name,
                    sheet=ws.title,
                    header_row=header_row,
                    data_rows=len(data_rows),
                    populated_columns=populated_columns,
                    max_columns=ws.max_column,
                    header_labels=headers,
                )
            )
        return summaries

    def _iter_sheet_data(self):
        for path in self._workbook_paths():
            workbook = self._load_workbook(path)
            for ws in workbook.worksheets:
                header_row, rows = self._extract_rows(ws)
                if not header_row:
                    continue
                headers = self._build_headers(rows[header_row - 1], sheet_name=ws.title)
                yield path.name, ws.title, headers, rows[header_row:]

    @staticmethod
    def _build_headers(header_row: tuple, *, sheet_name: str) -> list[str]:
        headers: list[str] = []
        for idx, value in enumerate(header_row, start=1):
            label = ExcelProfiler._stringify(value)
            headers.append(label or f"Unlabeled column ({sheet_name}, column {idx})")
        return headers

    @staticmethod
    def _extract_rows(ws) -> tuple[int | None, list[tuple]]:
        rows = list(ws.iter_rows(values_only=True))
        for row_index, row in enumerate(rows, start=1):
            if any(isinstance(cell, str) and "Issuing Entity" in cell for cell in row if cell):
                return row_index, rows
        return None, rows

    @staticmethod
    def _stringify(value) -> str:
        if value is None:
            return ""
        return " ".join(str(value).split())

    @staticmethod
    @lru_cache(maxsize=32)
    def _load_workbook(path: Path):
        return load_workbook(path, read_only=True, data_only=True)

    def _workbook_paths(self) -> list[Path]:
        return sorted(self.settings.excel_dir.glob("*.xlsx"))
