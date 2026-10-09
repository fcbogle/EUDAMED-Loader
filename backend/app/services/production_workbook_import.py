"""Profile-bound, assessed Production imports into the existing SQLite schema."""
from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import UTC, datetime
from functools import lru_cache
import hashlib
from io import BytesIO
import json
from pathlib import Path
import secrets
import sqlite3
from threading import Lock
import time

from openpyxl import load_workbook

from app.config import get_settings
from app.services.canonical_validation import CanonicalValidationService
from app.services.production_import_files import current_pair_lock
from app.services.production_workbook_adapter import accepted_snapshot, canonical_record
from app.services.production_workbook_preparation import CONTROLS
from app.services.testing_state_store import TestingStateStore
from app.services.workbook_import import WorkbookImportService


def dump(value) -> str:
    return json.dumps(value, sort_keys=True, ensure_ascii=False, default=str)


def sha(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def semantic(value):
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, dict):
        return {k: semantic(v) for k,v in value.items() if v is not None and v != ""}
    if isinstance(value, list):
        return [semantic(v) for v in value]
    return value


@dataclass
class PreparedImport:
    workbook_hash: str
    audit_hash: str
    audit: dict
    rows: list[dict]
    skipped: list[dict]
    database_marker: str
    report: dict
    expires_at: float
    configuration: tuple[str, ...]


class ProductionWorkbookImporter:
    _pending: dict[str, PreparedImport] = {}
    _mutex = Lock()
    assessment_seconds = 900

    def __init__(self):
        # Assessment does not instantiate a service which creates tables/directories.
        self.settings = get_settings()
        if self.settings.environment != 'prod':
            raise ValueError('Production workbook import is only available in Prod.')
        self.workbook = self.settings.production_import_workbook
        self.audit_path = self.settings.production_import_audit
        if self.workbook is None or self.audit_path is None:
            raise ValueError('Configure the Production workbook and audit paths.')
        if self.workbook.parent != self.audit_path.parent:
            raise ValueError('Production workbook and audit must share a directory.')
        if self.workbook.resolve() == self.audit_path.resolve():
            raise ValueError('Workbook and audit paths must differ.')

    def _pair(self):
        if (self.workbook.parent / '.production-import.lock').exists() or self.workbook.with_name('~$'+self.workbook.name).exists():
            raise ValueError('Production workbook is locked; close Excel or finish preparation and retry.')
        try:
            content, audit_bytes = self.workbook.read_bytes(), self.audit_path.read_bytes()
            audit = json.loads(audit_bytes)
        except (OSError, ValueError) as exc:
            raise ValueError('The configured Production workbook/audit pair is missing or unreadable.') from exc
        if audit.get('format_version') != 1 or audit.get('workbook_sha256') != sha(content):
            raise ValueError('Production workbook/audit mismatch. Finish preparation and assess again.')
        return content, audit_bytes, audit

    def _database_state(self, connection=None):
        path = self.settings.testing_state_db_path
        if connection is None and not path.exists():
            return 'absent', {}, None
        own = connection is None
        if own:
            connection = sqlite3.connect(path.resolve().as_uri()+'?mode=ro', uri=True)
            connection.row_factory = sqlite3.Row
            connection.execute('BEGIN')
        try:
            tables = {r[0] for r in connection.execute("SELECT name FROM sqlite_master WHERE type='table'")}
            existing = {}
            material = []
            for table in ('import_batch','device_subject','source_row','canonical_device_record','testing_subjects','testing_events'):
                if table in tables:
                    rows = [dict(r) for r in connection.execute(f'SELECT * FROM {table} ORDER BY id')]
                    material.append((table, rows))
            if 'device_subject' in tables and 'source_row' in tables:
                for row in connection.execute('SELECT ds.subject_key, sr.raw_payload_json FROM device_subject ds LEFT JOIN source_row sr ON sr.id=ds.current_source_row_id'):
                    payload = json.loads(row['raw_payload_json'] or '{}')
                    if 'production_import' not in payload:
                        raise ValueError('Production database contains identities without Production import provenance; reconcile before importing.')
                    existing[row['subject_key']] = {**payload['production_import'], 'prepared_values':payload.get('prepared_values',{})}
            latest = None
            if 'import_batch' in tables:
                row = connection.execute('SELECT * FROM import_batch ORDER BY id DESC LIMIT 1').fetchone()
                latest = dict(row) if row else None
            return sha(dump(material).encode()), existing, latest
        finally:
            if own:
                connection.close()

    def _prepare(self):
        content, audit_bytes, audit = self._pair()
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
        rows, skipped, identities = [], [], set()
        outputs = {tuple(r['key']): r for r in audit.get('output_rows', [])}
        if len(outputs) != len(audit.get('output_rows', [])):
            raise ValueError('Duplicate audit output identities.')
        exports = {tuple(r['key']): r for r in audit.get('export_occurrences', [])}
        if len(exports) != len(audit.get('export_occurrences', [])):
            raise ValueError('Duplicate exported identities in audit.')
        templates = {}
        for occurrence in audit.get('template_occurrences', []):
            templates.setdefault(tuple(occurrence['key']), []).append(occurrence)
        service = CanonicalValidationService()
        service.normalization_repository.rules_for_column = lru_cache(maxsize=None)(service.normalization_repository.rules_for_column)
        try:
            if workbook.sheetnames != ['To Register','Registered','Summary']:
                raise ValueError('Production workbook must contain To Register, Registered, Summary in that order.')
            for name in ('To Register','Registered'):
                sheet = workbook[name]
                iterator = sheet.iter_rows()
                headers = [c.value for c in next(iterator)]
                if any(not isinstance(h,str) for h in headers) or len(set(headers)) != len(headers) or not set(CONTROLS).issubset(headers):
                    raise ValueError(f'Invalid or missing {name} control headers.')
                for number, cells in enumerate(iterator, 2):
                    if not any(c.value is not None for c in cells):
                        continue
                    if any(c.data_type == 'f' for c in cells):
                        raise ValueError(f'Formulas are not permitted in prepared import data: {name}!{number}.')
                    row = dict(zip(headers, [c.value for c in cells]))
                    key = row['Issuing Entity'], row['UDI-DI']
                    if not all(isinstance(v,str) and v.strip()==v and v for v in key) or key in identities:
                        raise ValueError(f'Missing, non-text or duplicate device identity: {name}!{number}.')
                    identities.add(key)
                    recorded = outputs.get(key)
                    if recorded is None or recorded['tab'] != name or recorded['status'] != row['Registration Status']:
                        raise ValueError(f'Workbook classification differs from audit: {name}!{number}.')
                    review = row['Review Required']
                    if review not in ('Yes','No') or (review=='Yes') != bool(recorded['review_reasons']):
                        raise ValueError(f'Invalid review flag or contradictory audit: {name}!{number}.')
                    if review == 'Yes':
                        skipped.append({'issuer':key[0],'udi_di':key[1],'sheet':name,'row':number,'reasons':recorded['review_reasons']})
                        continue
                    status = 'Registered' if name=='Registered' else 'Not registered'
                    if row['Registration Status'] != status:
                        raise ValueError(f'Eligible row has invalid registration classification: {name}!{number}.')
                    snapshot = None
                    export = exports.get(key)
                    if name == 'Registered':
                        if export is None:
                            raise ValueError('Registered row has no exported accepted evidence.')
                        for field,value in export['fields'].items():
                            expected = dump(value) if isinstance(value,(dict,list)) else value
                            # Preparation uses ordinary JSON formatting for structured cells.
                            actual = row.get(field)
                            if isinstance(value,(dict,list)):
                                if json.loads(actual or 'null') != value:
                                    raise ValueError(f'Accepted workbook field differs from audit: {field}.')
                            elif str(actual or '') != str(expected or ''):
                                raise ValueError(f'Accepted workbook field differs from audit: {field}.')
                        snapshot = accepted_snapshot(export,self.settings)
                        if snapshot['device_record']['manufacturer_srn'] != self.settings.eudamed_manufacturer_srn_override:
                            raise ValueError('Accepted export manufacturer differs from configured Production manufacturer.')
                    else:
                        if export is not None or key not in templates:
                            raise ValueError('Not registered classification lacks template-only evidence.')
                        if not all(isinstance(row.get(k),str) and row[k].strip() for k in ('Proposed Basic UDI-DI','Proposed Parent Issuer')):
                            raise ValueError('Eligible template row needs an explicit proposed parent identity.')
                    record = canonical_record(row,self.settings,sheet=name,number=number,snapshot=snapshot,service=service)
                    # Existing downstream catalogue addressing cannot safely distinguish collisions.
                    if not record.catalogue_number:
                        raise ValueError(f'Eligible device lacks a catalogue identity: {key}.')
                    business = {k:v for k,v in row.items() if k.startswith(('Proposed: ','Proposed Parent: ','Accepted Parent: ','Accepted Device: ')) or k in ('Registration Status','Issuing Entity','UDI-DI','Proposed Basic UDI-DI','Proposed Parent Issuer')}
                    fingerprint = sha(dump(semantic(business)).encode())
                    rows.append({'key': key, 'subject_key': 'PROD|'+key[0]+'|'+key[1], 'sheet':name,'number':number,
                                 'values':row,'fingerprint':fingerprint,'record':record,'snapshot':snapshot,
                                 'provenance':{'templates':templates.get(key,[]),'export':export}})
            if identities != set(outputs):
                raise ValueError('Workbook output identities differ from audit.')
        finally:
            workbook.close()
        catalogue_keys = [(r['record'].product_family.casefold(),r['record'].product_variant.casefold(),r['record'].catalogue_number.casefold()) for r in rows]
        catalogue_counts = Counter(catalogue_keys)
        for row,address in zip(rows,catalogue_keys):
            if catalogue_counts[address] > 1:
                row['record'].xml_blockers.append('Multiple UDI-DIs share this model/catalogue address; exact identity selection needs review.')
                row['record'].xml_readiness.status = 'incomplete'

        marker, existing, latest = self._database_state()
        for row in rows:
            record = row['record']
            for key, metadata in existing.items():
                if key != row['subject_key'] and (metadata.get('family'),metadata.get('model'),metadata.get('catalogue')) == (record.product_family,record.product_variant,record.catalogue_number):
                    record.xml_blockers.append('Existing device shares this model/catalogue address; resolve identity selection before XML generation.')
                    record.xml_readiness.status='incomplete'
                    break
        differences = []
        for row in rows:
            previous = existing.get(row['subject_key'])
            if previous is None or previous['fingerprint'] == row['fingerprint']:
                continue
            before, after = previous['prepared_values'], row['values']
            fields = [k for k in set(before) | set(after) if k.startswith(('Proposed: ', 'Proposed Parent: ', 'Accepted Parent: ', 'Accepted Device: ')) or k in ('Registration Status', 'Issuing Entity', 'UDI-DI', 'Proposed Basic UDI-DI', 'Proposed Parent Issuer')]
            differences.append({'issuer':row['key'][0], 'udi_di':row['key'][1],
                                'fields':[{'field':k, 'before':before.get(k), 'after':after.get(k)} for k in sorted(fields) if semantic(before.get(k)) != semantic(after.get(k))]})
        unchanged = sum(r['subject_key'] in existing for r in rows)-len(differences)
        retained = sorted(set(existing)-{r['subject_key'] for r in rows})
        report = {'environment':'prod','workbook':self.workbook.name,'audit':self.audit_path.name,
                  'workbook_sha256':sha(content),'audit_sha256':sha(audit_bytes),
                  'eligible_count':len(rows),'skipped_count':len(skipped),'skipped_rows':skipped,
                  'eligible_by_sheet':dict(Counter(r['sheet'] for r in rows)),
                  'new_count':len(rows)-unchanged-len(differences),'unchanged_count':unchanged,
                  'differences':differences,'retained_count':len(retained),
                  'xml_ready_count':sum(r['record'].xml_readiness.status=='complete' for r in rows),
                  'xml_blocked_count':sum(r['record'].xml_readiness.status!='complete' for r in rows),
                  'xml_blocked_rows':[{'issuer':r['key'][0], 'udi_di':r['key'][1], 'model':r['record'].product_variant, 'catalogue':r['record'].catalogue_number, 'reasons':r['record'].xml_blockers} for r in rows if r['record'].xml_readiness.status!='complete'],
                  'summary':audit.get('summary_entries',[]), 'source_inputs':audit.get('inputs',{}), 'can_import':not differences,
                  'latest_import_batch_id':latest['id'] if latest else None}
        return PreparedImport(sha(content),sha(audit_bytes),audit,rows,skipped,marker,report,time.time()+self.assessment_seconds,self._configuration())

    def _configuration(self):
        return tuple(str(value) for value in (self.settings.environment, self.workbook.resolve(), self.audit_path.resolve(), self.settings.testing_state_db_path.resolve(), self.settings.eudamed_manufacturer_srn_override, self.settings.eudamed_authorised_representative_srn_override, self.settings.eudamed_suppress_authorised_representative, self.settings.eudamed_message_schema_version))

    def assess(self):
        prepared = self._prepare()
        token = secrets.token_urlsafe(32)
        with self._mutex:
            self._pending = type(self)._pending
            for previous in list(self._pending):
                if self._pending[previous].expires_at < time.time():
                    del self._pending[previous]
            # Bound memory: a fresh assessment supersedes old pending content.
            self._pending.clear()
            self._pending[token] = prepared
        return {**prepared.report,'assessment_token':token,'expires_in_seconds':self.assessment_seconds}

    def run_import(self, *, assessment_token=None, imported_by=None, label=None, notes=None):
        with self._mutex:
            prepared = self._pending.pop(assessment_token or '',None)
        if prepared is None or prepared.expires_at < time.time():
            raise ValueError('A current Production assessment confirmation is required.')
        if prepared.configuration != self._configuration():
            raise ValueError('Production configuration changed; assess again.')
        if not prepared.report['can_import']:
            raise ValueError('Existing devices differ; review the differences report before importing.')
        content,audit_bytes,_ = self._pair()
        if (sha(content),sha(audit_bytes)) != (prepared.workbook_hash,prepared.audit_hash):
            raise ValueError('Production files changed; assess again.')
        marker,existing,latest = self._database_state()
        if marker != prepared.database_marker:
            raise ValueError('Database state changed; assess again.')
        if prepared.report['new_count'] == 0 and latest is not None:
            return {**prepared.report, 'import_batch_id':latest['id'],'source_type':'production_workbook',
                    'label':latest['label'],'imported_at':latest['imported_at'],'workbook_count':0,
                    'source_row_count':0,'device_subject_count':0,'created_count':0,'already_imported':True}
        # Schema creation is allowed only on commit, before the atomic data transaction.
        WorkbookImportService()
        store = TestingStateStore()
        imported_at = datetime.now(UTC).isoformat()
        path = self.settings.testing_state_db_path
        with current_pair_lock(self.workbook.parent):
            # Protect against a preparation replacement between file verification and lock.
            if sha(self.workbook.read_bytes()) != prepared.workbook_hash or sha(self.audit_path.read_bytes()) != prepared.audit_hash:
                raise ValueError('Production files changed; assess again.')
            with sqlite3.connect(path) as connection:
                connection.row_factory=sqlite3.Row
                connection.execute('PRAGMA foreign_keys=ON')
                connection.execute('BEGIN IMMEDIATE')
                marker,existing,latest = self._database_state(connection)
                # Schema initialization changes an absent/empty DB marker, not existing operational data.
                if (prepared.database_marker != 'absent' and marker != prepared.database_marker) or (prepared.database_marker == 'absent' and (existing or latest)):
                    raise ValueError('Database changed during confirmation; assess again.')
                report={**prepared.report,'operator_notes':notes,'imported_by':imported_by,'imported_at':imported_at}
                batch=connection.execute('INSERT INTO import_batch(source_type,label,imported_at,imported_by,notes) VALUES(?,?,?,?,?)',
                    ('production_workbook',label or 'Production workbook import',imported_at,imported_by,dump(report))).lastrowid
                workbook_id=connection.execute('INSERT INTO source_workbook(import_batch_id,workbook_name,file_path,file_hash,loaded_at) VALUES(?,?,?,?,?)',
                    (batch,self.workbook.name,str(self.workbook),prepared.workbook_hash,imported_at)).lastrowid
                created=0
                for row in prepared.rows:
                    if row['subject_key'] in existing:
                        continue
                    record=row['record']
                    metadata={'fingerprint':row['fingerprint'],'issuer':row['key'][0],'udi_di':row['key'][1],
                              'registration_status':row['values']['Registration Status'],'family':record.product_family,'model':record.product_variant,'catalogue':record.catalogue_number,'workbook_sha256':prepared.workbook_hash,
                              'audit_sha256':prepared.audit_hash,'provenance':row['provenance']}
                    payload={'production_import':metadata,'prepared_values':row['values']}
                    source=connection.execute('INSERT INTO source_row(source_workbook_id,sheet_name,row_index,product_family,product_variant,catalogue_number,primary_udi_di,submission_operation,raw_payload_json,canonical_status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                        (workbook_id,row['sheet'],row['number'],record.product_family,record.product_variant,record.catalogue_number,record.primary_udi_di,record.submission_operation,dump(payload),'xml_ready' if not record.xml_blockers else 'xml_blocked',imported_at)).lastrowid
                    parent = row['values'].get('Proposed Basic UDI-DI') or row['values'].get('Accepted Basic UDI-DI')
                    subject=connection.execute('INSERT INTO device_subject(subject_key,product_family,product_variant,catalogue_number,primary_udi_di,basic_udi_di,current_source_row_id,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)',
                        (row['subject_key'],record.product_family,record.product_variant,record.catalogue_number,record.primary_udi_di,parent,source,imported_at,imported_at)).lastrowid
                    canonical=connection.execute('INSERT INTO canonical_device_record(device_subject_id,source_row_id,source_import_batch_id,canonical_version,canonical_status,completeness_status,xml_readiness_status,xml_ready,record_json,completeness_json,blockers_json) VALUES(?,?,?,?,?,?,?,?,?,?,?)',
                        (subject,source,batch,1,'xml_ready' if not record.xml_blockers else 'xml_blocked',record.completeness.status,record.xml_readiness.status,int(not record.xml_blockers),dump(record.model_dump(mode='json')),dump({'completeness':record.completeness.model_dump(),'xml_readiness':record.xml_readiness.model_dump()}),dump({'blockers':record.blockers,'xml_blockers':record.xml_blockers}))).lastrowid
                    for field in record.fields:
                        connection.execute('INSERT INTO canonical_field_value(canonical_device_record_id,canonical_path,field_status,required,value_json,source_headers_json,review_note) VALUES(?,?,?,?,?,?,?)',
                            (canonical,field.canonical_path,'populated' if field.value is not None else 'missing',int(field.required),dump(field.value) if field.value is not None else None,dump([field.source_detail]),field.review_note))
                    if row['snapshot'] is not None:
                        self._baseline(connection,store,row,subject,batch,imported_at)
                    created+=1
                # Keep a complete current projection, retaining absent/skipped existing devices.
                connection.execute('UPDATE canonical_device_record SET source_import_batch_id=?',(batch,))
                self._snapshot(connection,batch)
                report.update(import_batch_id=batch,created_count=created)
                connection.execute('UPDATE import_batch SET notes=? WHERE id=?',(dump(report),batch))
        return {**report,'source_type':'production_workbook','label':label or 'Production workbook import',
                'workbook_count':1,'source_row_count':created,'device_subject_count':created,'already_imported':False}

    @staticmethod
    def _baseline(connection,store,row,subject,batch,imported_at):
        record=row['record']; snapshot=row['snapshot']; full=snapshot['device_record']
        from app.services.identity import normalize_identity
        testing=connection.execute("INSERT INTO testing_subjects(subject_key,device_subject_id,normalized_product_family,normalized_product_variant,normalized_catalogue_number,normalized_primary_udi_di,normalized_basic_udi_di,product_family,product_variant,catalogue_number,primary_udi_di,basic_udi_di) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
            (row['subject_key'],subject,normalize_identity(record.product_family),normalize_identity(record.product_variant),normalize_identity(record.catalogue_number),normalize_identity(record.primary_udi_di),normalize_identity(full['basic_identifier_code']),record.product_family,record.product_variant,record.catalogue_number,record.primary_udi_di,full['basic_identifier_code'])).lastrowid
        snapshot={**snapshot,'device_record':{**full,'product_family':record.product_family,'product_variant':record.product_variant},
                  'import_batch_id':batch,'accepted_state_source':'production_export'}
        event=connection.execute("INSERT INTO testing_events(subject_id,event_index,message_type,status,event_kind,accepted_state_source,version,source_file_name,state_after_json,raw_event_json) VALUES(?,1,'PRODUCTION_EXPORT.SNAPSHOT','IMPORTED','BASELINE_IMPORT','production_export',?,?,?,?)",
            (testing,snapshot['version'],snapshot['export_source']['file'],dump(snapshot),dump({'import_batch_id':batch,'issuer':row['key'][0],'udi_di':row['key'][1],'parent_issuer':full['basic_identifier_issuing_entity'],'parent':full['basic_identifier_code'],'export_source':snapshot['export_source']}))).lastrowid
        market={'version':snapshot['market_info_version'],'market_countries':snapshot['market_countries'],'accepted_state_source':'production_export'}
        connection.execute("UPDATE testing_subjects SET device_subject_id=?,registration_status='child_registered',latest_successful_version=?,latest_successful_state_json=?,latest_successful_market_info_version=?,latest_successful_market_info_state_json=? WHERE id=?",
            (subject,snapshot['version'],dump(snapshot),market['version'],dump(market),testing))

    @staticmethod
    def _snapshot(connection,batch):
        records=[json.loads(r[0]) for r in connection.execute('SELECT record_json FROM canonical_device_record')]
        total=len(records);ready=sum(r['completeness']['status']=='complete' for r in records);xml_ready=sum(r['xml_readiness']['status']=='complete' for r in records)
        fields=records[0]['fields'] if records else []
        connection.execute('INSERT INTO canonical_projection_snapshot(source_import_batch_id,family_scope,scope_note,validation_note,total_source_records,validation_subset_records,excluded_records,matched_reference_records,tracked_required_fields,tracked_xml_required_fields,ready_records,blocked_records,xml_ready_records,xml_blocked_records) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',
            (batch,'Production devices','Current complete projection includes retained earlier devices.','Prepared workbook eligibility is separate from XML readiness.',total,total,0,total,sum(f['required'] for f in fields),sum(f['xml_required'] for f in fields),ready,total-ready,xml_ready,total-xml_ready))
